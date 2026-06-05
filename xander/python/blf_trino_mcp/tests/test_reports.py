from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from blf_trino_mcp.reports import (
    export_nl_query_to_excel,
    export_sql_to_excel,
    generate_bi_report,
)


class FakeTrinoClient:
    def query(self, sql: str) -> dict[str, Any]:
        if "sales" in sql:
            return {
                "columns": ["dt", "amount"],
                "rows": [("20260601", 12.5), ("20260602", 16.0)],
            }
        return {"columns": ["id", "name"], "rows": [(1, "a"), (2, "b")]}


def test_export_sql_to_excel_writes_workbook(tmp_path: Path) -> None:
    result = export_sql_to_excel(
        FakeTrinoClient(),
        export_dir=str(tmp_path),
        public_base_url="http://localhost:9011",
        sql="select id, name from t",
        file_name="orders",
    )

    assert result["success"] is True
    path = Path(result["summary"]["file_path"])
    assert path.exists()
    assert result["summary"]["download_url"].startswith("http://localhost:9011/files/")
    workbook = load_workbook(path)
    assert "summary" in workbook.sheetnames
    assert "data" in workbook.sheetnames
    assert workbook["data"]["A1"].value == "id"
    assert workbook["data"]["B2"].value == "a"


def test_export_nl_query_to_excel_requires_generated_sql(tmp_path: Path) -> None:
    result = export_nl_query_to_excel(
        FakeTrinoClient(),
        export_dir=str(tmp_path),
        public_base_url="http://localhost:9011",
        question="查订单",
        generated_sql="",
    )

    assert result["success"] is False
    assert result["error_type"] == "invalid_input"


def test_generate_bi_report_writes_excel_and_html(tmp_path: Path) -> None:
    result = generate_bi_report(
        FakeTrinoClient(),
        export_dir=str(tmp_path),
        public_base_url="http://localhost:9011",
        report_title="sales_report",
        queries=[{"name": "sales", "sql": "select dt, amount from sales"}],
        include_html=True,
    )

    assert result["success"] is True
    excel_path = Path(result["summary"]["excel"]["file_path"])
    html_path = Path(result["summary"]["html"]["file_path"])
    assert excel_path.exists()
    assert html_path.exists()
    assert "sales" in load_workbook(excel_path).sheetnames
    assert "sales_report" in html_path.read_text(encoding="utf-8")
