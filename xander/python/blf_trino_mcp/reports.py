"""Excel and HTML report generation for BLF Trino MCP."""

from __future__ import annotations

import html
import json
import re
import uuid
from datetime import datetime
from pathlib import Path
from typing import Any

from .hive import MAX_QUERY_ROWS
from .tools import _json_safe_value, _query_with_limit
from .trino_client import TrinoClient

MAX_EXPORT_ROWS = MAX_QUERY_ROWS
_SAFE_NAME_RE = re.compile(r"[^a-zA-Z0-9_.-]+")


def export_sql_to_excel(
    client: TrinoClient,
    *,
    export_dir: str,
    public_base_url: str,
    sql: str,
    row_limit: int = 10000,
    file_name: str = "",
    sheet_name: str = "data",
    include_summary: bool = True,
) -> dict[str, Any]:
    """Run one SQL query and export the result to an Excel file."""
    try:
        query = _query_with_limit(
            client,
            sql,
            row_limit=min(max(int(row_limit), 1), MAX_EXPORT_ROWS),
        )
        export = _write_excel_report(
            export_dir=export_dir,
            public_base_url=public_base_url,
            report_title=file_name or "trino_export",
            datasets=[
                {
                    "name": sheet_name,
                    "sql": query["sql"],
                    "columns": query["columns"],
                    "rows": query["rows"],
                    "row_count": query["row_count"],
                    "truncated": query["truncated"],
                }
            ],
            include_summary=include_summary,
            file_name=file_name,
        )
        return {
            "success": True,
            "summary": {
                **export,
                "query": {
                    "sql": query["sql"],
                    "row_count": query["row_count"],
                    "truncated": query["truncated"],
                },
            },
            "risks": ["结果已按 row_limit 截断"] if query["truncated"] else [],
            "evidence": {"interface": "Trino SQL export to Excel"},
        }
    except Exception as exc:
        return _error_response(exc, sql=sql)


def export_nl_query_to_excel(
    client: TrinoClient,
    *,
    export_dir: str,
    public_base_url: str,
    question: str,
    generated_sql: str,
    row_limit: int = 10000,
    file_name: str = "",
) -> dict[str, Any]:
    """Export an Agent-generated SQL query to Excel."""
    try:
        if not generated_sql.strip():
            raise ValueError("generated_sql is required for Excel export")
        result = export_sql_to_excel(
            client,
            export_dir=export_dir,
            public_base_url=public_base_url,
            sql=generated_sql,
            row_limit=row_limit,
            file_name=file_name or _short_file_stem(question),
            sheet_name="data",
            include_summary=True,
        )
        result.setdefault("summary", {})["question"] = question
        result.setdefault("evidence", {})["sql_source"] = "generated_sql"
        return result
    except Exception as exc:
        return _error_response(exc, question=question, generated_sql=generated_sql)


def generate_bi_report(
    client: TrinoClient,
    *,
    export_dir: str,
    public_base_url: str,
    report_title: str,
    queries: list[dict[str, Any]],
    row_limit: int = 10000,
    include_html: bool = True,
) -> dict[str, Any]:
    """Run multiple SQL queries and generate an Excel workbook plus optional HTML report."""
    try:
        if not queries:
            raise ValueError("queries cannot be empty")
        datasets = []
        risks = []
        max_rows = min(max(int(row_limit), 1), MAX_EXPORT_ROWS)
        for index, item in enumerate(queries, start=1):
            sql = str(item.get("sql") or "").strip()
            if not sql:
                raise ValueError(f"queries[{index}].sql cannot be empty")
            name = str(item.get("name") or f"query_{index}")
            query = _query_with_limit(client, sql, row_limit=max_rows)
            if query["truncated"]:
                risks.append(f"{name} 结果已按 row_limit 截断")
            datasets.append(
                {
                    "name": name,
                    "sql": query["sql"],
                    "columns": query["columns"],
                    "rows": query["rows"],
                    "row_count": query["row_count"],
                    "truncated": query["truncated"],
                }
            )
        excel = _write_excel_report(
            export_dir=export_dir,
            public_base_url=public_base_url,
            report_title=report_title,
            datasets=datasets,
            include_summary=True,
            file_name=report_title,
        )
        html_report = (
            _write_html_report(
                export_dir=export_dir,
                public_base_url=public_base_url,
                report_title=report_title,
                datasets=datasets,
            )
            if include_html
            else None
        )
        return {
            "success": True,
            "summary": {
                "report_title": report_title,
                "excel": excel,
                "html": html_report,
                "datasets": [
                    {
                        "name": dataset["name"],
                        "row_count": dataset["row_count"],
                        "truncated": dataset["truncated"],
                    }
                    for dataset in datasets
                ],
            },
            "risks": risks,
            "evidence": {"interface": "Trino BI report export"},
        }
    except Exception as exc:
        return _error_response(exc, report_title=report_title)


def _write_excel_report(
    *,
    export_dir: str,
    public_base_url: str,
    report_title: str,
    datasets: list[dict[str, Any]],
    include_summary: bool,
    file_name: str = "",
) -> dict[str, Any]:
    try:
        from openpyxl import Workbook
        from openpyxl.chart import BarChart, Reference
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise RuntimeError(
            "Missing dependency: openpyxl. Install it with: python3 -m pip install openpyxl"
        ) from exc

    export_path = Path(export_dir)
    export_path.mkdir(parents=True, exist_ok=True)
    stem = _unique_stem(file_name or report_title)
    workbook_path = export_path / f"{stem}.xlsx"

    workbook = Workbook()
    if include_summary:
        summary = workbook.active
        summary.title = "summary"
        _write_summary_sheet(summary, report_title, datasets)
    else:
        workbook.remove(workbook.active)

    for dataset in datasets:
        sheet = workbook.create_sheet(_sheet_title(str(dataset["name"])))
        _write_data_sheet(
            sheet,
            columns=[str(item) for item in dataset["columns"]],
            rows=dataset["rows"],
            get_column_letter=get_column_letter,
            font_cls=Font,
            fill_cls=PatternFill,
            alignment_cls=Alignment,
        )

    if datasets:
        chart_sheet = workbook.create_sheet("charts")
        _write_chart_sheet(
            chart_sheet,
            workbook[_sheet_title(str(datasets[0]["name"]))],
            BarChart,
            Reference,
        )

    workbook.save(workbook_path)
    return _file_payload(workbook_path, public_base_url)


def _write_summary_sheet(sheet: Any, report_title: str, datasets: list[dict[str, Any]]) -> None:
    from openpyxl.styles import Font, PatternFill

    sheet["A1"] = report_title
    sheet["A1"].font = Font(bold=True, size=16)
    sheet["A2"] = "生成时间"
    sheet["B2"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    sheet.append([])
    sheet.append(["数据集", "行数", "是否截断", "SQL"])
    for cell in sheet[4]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
    for dataset in datasets:
        sheet.append(
            [
                dataset["name"],
                dataset["row_count"],
                "是" if dataset["truncated"] else "否",
                dataset["sql"],
            ]
        )
    sheet.column_dimensions["A"].width = 24
    sheet.column_dimensions["B"].width = 12
    sheet.column_dimensions["C"].width = 12
    sheet.column_dimensions["D"].width = 100


def _write_data_sheet(
    sheet: Any,
    *,
    columns: list[str],
    rows: list[dict[str, Any]],
    get_column_letter: Any,
    font_cls: Any,
    fill_cls: Any,
    alignment_cls: Any,
) -> None:
    header_fill = fill_cls("solid", fgColor="1F4E78")
    header_font = font_cls(bold=True, color="FFFFFF")
    sheet.append(columns or ["row"])
    for cell in sheet[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = alignment_cls(horizontal="center")
    for row in rows:
        if columns:
            sheet.append([_excel_cell_value(row.get(column)) for column in columns])
        else:
            sheet.append([_excel_cell_value(row.get("row"))])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, column in enumerate(columns or ["row"], start=1):
        letter = get_column_letter(index)
        max_len = max([len(str(column))] + [len(str(row.get(column, ""))) for row in rows[:200]])
        sheet.column_dimensions[letter].width = min(max(max_len + 2, 10), 60)


def _write_chart_sheet(chart_sheet: Any, data_sheet: Any, chart_cls: Any, reference_cls: Any) -> None:
    if data_sheet.max_row < 2 or data_sheet.max_column < 2:
        chart_sheet["A1"] = "没有足够的数据生成图表"
        return
    numeric_col = None
    for col in range(2, data_sheet.max_column + 1):
        values = [data_sheet.cell(row=row, column=col).value for row in range(2, min(data_sheet.max_row, 20) + 1)]
        if any(isinstance(value, (int, float)) for value in values):
            numeric_col = col
            break
    if numeric_col is None:
        chart_sheet["A1"] = "未找到可用于图表的数值列"
        return
    chart = chart_cls()
    chart.type = "bar"
    chart.title = f"{data_sheet.title} 概览"
    chart.y_axis.title = str(data_sheet.cell(row=1, column=numeric_col).value)
    chart.x_axis.title = str(data_sheet.cell(row=1, column=1).value)
    max_row = min(data_sheet.max_row, 20)
    data = reference_cls(data_sheet, min_col=numeric_col, min_row=1, max_row=max_row)
    categories = reference_cls(data_sheet, min_col=1, min_row=2, max_row=max_row)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    chart_sheet.add_chart(chart, "A1")


def _write_html_report(
    *,
    export_dir: str,
    public_base_url: str,
    report_title: str,
    datasets: list[dict[str, Any]],
) -> dict[str, Any]:
    export_path = Path(export_dir)
    export_path.mkdir(parents=True, exist_ok=True)
    html_path = export_path / f"{_unique_stem(report_title)}.html"
    parts = [
        "<!doctype html><html><head><meta charset='utf-8'>",
        f"<title>{html.escape(report_title)}</title>",
        "<style>body{font-family:Arial,sans-serif;margin:24px;color:#172033}"
        "table{border-collapse:collapse;width:100%;margin:16px 0}"
        "th,td{border:1px solid #d8dee9;padding:6px 8px;font-size:13px}"
        "th{background:#1f4e78;color:white;text-align:left}"
        ".meta{color:#5b6472}.sql{white-space:pre-wrap;background:#f6f8fa;padding:12px}</style>",
        "</head><body>",
        f"<h1>{html.escape(report_title)}</h1>",
        f"<p class='meta'>生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}</p>",
    ]
    for dataset in datasets:
        parts.append(f"<h2>{html.escape(str(dataset['name']))}</h2>")
        parts.append(
            f"<p class='meta'>行数：{dataset['row_count']}，截断：{'是' if dataset['truncated'] else '否'}</p>"
        )
        parts.append(f"<div class='sql'>{html.escape(str(dataset['sql']))}</div>")
        parts.append(_html_table(dataset["columns"], dataset["rows"][:200]))
    parts.append("</body></html>")
    html_path.write_text("\n".join(parts), encoding="utf-8")
    return _file_payload(html_path, public_base_url)


def _html_table(columns: list[str], rows: list[dict[str, Any]]) -> str:
    safe_columns = columns or ["row"]
    parts = ["<table><thead><tr>"]
    parts.extend(f"<th>{html.escape(str(column))}</th>" for column in safe_columns)
    parts.append("</tr></thead><tbody>")
    for row in rows:
        parts.append("<tr>")
        for column in safe_columns:
            parts.append(f"<td>{html.escape(str(row.get(column, '')))}</td>")
        parts.append("</tr>")
    parts.append("</tbody></table>")
    return "".join(parts)


def _excel_cell_value(value: Any) -> Any:
    if isinstance(value, (list, dict)):
        return json.dumps(value, ensure_ascii=False)
    return value


def _file_payload(path: Path, public_base_url: str) -> dict[str, Any]:
    return {
        "file_name": path.name,
        "file_path": str(path),
        "download_url": f"{public_base_url.rstrip('/')}/files/{path.name}",
        "bytes": path.stat().st_size,
    }


def _unique_stem(value: str) -> str:
    base = _short_file_stem(value)
    return f"{base}_{datetime.now().strftime('%Y%m%d_%H%M%S')}_{uuid.uuid4().hex[:8]}"


def _short_file_stem(value: str) -> str:
    cleaned = _SAFE_NAME_RE.sub("_", (value or "report").strip()).strip("._")
    return (cleaned or "report")[:60]


def _sheet_title(value: str) -> str:
    title = re.sub(r"[][\\\\:*?/]", "_", (value or "data").strip())[:31]
    return title or "data"


def _error_response(error: Exception, **extra: Any) -> dict[str, Any]:
    return {
        "success": False,
        "error_type": "report_error" if not isinstance(error, ValueError) else "invalid_input",
        "message": str(error),
        **extra,
    }
