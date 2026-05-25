#!/usr/bin/env python3
"""Check Hive dataset availability and update data_availability_flag.

The script is designed for Jenkins jobs that pass ``TABLE_NAMES``. It is
read-mostly: failed checks only produce reasons in logs/reports; successful
checks merge availability flags into ``blf.data.warehouse.data_availability_flag``.
"""

from __future__ import annotations

import argparse
import json
import os
import re
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any, Iterable, Optional

from .field_lineage_datahub_reader import (
    fetch_structured_properties,
    make_hive_dataset_urn,
    strip_markdown_code_fence,
)
from .query_upstream_lineage import urn_to_table_name
from .structured_properties import (
    URN_DATA_AVAILABILITY_FLAG,
    URN_ETL_SCRIPT,
    URN_EXECUTE_SHELL,
    URN_SCHEDULE_URL,
)
from .table_documentation_full_discovery import (
    _parse_tab_rows,
    _run_mysql_query,
    parse_hive_dataset_urn,
    table_name_matches_prefix,
)

FLAG_DDL = "DDL"
FLAG_TABLE_LINEAGE = "表血缘"
FLAG_FIELD_LINEAGE = "字段血缘"
FLAG_ORDER = [FLAG_DDL, FLAG_TABLE_LINEAGE, FLAG_FIELD_LINEAGE]

LABELS = {
    URN_ETL_SCRIPT: "Etl Script",
    URN_SCHEDULE_URL: "Schedule URL",
    URN_EXECUTE_SHELL: "Execute Shell",
}

STRUCTURED_PROPERTY_RULES = {
    URN_ETL_SCRIPT: "structuredProperties: Etl Script",
    URN_SCHEDULE_URL: "structuredProperties: Schedule URL",
    URN_EXECUTE_SHELL: "structuredProperties: Execute Shell",
}


@dataclass
class AvailabilityResult:
    table_name: str
    dataset_urn: str
    dataset_type: str
    status: str = "OK"
    passed_flags: set[str] = field(default_factory=set)
    existing_flags: set[str] = field(default_factory=set)
    final_flags: list[str] = field(default_factory=list)
    reason: str = ""
    write_status: str = "SKIP"
    lineage_documented_upstreams: list[str] = field(default_factory=list)
    lineage_existing_upstreams: list[str] = field(default_factory=list)
    lineage_missing_upstreams: list[str] = field(default_factory=list)
    lineage_extra_upstreams: list[str] = field(default_factory=list)
    error: str = ""


def parse_table_names(raw: str) -> list[str]:
    out: list[str] = []
    seen: set[str] = set()
    for line in raw.replace("\r", "\n").splitlines():
        body = line.split("#", 1)[0].strip()
        if not body:
            continue
        for item in body.split(","):
            name = item.strip().strip("'\"").lower()
            if not name:
                continue
            if "." not in name:
                name = f"default.{name}"
            if name not in seen:
                seen.add(name)
                out.append(name)
    return out


def discover_table_names_by_prefix_from_mysql(
    table_prefix: str,
    *,
    platform_instance: str,
    env: str,
) -> list[str]:
    prefix = table_prefix.strip()
    if not prefix:
        return []
    sql = (
        "select distinct urn "
        "from metadata_aspect_v2 "
        "where version=0 "
        "and urn like 'urn:li:dataset:(urn:li:dataPlatform:hive,%' "
        "order by urn"
    )
    table_names: set[str] = set()
    for (urn,) in _parse_tab_rows(_run_mysql_query(sql), 1):
        table_name = parse_hive_dataset_urn(urn, platform_instance=platform_instance, env=env)
        if table_name and table_name_matches_prefix(table_name, prefix):
            table_names.add(table_name)
    return sorted(table_names)


def is_meaningful_text(value: str) -> bool:
    text = strip_markdown_code_fence(value or "").strip()
    return bool(text) and text.lower() not in {"无", "null", "none"}


def _iter_property_assignments(payload: dict[str, Any]) -> Iterable[dict[str, Any]]:
    current: Any = payload
    for key in ("structuredProperties", "value"):
        if isinstance(current, dict) and key in current:
            current = current[key]
    if isinstance(current, dict) and isinstance(current.get("properties"), list):
        for item in current["properties"]:
            if isinstance(item, dict):
                yield item


def _string_values(assignment: dict[str, Any]) -> list[str]:
    values = assignment.get("values")
    if not isinstance(values, list):
        return []
    out: list[str] = []
    for value in values:
        if isinstance(value, dict) and isinstance(value.get("string"), str):
            out.append(value["string"])
        elif isinstance(value, str):
            out.append(value)
    return out


def extract_structured_values(payload: dict[str, Any]) -> dict[str, list[str]]:
    by_urn: dict[str, list[str]] = {}
    for assignment in _iter_property_assignments(payload):
        property_urn = assignment.get("propertyUrn")
        if isinstance(property_urn, str):
            by_urn[property_urn] = _string_values(assignment)
    return by_urn


def first_structured_value(values: dict[str, list[str]], property_urn: str) -> str:
    for value in values.get(property_urn, []):
        if isinstance(value, str):
            return value
    return ""


def extract_existing_flags(values: dict[str, list[str]]) -> set[str]:
    flags: set[str] = set()
    for value in values.get(URN_DATA_AVAILABILITY_FLAG, []):
        for token in re.split(r"[,，/、\n]+", value):
            token = token.strip()
            if token:
                flags.add(token)
    return flags


def merge_availability_flags(existing: Iterable[str], passed: Iterable[str]) -> list[str]:
    merged = {v for v in existing if v}
    merged.update(v for v in passed if v)
    ordered = [flag for flag in FLAG_ORDER if flag in merged]
    ordered.extend(sorted(merged - set(ordered)))
    return ordered


def _aspect_url(gms_url: str, dataset_urn: str, aspect_name: str) -> str:
    encoded = urllib.parse.quote(dataset_urn, safe="")
    return f"{gms_url.rstrip('/')}/openapi/v3/entity/dataset/{encoded}/{aspect_name}"


def fetch_aspect_payload(
    gms_url: str,
    dataset_urn: str,
    aspect_name: str,
    token: Optional[str] = None,
    timeout_sec: int = 60,
) -> dict[str, Any]:
    req = urllib.request.Request(
        _aspect_url(gms_url, dataset_urn, aspect_name),
        method="GET",
        headers={"Accept": "application/json"},
    )
    if token:
        req.add_header("Authorization", f"Bearer {token}")
    try:
        with urllib.request.urlopen(req, timeout=timeout_sec) as resp:
            return json.loads(resp.read().decode("utf-8"))
    except urllib.error.HTTPError as exc:
        if exc.code == 404:
            return {}
        detail = exc.read().decode("utf-8", errors="replace")
        raise RuntimeError(f"GET {aspect_name} HTTP {exc.code}: {detail}") from exc


def fetch_structured_properties_or_empty(
    gms_url: str,
    dataset_urn: str,
    token: Optional[str] = None,
) -> dict[str, Any]:
    try:
        return fetch_structured_properties(gms_url, dataset_urn, token=token)
    except RuntimeError as exc:
        if "HTTP 404" in str(exc):
            return {"structuredProperties": {"value": {"properties": []}}}
        raise


def unwrap_aspect(payload: dict[str, Any], aspect_name: str) -> dict[str, Any]:
    current: Any = payload
    for key in (aspect_name, "value"):
        if isinstance(current, dict) and key in current:
            current = current[key]
    return current if isinstance(current, dict) else {}


def schema_field_count(payload: dict[str, Any]) -> int:
    aspect = unwrap_aspect(payload, "schemaMetadata")
    fields = aspect.get("fields")
    return len(fields) if isinstance(fields, list) else 0


def view_logic_from_payload(payload: dict[str, Any]) -> str:
    aspect = unwrap_aspect(payload, "viewProperties")
    view_logic = aspect.get("viewLogic")
    return view_logic if isinstance(view_logic, str) else ""


def editable_description_from_payload(payload: dict[str, Any]) -> str:
    aspect = unwrap_aspect(payload, "editableDatasetProperties")
    description = aspect.get("description")
    return description if isinstance(description, str) else ""


def is_deprecated_from_payload(payload: dict[str, Any], aspect_name: str = "deprecation") -> bool:
    aspect = unwrap_aspect(payload, aspect_name)
    return aspect.get("deprecated") is True


def upstream_tables_from_payload(
    payload: dict[str, Any],
    *,
    platform_instance: str,
) -> set[str]:
    aspect = unwrap_aspect(payload, "upstreamLineage")
    upstreams = aspect.get("upstreams")
    out: set[str] = set()
    if not isinstance(upstreams, list):
        return out
    for upstream in upstreams:
        if not isinstance(upstream, dict):
            continue
        dataset_urn = upstream.get("dataset")
        if isinstance(dataset_urn, str) and dataset_urn:
            out.add(urn_to_table_name(dataset_urn, platform_instance).lower())
    return out


def extract_data_source_section(description: str) -> str:
    lines = (description or "").splitlines()
    start: Optional[int] = None
    for idx, line in enumerate(lines):
        if re.match(r"^\s{0,3}#{1,6}\s*4[.、]\s*数据来源\s*$", line.strip()):
            start = idx
            break
    if start is None:
        for idx, line in enumerate(lines):
            if re.search(r"4[.、]\s*数据来源", line):
                start = idx
                break
    if start is None:
        return ""
    end = len(lines)
    for idx in range(start + 1, len(lines)):
        if re.match(r"^\s{0,3}#{1,6}\s*\d+[.、]\s+", lines[idx].strip()):
            end = idx
            break
    return "\n".join(lines[start:end]).strip()


def _clean_table_token(value: str) -> str:
    text = value.strip().strip("`").strip()
    text = re.sub(r"<br\s*/?>", "\n", text, flags=re.I)
    text = re.sub(r"\s+", " ", text)
    return text.lower()


def _normalize_table_token(value: str) -> str:
    token = _clean_table_token(value)
    if not token:
        return ""
    if "." in token:
        db_name, table_name = token.rsplit(".", 1)
        if table_name.startswith("not_verified_"):
            table_name = table_name.removeprefix("not_verified_")
        return f"{db_name}.{table_name}"
    if token.startswith("not_verified_"):
        return token.removeprefix("not_verified_")
    return token


def parse_documented_upstreams(description: str) -> tuple[bool, set[str]]:
    section = extract_data_source_section(description)
    if not section:
        return False, set()

    upstreams: set[str] = set()
    for raw_line in section.splitlines():
        line = raw_line.strip()
        candidates: list[str] = []
        if line.startswith("|"):
            cells = [cell.strip() for cell in line.strip("|").split("|")]
            if cells:
                first = _clean_table_token(cells[0])
                if first and first not in {"---", "上游表"} and not set(first) <= {"-", " "}:
                    candidates.append(first)
        else:
            match = re.match(r"^(?:[-*+]|\d+[.、])\s+`([^`]+)`", line)
            if match:
                candidates.append(_clean_table_token(match.group(1)))

        for candidate in candidates:
            for token in re.split(r"[,，、\n]+", candidate):
                token = _normalize_table_token(token)
                if token:
                    upstreams.add(token)

    return True, upstreams


def resolve_documented_upstreams(
    documented: set[str],
    existing: set[str],
) -> tuple[set[str], list[str]]:
    by_table: dict[str, set[str]] = {}
    for table_name in existing:
        by_table.setdefault(table_name.rsplit(".", 1)[-1], set()).add(table_name)

    resolved: set[str] = set()
    errors: list[str] = []
    for ref in sorted(documented):
        if "." in ref:
            resolved.add(ref)
            continue
        matches = sorted(by_table.get(ref, set()))
        if len(matches) == 1:
            resolved.add(matches[0])
        elif len(matches) > 1:
            errors.append(f"Documentation 上游表 {ref} 无库名且匹配多个 DataHub upstream: {', '.join(matches)}")
        else:
            resolved.add(ref)
    return resolved, errors


def evaluate_table_lineage(
    documentation: str,
    upstreams: set[str],
    target_table_name: str = "",
) -> tuple[bool, set[str], list[str], list[str], list[str], str]:
    section_found, documented_raw = parse_documented_upstreams(documentation)
    if not section_found:
        return (
            False,
            set(),
            [],
            sorted(upstreams),
            [],
            '表血缘: FAIL [LINEAGE-DOC-SECTION] editableDatasetProperties.description 缺少 "4. 数据来源" 小节',
        )
    if not documented_raw:
        return (
            False,
            set(),
            [],
            sorted(upstreams),
            [],
            '表血缘: FAIL [LINEAGE-DOC-PARSE] editableDatasetProperties.description 的 "4. 数据来源" 小节未解析到上游表；支持表格第一列或编号列表中的反引号表名',
        )
    documented, errors = resolve_documented_upstreams(documented_raw, upstreams)
    normalized_target = _normalize_table_token(target_table_name)
    if normalized_target:
        target_table = normalized_target.rsplit(".", 1)[-1]
        documented = _exclude_target_table(documented, normalized_target, target_table)
        upstreams = _exclude_target_table(upstreams, normalized_target, target_table)
    if errors:
        return (
            False,
            documented,
            [],
            [],
            errors,
            "表血缘: FAIL [LINEAGE-DOC-AMBIGUOUS] Documentation 上游表存在无库名歧义；请在 4. 数据来源 中补全 db.table",
        )
    missing = sorted(documented - upstreams)
    extra = sorted(upstreams - documented)
    if missing or extra:
        reason = (
            "表血缘: FAIL [LINEAGE-DIFF] Documentation 的 4. 数据来源 与 DataHub upstreamLineage 不一致 "
            f"missing_in_datahub={missing} extra_in_datahub={extra}"
        )
        return False, documented, missing, extra, [], reason
    return True, documented, [], [], [], "表血缘: PASS"


def _exclude_target_table(
    table_names: set[str],
    normalized_target: str,
    target_table: str,
) -> set[str]:
    return {
        table_name
        for table_name in table_names
        if table_name != normalized_target and table_name.rsplit(".", 1)[-1] != target_table
    }


def evaluate_dataset_availability(
    *,
    table_name: str,
    dataset_urn: str,
    is_view: bool,
    is_deprecated: bool = False,
    structured_values: dict[str, str] | dict[str, list[str]],
    schema_field_count: int,
    view_logic: str,
    documentation: str,
    upstreams: set[str],
    existing_flags: set[str],
) -> AvailabilityResult:
    dataset_type = "view" if is_view else "table"
    normalized_target = _normalize_table_token(table_name)
    filtered_upstreams = (
        _exclude_target_table(upstreams, normalized_target, normalized_target.rsplit(".", 1)[-1])
        if normalized_target
        else upstreams
    )
    result = AvailabilityResult(
        table_name=table_name,
        dataset_urn=dataset_urn,
        dataset_type=dataset_type,
        existing_flags=set(existing_flags),
        lineage_existing_upstreams=sorted(filtered_upstreams),
    )
    reasons: list[str] = []

    if is_deprecated:
        result.passed_flags.update({FLAG_DDL, FLAG_TABLE_LINEAGE})
        result.final_flags = merge_availability_flags(existing_flags, result.passed_flags)
        result.reason = "Dataset 已废弃，跳过可用性检查"
        return result

    def _value(urn: str) -> str:
        raw = structured_values.get(urn, "")  # type: ignore[arg-type]
        if isinstance(raw, list):
            return raw[0] if raw else ""
        return raw

    if is_view:
        if is_meaningful_text(view_logic):
            result.passed_flags.add(FLAG_DDL)
            reasons.append("DDL: PASS")
        else:
            reasons.append("DDL: FAIL [DDL-VIEW-LOGIC] viewProperties.viewLogic 缺失或为空")

        if upstreams:
            result.passed_flags.add(FLAG_TABLE_LINEAGE)
            reasons.append("表血缘: PASS")
        else:
            reasons.append("表血缘: FAIL [LINEAGE-UPSTREAM-ASPECT] view 的 upstreamLineage.upstreams 为空")
    else:
        missing_props = [
            f"{STRUCTURED_PROPERTY_RULES[urn]} ({urn})"
            for urn in LABELS
            if not is_meaningful_text(_value(urn))
        ]
        if missing_props:
            reasons.append(
                "DDL: FAIL [DDL-STRUCTURED-PROPERTIES] "
                + "、".join(f"{label} 缺失或为空" for label in missing_props)
            )
        elif schema_field_count <= 0:
            reasons.append("DDL: FAIL [DDL-SCHEMA-FIELDS] schemaMetadata.fields 为空")
        else:
            result.passed_flags.add(FLAG_DDL)
            reasons.append("DDL: PASS")

        ok, documented, missing, extra, errors, lineage_reason = evaluate_table_lineage(
            documentation,
            upstreams,
            target_table_name=table_name,
        )
        result.lineage_documented_upstreams = sorted(documented)
        result.lineage_missing_upstreams = missing
        result.lineage_extra_upstreams = extra
        if ok:
            result.passed_flags.add(FLAG_TABLE_LINEAGE)
        reasons.append(lineage_reason)
        reasons.extend(errors)

    result.final_flags = merge_availability_flags(existing_flags, result.passed_flags)
    result.reason = "；".join(reasons)
    return result


def patch_data_availability_flags(
    gms_url: str,
    dataset_urn: str,
    flags: list[str],
    token: Optional[str] = None,
) -> None:
    patch_ops = [
        {
            "op": "add",
            "path": f"/properties/{URN_DATA_AVAILABILITY_FLAG}",
            "value": {
                "propertyUrn": URN_DATA_AVAILABILITY_FLAG,
                "values": [{"string": flag} for flag in flags],
            },
        }
    ]
    body = {
        "patch": patch_ops,
        "arrayPrimaryKeys": {"properties": ["propertyUrn"]},
    }
    data = json.dumps(body, ensure_ascii=False).encode("utf-8")
    req = urllib.request.Request(
        f"{gms_url.rstrip('/')}/openapi/v3/entity/dataset/{urllib.parse.quote(dataset_urn, safe='')}/structuredProperties",
        data=data,
        method="PATCH",
        headers={
            "Content-Type": "application/json-patch+json",
            "Accept": "application/json",
        },
    )
    if token:
        req.add_header("Authorization", f"Bearer {token}")
    try:
        with urllib.request.urlopen(req, timeout=180) as resp:
            resp.read()
    except urllib.error.HTTPError as exc:
        detail = exc.read().decode("utf-8", errors="replace")
        raise RuntimeError(f"PATCH data_availability_flag HTTP {exc.code}: {detail}") from exc


def check_one_dataset(
    table_name: str,
    *,
    gms_url: str,
    token: Optional[str],
    platform_instance: str,
    env: str,
    dry_run: bool,
) -> AvailabilityResult:
    normalized = table_name.strip().lower()
    dataset_urn = make_hive_dataset_urn(normalized, platform_instance, env)
    structured_payload = fetch_structured_properties_or_empty(gms_url, dataset_urn, token=token)
    structured_values = extract_structured_values(structured_payload)
    existing_flags = extract_existing_flags(structured_values)

    deprecation_payload = fetch_aspect_payload(gms_url, dataset_urn, "deprecation", token=token)
    is_deprecated = is_deprecated_from_payload(deprecation_payload)
    if is_deprecated:
        result = evaluate_dataset_availability(
            table_name=normalized,
            dataset_urn=dataset_urn,
            is_view=False,
            is_deprecated=True,
            structured_values=structured_values,
            schema_field_count=0,
            view_logic="",
            documentation="",
            upstreams=set(),
            existing_flags=existing_flags,
        )
        if set(result.final_flags) == existing_flags:
            result.write_status = "NO_CHANGE"
        elif dry_run:
            result.write_status = "DRY_RUN"
        else:
            patch_data_availability_flags(gms_url, dataset_urn, result.final_flags, token=token)
            result.write_status = "UPDATED"
        return result

    view_payload = fetch_aspect_payload(gms_url, dataset_urn, "viewProperties", token=token)
    view_logic = view_logic_from_payload(view_payload)
    is_view = is_meaningful_text(view_logic)
    schema_fields = schema_field_count(fetch_aspect_payload(gms_url, dataset_urn, "schemaMetadata", token=token))
    documentation = editable_description_from_payload(
        fetch_aspect_payload(gms_url, dataset_urn, "editableDatasetProperties", token=token)
    )
    upstreams = upstream_tables_from_payload(
        fetch_aspect_payload(gms_url, dataset_urn, "upstreamLineage", token=token),
        platform_instance=platform_instance,
    )

    result = evaluate_dataset_availability(
        table_name=normalized,
        dataset_urn=dataset_urn,
        is_view=is_view,
        structured_values=structured_values,
        schema_field_count=schema_fields,
        view_logic=view_logic,
        documentation=documentation,
        upstreams=upstreams,
        existing_flags=existing_flags,
    )
    if set(result.final_flags) == existing_flags:
        result.write_status = "NO_CHANGE"
    elif dry_run:
        result.write_status = "DRY_RUN"
    else:
        patch_data_availability_flags(gms_url, dataset_urn, result.final_flags, token=token)
        result.write_status = "UPDATED"
    return result


def result_to_row(result: AvailabilityResult) -> dict[str, Any]:
    row = asdict(result)
    row["passed_flags"] = sorted(result.passed_flags)
    row["existing_flags"] = sorted(result.existing_flags)
    return row


def write_jsonl_report(path: str, rows: list[dict[str, Any]]) -> None:
    out = Path(path)
    out.parent.mkdir(parents=True, exist_ok=True)
    with out.open("w", encoding="utf-8") as fh:
        for row in rows:
            fh.write(json.dumps(row, ensure_ascii=False) + "\n")


def write_xlsx_report(path: str, rows: list[dict[str, Any]]) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise RuntimeError("请先安装 openpyxl") from exc

    wb = Workbook()
    ws = wb.active
    ws.title = "availability"
    headers = [
        "table_name",
        "dataset_type",
        "status",
        "passed_flags",
        "existing_flags",
        "final_flags",
        "write_status",
        "reason",
        "lineage_documented_upstreams",
        "lineage_existing_upstreams",
        "lineage_missing_upstreams",
        "lineage_extra_upstreams",
        "error",
        "dataset_urn",
    ]
    ws.append(headers)
    for row in rows:
        ws.append([
            json.dumps(row.get(key), ensure_ascii=False) if isinstance(row.get(key), (list, dict)) else row.get(key, "")
            for key in headers
        ])
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9E1F2")
    for col_idx, _ in enumerate(headers, start=1):
        col_letter = get_column_letter(col_idx)
        max_len = 10
        for cell in ws[col_letter]:
            if cell.value:
                max_len = min(max(max_len, len(str(cell.value))), 80)
        ws.column_dimensions[col_letter].width = max_len + 2
    out = Path(path)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)


def run(
    table_names: list[str],
    *,
    gms_url: str,
    token: Optional[str],
    platform_instance: str,
    env: str,
    dry_run: bool,
    jsonl_path: str,
    xlsx_path: str,
) -> int:
    rows: list[dict[str, Any]] = []
    failures = 0
    for idx, table_name in enumerate(table_names, start=1):
        started = time.time()
        print(f"[INFO] [{idx}/{len(table_names)}] 检查 {table_name}", flush=True)
        try:
            result = check_one_dataset(
                table_name,
                gms_url=gms_url,
                token=token,
                platform_instance=platform_instance,
                env=env,
                dry_run=dry_run,
            )
        except Exception as exc:
            failures += 1
            normalized = table_name.strip().lower()
            result = AvailabilityResult(
                table_name=normalized,
                dataset_urn=make_hive_dataset_urn(normalized, platform_instance, env),
                dataset_type="unknown",
                status="FAIL",
                reason="DataHub 读取或写入失败",
                error=str(exc),
            )
        row = result_to_row(result)
        row["elapsed"] = round(time.time() - started, 2)
        rows.append(row)
        print(
            f"[{result.status}] table={result.table_name} type={result.dataset_type} "
            f"passed={sorted(result.passed_flags)} final={result.final_flags} write={result.write_status}",
            flush=True,
        )
        if result.reason:
            print(f"    reason={result.reason}", flush=True)
        if result.error:
            print(f"    error={result.error}", flush=True)

    write_jsonl_report(jsonl_path, rows)
    write_xlsx_report(xlsx_path, rows)
    print(f"[DONE] jsonl: {jsonl_path}", flush=True)
    print(f"[DONE] xlsx: {xlsx_path}", flush=True)
    return 1 if failures else 0


def parse_args(argv: Optional[list[str]] = None) -> argparse.Namespace:
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--table-name", action="append", default=[])
    p.add_argument("--table-list-file")
    p.add_argument("--table-prefix", default=os.getenv("TABLE_PRE", "").strip())
    p.add_argument("--datahub-gms", default=os.getenv("DATAHUB_GMS_URL", "http://localhost:8080"))
    p.add_argument("--token", default=os.getenv("DATAHUB_GMS_TOKEN") or None)
    p.add_argument("--platform-instance", default=os.getenv("BLF_DATAHUB_PLATFORM_INSTANCE", "blf-prod-hive"))
    p.add_argument("--env", default=os.getenv("DATAHUB_ENV", "PROD"))
    p.add_argument("--dry-run", action="store_true", default=os.getenv("DRY_RUN", "1") == "1")
    p.add_argument("--jsonl", required=True)
    p.add_argument("--xlsx", required=True)
    return p.parse_args(argv)


def main(argv: Optional[list[str]] = None) -> int:
    args = parse_args(argv)
    table_names: list[str] = []
    if args.table_list_file:
        table_names.extend(parse_table_names(Path(args.table_list_file).read_text(encoding="utf-8")))
    table_names.extend(parse_table_names("\n".join(args.table_name)))
    table_names.extend(parse_table_names(os.getenv("TABLE_NAMES", "")))
    table_names = list(dict.fromkeys(table_names))
    if not table_names and args.table_prefix:
        table_names = discover_table_names_by_prefix_from_mysql(
            args.table_prefix,
            platform_instance=args.platform_instance,
            env=args.env,
        )
        print(f"[INFO] TABLE_PRE={args.table_prefix} discovered {len(table_names)} datasets", flush=True)
    if not table_names:
        print("ERROR: 请通过 TABLE_NAMES、--table-name、--table-list-file 或 TABLE_PRE/--table-prefix 提供表名", file=sys.stderr)
        return 2
    return run(
        table_names,
        gms_url=args.datahub_gms,
        token=args.token,
        platform_instance=args.platform_instance,
        env=args.env,
        dry_run=args.dry_run,
        jsonl_path=args.jsonl,
        xlsx_path=args.xlsx,
    )


if __name__ == "__main__":
    raise SystemExit(main())
