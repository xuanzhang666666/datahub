#!/usr/bin/env python3
"""Audit current DataHub Hive table-level lineage quality.

The audit is read-only. It scans DataHub dataset upstreamLineage aspects and
checks whether lineage endpoints still exist in DataHub and Hive.
"""

from __future__ import annotations

import argparse
import json
import os
import re
import shutil
import subprocess
import sys
from collections import Counter
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any, Dict, Iterable, List, Mapping, Optional, Sequence, Set

from .field_lineage_datahub_reader import make_hive_dataset_urn
from .hive_table_existence import query_hive_existing_fqtns
from .query_upstream_lineage import urn_to_table_name
from .check_dataset_availability import (
    is_meaningful_text,
    parse_documented_upstreams,
)
from .field_lineage_datahub_reader import strip_markdown_code_fence
from .structured_properties import (
    URN_DATA_AVAILABILITY_FLAG,
    URN_ETL_SCRIPT,
    sort_data_availability_flags,
)

DEFAULT_PLATFORM_INSTANCE = "blf-prod-hive"
DEFAULT_ENV = "PROD"
FLAG_DDL = "DDL"
FLAG_TABLE_LINEAGE = "表血缘"
FLAG_FIELD_LINEAGE = "字段血缘"

ISSUE_TARGET_NOT_IN_HIVE = "target_not_in_hive"
ISSUE_UPSTREAM_DATAHUB_ENTITY_NOT_FOUND = "upstream_datahub_entity_not_found"
ISSUE_UPSTREAM_NOT_IN_HIVE = "upstream_not_in_hive"
ISSUE_SELF_DEPENDENCY = "self_dependency"
ISSUE_UPSTREAM_NOT_HIVE_PLATFORM = "upstream_not_hive_platform"
ISSUE_NON_ODS_NO_UPSTREAM = "non_ods_no_upstream"
ISSUE_MISSING_ETL_SCRIPT = "missing_etl_script"
ISSUE_VIEW_NO_UPSTREAM = "view_no_upstream"
ISSUE_ETL_SCRIPT_NO_UPSTREAM = "etl_script_no_upstream"
ISSUE_DOCUMENTATION_MISSING_DATA_SOURCE_SECTION = "documentation_missing_data_source_section"
ISSUE_DOCUMENTATION_SOURCE_PARSE_EMPTY = "documentation_source_parse_empty"
ISSUE_DOCUMENTATION_LINEAGE_MISMATCH = "documentation_lineage_mismatch"
ISSUE_DOCUMENTATION_UPSTREAM_NOT_IN_HIVE = "documentation_upstream_not_in_hive"
ISSUE_ETL_SCRIPT_PLACEHOLDER = "etl_script_placeholder"
ISSUE_ETL_SCRIPT_HAS_UNRESOLVED_VARS = "etl_script_has_unresolved_vars"
ISSUE_ETL_SCRIPT_HAS_NO_DOCUMENTATION = "etl_script_has_no_documentation"
ISSUE_FLAG_CLAIMS_LINEAGE_BUT_NO_UPSTREAM = "flag_claims_lineage_but_no_upstream"
ISSUE_FLAG_MISSING_LINEAGE_BUT_HAS_UPSTREAM = "flag_missing_lineage_but_has_upstream"
ISSUE_FLAG_CLAIMS_DDL_BUT_NO_SCHEMA_OR_VIEW_LOGIC = "flag_claims_ddl_but_no_schema_or_view_logic"
ISSUE_VIEW_FLAG_INCOMPLETE = "view_flag_incomplete"
ISSUE_VIEW_MISSING_DEFINITION = "view_missing_definition"
ISSUE_VIEW_DEFINITION_NO_UPSTREAM = "view_definition_no_upstream"

_LINEAGE_TABLE_PREFIXES_WITH_UPSTREAM = (
    "app_",
    "dm_",
    "dim_",
    "dw_",
    "dwa_",
    "dwd_",
    "mid_",
    "pdw_",
)
_TABLE_PREFIXES_ALLOW_MISSING_ETL = ("ods", "ai", "app")
_TABLE_PREFIXES_REQUIRE_ETL = ("dwa", "dwd", "pdim", "dim", "pdw", "mid", "dm", "dw")

_ISSUE_DEFAULTS: Dict[str, tuple[str, str]] = {
    ISSUE_TARGET_NOT_IN_HIVE: ("P0", "先确认该 DataHub dataset 是否已废弃；若仍在线，重跑 Hive ingest。"),
    ISSUE_SELF_DEPENDENCY: ("P0", "检查 LLM/手工血缘输出，移除目标表指向自身的 upstreamLineage。"),
    ISSUE_VIEW_NO_UPSTREAM: ("P0", "对该 view 重跑 Hive ingest/full lineage，或手工补充 upstreamLineage。"),
    ISSUE_VIEW_DEFINITION_NO_UPSTREAM: ("P0", "view 有定义但无上游，优先用 View Definition 重建 upstreamLineage。"),
    ISSUE_NON_ODS_NO_UPSTREAM: ("P0", "非源层表应有上游；重跑作业血缘同步或手工补血缘。"),
    ISSUE_ETL_SCRIPT_NO_UPSTREAM: ("P1", "Etl Script 已存在但无血缘，重跑 job 血缘同步并检查 LLM/raw 报告。"),
    ISSUE_UPSTREAM_NOT_IN_HIVE: ("P1", "确认上游表是否已下线或表名解析错误；必要时从 DataHub 血缘中移除。"),
    ISSUE_UPSTREAM_DATAHUB_ENTITY_NOT_FOUND: ("P1", "对上游表执行 Hive ingest，或修正上游 URN。"),
    ISSUE_UPSTREAM_NOT_HIVE_PLATFORM: ("P1", "检查上游平台/env 是否写错，保持 Hive platform instance 一致。"),
    ISSUE_DOCUMENTATION_LINEAGE_MISMATCH: ("P1", "对齐 Documentation「4. 数据来源」与 DataHub upstreamLineage。"),
    ISSUE_DOCUMENTATION_UPSTREAM_NOT_IN_HIVE: ("P1", "修正文档中的上游表，或先同步该 Hive 表。"),
    ISSUE_FLAG_CLAIMS_LINEAGE_BUT_NO_UPSTREAM: ("P1", "data_availability_flag 声称有表血缘但 upstreamLineage 为空，重建血缘或修正 flag。"),
    ISSUE_FLAG_CLAIMS_DDL_BUT_NO_SCHEMA_OR_VIEW_LOGIC: ("P1", "flag 声称有 DDL，但 schema/viewLogic 缺失，重跑 Hive ingest 或修正 flag。"),
    ISSUE_MISSING_ETL_SCRIPT: ("P2", "按作业重跑血缘同步，写入 Etl Script structured property。"),
    ISSUE_DOCUMENTATION_MISSING_DATA_SOURCE_SECTION: ("P2", "重跑 Documentation 生成，确保包含「4. 数据来源」小节。"),
    ISSUE_DOCUMENTATION_SOURCE_PARSE_EMPTY: ("P2", "检查 Documentation 第 4 节表格格式，第一列应为 db.table。"),
    ISSUE_ETL_SCRIPT_PLACEHOLDER: ("P2", "Etl Script 是占位符，重新解析调度作业并写入真实脚本。"),
    ISSUE_ETL_SCRIPT_HAS_UNRESOLVED_VARS: ("P2", "Etl Script 仍含 ${VAR}，需要从 shell_command 参数渲染后重跑血缘。"),
    ISSUE_ETL_SCRIPT_HAS_NO_DOCUMENTATION: ("P2", "Etl Script 已有但 Documentation 缺失，重跑文档生成任务。"),
    ISSUE_FLAG_MISSING_LINEAGE_BUT_HAS_UPSTREAM: ("P2", "已有 upstreamLineage 但 flag 缺表血缘，重跑 availability flag 更新。"),
    ISSUE_VIEW_FLAG_INCOMPLETE: ("P2", "view 已有 View Definition 与上游，更新 flag 为 DDL/表血缘/字段血缘。"),
    ISSUE_VIEW_MISSING_DEFINITION: ("P2", "对该 view 执行 Hive ingest full 模式，补齐 viewProperties.viewLogic。"),
}


@dataclass(frozen=True)
class QualityIssue:
    issue_type: str
    target_table: str
    upstream_table: str = ""
    target_urn: str = ""
    upstream_urn: str = ""
    reason: str = ""
    detail: str = ""
    severity: str = ""
    fix_hint: str = ""
    evidence: str = ""

    def __post_init__(self) -> None:
        severity, fix_hint = _ISSUE_DEFAULTS.get(
            self.issue_type,
            ("P3", "根据 issue_type 检查对应 DataHub aspect 并修复。"),
        )
        if not self.severity:
            object.__setattr__(self, "severity", severity)
        if not self.fix_hint:
            object.__setattr__(self, "fix_hint", fix_hint)
        if not self.evidence:
            evidence = self.detail or self.upstream_urn or self.target_urn
            object.__setattr__(self, "evidence", evidence)


@dataclass(frozen=True)
class AuditResult:
    summary: Dict[str, Any]
    issues: List[QualityIssue]


def _is_hive_dataset_urn(urn: str, platform_instance: str, env: str) -> bool:
    return (
        urn.startswith("urn:li:dataset:(urn:li:dataPlatform:hive,")
        and f",{env})" in urn
        and f",{platform_instance}." in urn
    )


def _should_expect_upstream(table_name: str) -> bool:
    table = table_name.rsplit(".", 1)[-1].lower()
    return table.startswith(_LINEAGE_TABLE_PREFIXES_WITH_UPSTREAM)


def _should_expect_etl_script(table_name: str) -> bool:
    table = table_name.rsplit(".", 1)[-1].lower()
    if table.startswith(_TABLE_PREFIXES_ALLOW_MISSING_ETL):
        return False
    return table.startswith(_TABLE_PREFIXES_REQUIRE_ETL)


def _strip_markdown_code_fence(value: str) -> str:
    text = (value or "").strip()
    if text.startswith("```") and text.endswith("```"):
        lines = text.splitlines()
        if len(lines) >= 2:
            return "\n".join(lines[1:-1]).strip()
    return text


def _has_meaningful_etl_script(metadata: str) -> bool:
    try:
        payload = json.loads(metadata)
    except json.JSONDecodeError:
        return False
    properties = payload.get("properties")
    if not isinstance(properties, list):
        return False
    for prop in properties:
        if not isinstance(prop, dict) or prop.get("propertyUrn") != URN_ETL_SCRIPT:
            continue
        values = prop.get("values")
        if not isinstance(values, list):
            continue
        for value in values:
            if isinstance(value, dict):
                text = value.get("string", "")
            elif isinstance(value, str):
                text = value
            else:
                text = ""
            stripped = _strip_markdown_code_fence(text)
            if stripped and stripped != "无":
                return True
    return False


def _has_meaningful_etl_script_value(etl_script: str) -> bool:
    stripped = strip_markdown_code_fence(etl_script or "").strip()
    return is_meaningful_text(etl_script) and stripped != "无"


def _metadata_aspect(metadata: str) -> Dict[str, Any]:
    try:
        payload = json.loads(metadata)
    except json.JSONDecodeError:
        return {}
    if not isinstance(payload, dict):
        return {}
    value = payload.get("value")
    return value if isinstance(value, dict) else payload


def _first_string_value(prop: Mapping[str, Any]) -> str:
    values = prop.get("values")
    if not isinstance(values, list):
        return ""
    for item in values:
        if isinstance(item, dict):
            text = item.get("string")
        else:
            text = item
        if isinstance(text, str) and text.strip():
            return text.strip()
    return ""


def _all_string_values(prop: Mapping[str, Any]) -> List[str]:
    values = prop.get("values")
    if not isinstance(values, list):
        return []
    out: List[str] = []
    for item in values:
        if isinstance(item, dict):
            text = item.get("string")
        else:
            text = item
        if isinstance(text, str) and text.strip():
            out.append(text.strip())
    return out


def _structured_property_values(metadata: str) -> Dict[str, List[str]]:
    aspect = _metadata_aspect(metadata)
    props = aspect.get("properties")
    if not isinstance(props, list):
        return {}
    out: Dict[str, List[str]] = {}
    for prop in props:
        if not isinstance(prop, dict):
            continue
        urn = prop.get("propertyUrn")
        if isinstance(urn, str) and urn:
            out[urn] = _all_string_values(prop)
    return out


def _extract_etl_script(metadata: str) -> str:
    values = _structured_property_values(metadata).get(URN_ETL_SCRIPT, [])
    return values[0] if values else ""


def _extract_data_availability_flags(metadata: str) -> Set[str]:
    values = _structured_property_values(metadata).get(URN_DATA_AVAILABILITY_FLAG, [])
    return set(sort_data_availability_flags(values))


def _extract_editable_description(metadata: str) -> str:
    aspect = _metadata_aspect(metadata)
    description = aspect.get("description")
    return description if isinstance(description, str) else ""


def _extract_view_logic(metadata: str) -> str:
    aspect = _metadata_aspect(metadata)
    view_logic = aspect.get("viewLogic")
    return view_logic if isinstance(view_logic, str) else ""


def _extract_schema_field_count(metadata: str) -> int:
    aspect = _metadata_aspect(metadata)
    fields = aspect.get("fields")
    return len(fields) if isinstance(fields, list) else 0


def _contains_unresolved_shell_vars(text: str) -> bool:
    return bool(re.search(r"\$\{[A-Za-z_][A-Za-z0-9_]*\}|\$[A-Za-z_][A-Za-z0-9_]*\b", text or ""))


def _collect_needed_hive_tables(
    dataset_urns: Set[str],
    upstreams_by_target: Mapping[str, Sequence[str]],
    platform_instance: str,
    env: str,
) -> Set[str]:
    needed = {
        urn_to_table_name(urn, platform_instance).lower()
        for urn in dataset_urns
        if _is_hive_dataset_urn(urn, platform_instance, env)
    }
    for upstreams in upstreams_by_target.values():
        for upstream_urn in upstreams:
            if _is_hive_dataset_urn(upstream_urn, platform_instance, env):
                needed.add(urn_to_table_name(upstream_urn, platform_instance).lower())
    return {x for x in needed if "." in x}


def _collect_documented_hive_tables(documentation_by_urn: Mapping[str, str]) -> Set[str]:
    needed: Set[str] = set()
    for description in documentation_by_urn.values():
        section_found, upstreams = parse_documented_upstreams(description)
        if not section_found:
            continue
        for upstream in upstreams:
            if "." in upstream:
                needed.add(upstream.lower())
    return needed


def evaluate_lineage_quality(
    dataset_urns: Set[str],
    upstreams_by_target: Mapping[str, Sequence[str]],
    hive_existing_fqtns: Set[str],
    *,
    platform_instance: str = DEFAULT_PLATFORM_INSTANCE,
    env: str = DEFAULT_ENV,
    include_no_upstream: bool = False,
    check_datahub_entity_existence: bool = True,
    view_dataset_urns: Optional[Set[str]] = None,
    etl_script_dataset_urns: Optional[Set[str]] = None,
    documentation_by_urn: Optional[Mapping[str, str]] = None,
    data_availability_flags_by_urn: Optional[Mapping[str, Set[str]]] = None,
    view_logic_by_urn: Optional[Mapping[str, str]] = None,
    schema_field_count_by_urn: Optional[Mapping[str, int]] = None,
    etl_script_by_urn: Optional[Mapping[str, str]] = None,
) -> AuditResult:
    """Classify lineage quality issues from already fetched DataHub/Hive facts."""
    normalized_dataset_urns = {urn.strip() for urn in dataset_urns if urn and urn.strip()}
    normalized_hive = {name.strip().lower() for name in hive_existing_fqtns if name.strip()}
    normalized_views = {urn.strip() for urn in (view_dataset_urns or set()) if urn and urn.strip()}
    normalized_etl = {urn.strip() for urn in (etl_script_dataset_urns or set()) if urn and urn.strip()}
    has_documentation_facts = documentation_by_urn is not None
    has_flag_facts = data_availability_flags_by_urn is not None
    has_view_logic_facts = view_logic_by_urn is not None
    has_schema_facts = schema_field_count_by_urn is not None
    has_etl_script_values = etl_script_by_urn is not None
    documentation = documentation_by_urn or {}
    data_flags = data_availability_flags_by_urn or {}
    view_logic = view_logic_by_urn or {}
    schema_field_counts = schema_field_count_by_urn or {}
    etl_scripts = etl_script_by_urn or {}
    issues: List[QualityIssue] = []
    upstream_edge_count = 0

    for target_urn in sorted(normalized_dataset_urns):
        target_table = urn_to_table_name(target_urn, platform_instance).lower()
        upstreams = list(upstreams_by_target.get(target_urn, []))
        upstream_tables = {
            urn_to_table_name(upstream_urn, platform_instance).lower()
            for upstream_urn in upstreams
            if _is_hive_dataset_urn(upstream_urn, platform_instance, env)
        }
        target_is_view = target_urn in normalized_views
        target_has_upstream = bool(upstreams)
        target_flags = data_flags.get(target_urn, set())
        target_view_logic = view_logic.get(target_urn, "")
        target_schema_field_count = schema_field_counts.get(target_urn, 0)
        target_etl_script = etl_scripts.get(target_urn, "")

        if target_table not in normalized_hive:
            issues.append(
                QualityIssue(
                    issue_type=ISSUE_TARGET_NOT_IN_HIVE,
                    target_table=target_table,
                    target_urn=target_urn,
                    reason="目标表在 Hive information_schema 中不存在",
                )
            )

        if (
            etl_script_dataset_urns is not None
            and not target_is_view
            and _should_expect_etl_script(target_table)
            and target_urn not in normalized_etl
        ):
            issues.append(
                QualityIssue(
                    issue_type=ISSUE_MISSING_ETL_SCRIPT,
                    target_table=target_table,
                    target_urn=target_urn,
                    reason="非 view 且按表名前缀应维护 Etl Script，但结构化属性 blf.data.warehouse.etl_script 缺失或无有效内容",
                    detail="允许缺失前缀: ods, ai, app；要求存在前缀: dwa, dwd, pdim, dim, pdw, mid, dm, dw",
                )
            )

        if target_is_view and not upstreams:
            issues.append(
                QualityIssue(
                    issue_type=ISSUE_VIEW_NO_UPSTREAM,
                    target_table=target_table,
                    target_urn=target_urn,
                    reason="view dataset 没有 DataHub 表级上游血缘",
                )
            )

        if (
            etl_script_dataset_urns is not None
            and not target_is_view
            and target_urn in normalized_etl
            and not upstreams
            and _should_expect_etl_script(target_table)
        ):
            issues.append(
                QualityIssue(
                    issue_type=ISSUE_ETL_SCRIPT_NO_UPSTREAM,
                    target_table=target_table,
                    target_urn=target_urn,
                    reason="Etl Script 有有效内容且按表名前缀应建立上游血缘，但没有 DataHub 表级上游表",
                    detail="允许无上游前缀: ods, ai, app；要求有上游前缀: dwa, dwd, pdim, dim, pdw, mid, dm, dw",
                )
            )

        if include_no_upstream and not upstreams and _should_expect_upstream(target_table):
            issues.append(
                QualityIssue(
                    issue_type=ISSUE_NON_ODS_NO_UPSTREAM,
                    target_table=target_table,
                    target_urn=target_urn,
                    reason="非 ODS/源层表没有 DataHub 表级上游血缘",
                )
            )

        if target_is_view and has_view_logic_facts:
            if not is_meaningful_text(target_view_logic):
                issues.append(
                    QualityIssue(
                        issue_type=ISSUE_VIEW_MISSING_DEFINITION,
                        target_table=target_table,
                        target_urn=target_urn,
                        reason="view dataset 缺少 View Definition（viewProperties.viewLogic 为空）",
                        detail="viewProperties.viewLogic",
                    )
                )
            elif not target_has_upstream:
                issues.append(
                    QualityIssue(
                        issue_type=ISSUE_VIEW_DEFINITION_NO_UPSTREAM,
                        target_table=target_table,
                        target_urn=target_urn,
                        reason="view 有 View Definition，但没有 DataHub 表级上游血缘",
                        detail="viewProperties.viewLogic 已填写；upstreamLineage.upstreams 为空",
                    )
                )

        if has_etl_script_values and target_etl_script:
            stripped_etl = strip_markdown_code_fence(target_etl_script).strip()
            if not is_meaningful_text(target_etl_script) or stripped_etl == "无":
                issues.append(
                    QualityIssue(
                        issue_type=ISSUE_ETL_SCRIPT_PLACEHOLDER,
                        target_table=target_table,
                        target_urn=target_urn,
                        reason="Etl Script 内容为空或为占位符",
                        detail=f"Etl Script={stripped_etl[:100]!r}",
                    )
                )
            elif _contains_unresolved_shell_vars(stripped_etl):
                issues.append(
                    QualityIssue(
                        issue_type=ISSUE_ETL_SCRIPT_HAS_UNRESOLVED_VARS,
                        target_table=target_table,
                        target_urn=target_urn,
                        reason="Etl Script 中仍存在未替换 shell 变量，可能导致 LLM 解析出无效表名",
                        detail="检测到 ${VAR} 或 $VAR",
                    )
                )
            if has_documentation_facts and not is_meaningful_text(documentation.get(target_urn, "")):
                issues.append(
                    QualityIssue(
                        issue_type=ISSUE_ETL_SCRIPT_HAS_NO_DOCUMENTATION,
                        target_table=target_table,
                        target_urn=target_urn,
                        reason="Etl Script 有有效内容，但 Documentation 为空",
                        detail="editableDatasetProperties.description",
                    )
                )

        description = documentation.get(target_urn, "")
        if (
            has_documentation_facts
            and description
            and not target_is_view
            and (target_urn in normalized_etl or _should_expect_etl_script(target_table))
        ):
            section_found, documented_upstreams = parse_documented_upstreams(description)
            if not section_found:
                issues.append(
                    QualityIssue(
                        issue_type=ISSUE_DOCUMENTATION_MISSING_DATA_SOURCE_SECTION,
                        target_table=target_table,
                        target_urn=target_urn,
                        reason="Documentation 缺少「4. 数据来源」小节",
                        detail='editableDatasetProperties.description / "4. 数据来源"',
                    )
                )
            elif not documented_upstreams:
                issues.append(
                    QualityIssue(
                        issue_type=ISSUE_DOCUMENTATION_SOURCE_PARSE_EMPTY,
                        target_table=target_table,
                        target_urn=target_urn,
                        reason="Documentation「4. 数据来源」存在但未解析到上游表",
                        detail='第 4 节第一列应填写 `db.table`',
                    )
                )
            else:
                documented_norm = {item.lower() for item in documented_upstreams}
                missing_in_lineage = sorted(documented_norm - upstream_tables)
                extra_in_lineage = sorted(upstream_tables - documented_norm)
                if not target_has_upstream:
                    issues.append(
                        QualityIssue(
                            issue_type=ISSUE_DOCUMENTATION_LINEAGE_MISMATCH,
                            target_table=target_table,
                            target_urn=target_urn,
                            reason="Documentation 写了上游表，但 DataHub upstreamLineage 为空",
                            detail=f"documented={missing_in_lineage[:20]}",
                            evidence=json.dumps(
                                {"documented": missing_in_lineage[:50], "lineage": []},
                                ensure_ascii=False,
                            ),
                        )
                    )
                elif missing_in_lineage or extra_in_lineage:
                    issues.append(
                        QualityIssue(
                            issue_type=ISSUE_DOCUMENTATION_LINEAGE_MISMATCH,
                            target_table=target_table,
                            target_urn=target_urn,
                            reason="Documentation「4. 数据来源」与 DataHub upstreamLineage 不一致",
                            detail=f"missing={missing_in_lineage[:10]} extra={extra_in_lineage[:10]}",
                            evidence=json.dumps(
                                {
                                    "missing_in_lineage": missing_in_lineage[:50],
                                    "extra_in_lineage": extra_in_lineage[:50],
                                },
                                ensure_ascii=False,
                            ),
                        )
                    )
                for documented_table in sorted(documented_norm):
                    if "." in documented_table and documented_table not in normalized_hive:
                        issues.append(
                            QualityIssue(
                                issue_type=ISSUE_DOCUMENTATION_UPSTREAM_NOT_IN_HIVE,
                                target_table=target_table,
                                upstream_table=documented_table,
                                target_urn=target_urn,
                                reason="Documentation「4. 数据来源」中的上游表在 Hive information_schema 中不存在",
                            )
                        )

        if has_flag_facts and FLAG_TABLE_LINEAGE in target_flags and not target_has_upstream:
            issues.append(
                QualityIssue(
                    issue_type=ISSUE_FLAG_CLAIMS_LINEAGE_BUT_NO_UPSTREAM,
                    target_table=target_table,
                    target_urn=target_urn,
                    reason="data_availability_flag 包含「表血缘」，但 upstreamLineage 为空",
                    detail=f"flags={sorted(target_flags)}",
                )
            )
        if has_flag_facts and FLAG_TABLE_LINEAGE not in target_flags and target_has_upstream:
            issues.append(
                QualityIssue(
                    issue_type=ISSUE_FLAG_MISSING_LINEAGE_BUT_HAS_UPSTREAM,
                    target_table=target_table,
                    target_urn=target_urn,
                    reason="DataHub 已有上游血缘，但 data_availability_flag 缺少「表血缘」",
                    detail=f"flags={sorted(target_flags)} upstream_count={len(upstreams)}",
                )
            )
        if has_flag_facts and FLAG_DDL in target_flags:
            has_ddl = (
                is_meaningful_text(target_view_logic)
                if target_is_view
                else target_schema_field_count > 0
            )
            if (
                (target_is_view and has_view_logic_facts)
                or (not target_is_view and has_schema_facts)
            ) and not has_ddl:
                issues.append(
                    QualityIssue(
                        issue_type=ISSUE_FLAG_CLAIMS_DDL_BUT_NO_SCHEMA_OR_VIEW_LOGIC,
                        target_table=target_table,
                        target_urn=target_urn,
                        reason="data_availability_flag 包含 DDL，但 schemaMetadata/viewLogic 不完整",
                        detail=f"schema_fields={target_schema_field_count} viewLogic_len={len(target_view_logic)}",
                    )
                )
        if (
            has_flag_facts
            and has_view_logic_facts
            and target_is_view
            and is_meaningful_text(target_view_logic)
            and target_has_upstream
            and not {FLAG_DDL, FLAG_TABLE_LINEAGE, FLAG_FIELD_LINEAGE}.issubset(target_flags)
        ):
            issues.append(
                QualityIssue(
                    issue_type=ISSUE_VIEW_FLAG_INCOMPLETE,
                    target_table=target_table,
                    target_urn=target_urn,
                    reason="view 已有 View Definition 和上游血缘，但 data_availability_flag 未标齐三项",
                    detail=f"flags={sorted(target_flags)}",
                )
            )

        for upstream_urn in upstreams:
            if not upstream_urn:
                continue
            upstream_edge_count += 1
            upstream_table = urn_to_table_name(upstream_urn, platform_instance).lower()

            if upstream_urn == target_urn or upstream_table == target_table:
                issues.append(
                    QualityIssue(
                        issue_type=ISSUE_SELF_DEPENDENCY,
                        target_table=target_table,
                        upstream_table=upstream_table,
                        target_urn=target_urn,
                        upstream_urn=upstream_urn,
                        reason="目标表依赖自身",
                    )
                )
                continue

            if not _is_hive_dataset_urn(upstream_urn, platform_instance, env):
                issues.append(
                    QualityIssue(
                        issue_type=ISSUE_UPSTREAM_NOT_HIVE_PLATFORM,
                        target_table=target_table,
                        upstream_table=upstream_table,
                        target_urn=target_urn,
                        upstream_urn=upstream_urn,
                        reason="上游不是当前 Hive platform instance/env 的 dataset",
                    )
                )
                continue

            if check_datahub_entity_existence and upstream_urn not in normalized_dataset_urns:
                issues.append(
                    QualityIssue(
                        issue_type=ISSUE_UPSTREAM_DATAHUB_ENTITY_NOT_FOUND,
                        target_table=target_table,
                        upstream_table=upstream_table,
                        target_urn=target_urn,
                        upstream_urn=upstream_urn,
                        reason="上游 DataHub dataset 实体不存在或已软删除",
                    )
                )

            if upstream_table not in normalized_hive:
                issues.append(
                    QualityIssue(
                        issue_type=ISSUE_UPSTREAM_NOT_IN_HIVE,
                        target_table=target_table,
                        upstream_table=upstream_table,
                        target_urn=target_urn,
                        upstream_urn=upstream_urn,
                        reason="上游表在 Hive information_schema 中不存在",
                    )
                )

    issue_counts = Counter(issue.issue_type for issue in issues)
    severity_counts = Counter(issue.severity for issue in issues)
    summary: Dict[str, Any] = {
        "scanned_dataset_count": len(normalized_dataset_urns),
        "lineage_target_count": len(upstreams_by_target),
        "upstream_edge_count": upstream_edge_count,
        "issue_count": len(issues),
        "issue_counts_by_type": dict(sorted(issue_counts.items())),
        "issue_counts_by_severity": dict(sorted(severity_counts.items())),
        "datahub_entity_existence_check_complete": check_datahub_entity_existence,
        "etl_script_presence_check_complete": etl_script_dataset_urns is not None,
        "documentation_check_complete": documentation_by_urn is not None,
        "availability_flag_check_complete": data_availability_flags_by_urn is not None,
        "view_definition_check_complete": view_logic_by_urn is not None,
        "schema_metadata_check_complete": schema_field_count_by_urn is not None,
        "view_dataset_count": len(normalized_views),
        "dataset_with_etl_script_count": len(normalized_etl),
        "dataset_with_documentation_count": len(documentation),
        "dataset_with_availability_flag_count": len(data_flags),
        "dataset_with_view_definition_count": sum(1 for value in view_logic.values() if is_meaningful_text(value)),
    }
    return AuditResult(summary=summary, issues=issues)


def _mysql_host_base_cmd() -> List[str]:
    mysql_bin = os.getenv("DATAHUB_MYSQL_BIN") or shutil.which("mysql") or "/opt/anaconda3/bin/mysql"
    host = os.getenv("DATAHUB_MYSQL_HOST", "127.0.0.1")
    port = os.getenv("DATAHUB_MYSQL_PORT", "3306")
    user = os.getenv("DATAHUB_MYSQL_USER", "root")
    password = os.getenv("DATAHUB_MYSQL_PASSWORD", "datahub")
    database = os.getenv("DATAHUB_MYSQL_DATABASE", "datahub")
    return [
        mysql_bin,
        "-h",
        host,
        "-P",
        port,
        f"-u{user}",
        f"-p{password}",
        "-D",
        database,
        "--batch",
        "--raw",
        "--skip-column-names",
    ]


def _mysql_docker_base_cmd() -> List[str]:
    container = os.getenv("DATAHUB_MYSQL_CONTAINER", "datahub-mysql-1")
    user = os.getenv("DATAHUB_MYSQL_USER", "root")
    password = os.getenv("DATAHUB_MYSQL_PASSWORD", "datahub")
    database = os.getenv("DATAHUB_MYSQL_DATABASE", "datahub")
    return [
        "docker",
        "exec",
        container,
        "mysql",
        f"-u{user}",
        f"-p{password}",
        "-D",
        database,
        "--batch",
        "--raw",
        "--skip-column-names",
    ]


def _mysql_candidate_cmds() -> List[List[str]]:
    mode = os.getenv("DATAHUB_MYSQL_MODE", "host").strip().lower()
    if mode == "docker":
        return [_mysql_docker_base_cmd()]
    if mode == "host":
        return [_mysql_host_base_cmd(), _mysql_docker_base_cmd()]
    return [_mysql_host_base_cmd(), _mysql_docker_base_cmd()]


def _run_mysql_query(sql: str) -> str:
    errors: List[str] = []
    for base_cmd in _mysql_candidate_cmds():
        proc = subprocess.run(base_cmd + ["-e", sql], capture_output=True)
        if proc.returncode == 0:
            return proc.stdout.decode("utf-8", errors="replace")
        stderr = proc.stderr.decode("utf-8", errors="replace").strip()
        stdout = proc.stdout.decode("utf-8", errors="replace").strip()
        errors.append(f"cmd={base_cmd[0]} exit={proc.returncode}: {(stderr or stdout or 'no output')[:1000]}")
    raise RuntimeError("MySQL 查询失败: " + " | ".join(errors)[:2000])


def _parse_tab_rows(output: str, expected_columns: int) -> List[tuple[str, ...]]:
    rows: List[tuple[str, ...]] = []
    for raw in output.splitlines():
        if not raw.strip():
            continue
        parts = raw.split("\t", expected_columns - 1)
        if len(parts) == expected_columns:
            rows.append(tuple(parts))
    return rows


def fetch_view_dataset_urns_from_mysql(platform_instance: str, env: str) -> Set[str]:
    sql = (
        "select urn "
        "from metadata_aspect_v2 "
        "where aspect='viewProperties' "
        "and version=0 "
        "and urn like 'urn:li:dataset:(urn:li:dataPlatform:hive,"
        f"{platform_instance}.%,{env})'"
    )
    return {row[0] for row in _parse_tab_rows(_run_mysql_query(sql), 1)}


def fetch_etl_script_dataset_urns_from_mysql(platform_instance: str, env: str) -> Set[str]:
    sql = (
        "select urn, metadata "
        "from metadata_aspect_v2 "
        "where aspect='structuredProperties' "
        "and version=0 "
        "and urn like 'urn:li:dataset:(urn:li:dataPlatform:hive,"
        f"{platform_instance}.%,{env})'"
    )
    out: Set[str] = set()
    for urn, metadata in _parse_tab_rows(_run_mysql_query(sql), 2):
        if _has_meaningful_etl_script(metadata):
            out.add(urn)
    return out


def fetch_structured_properties_from_mysql(
    platform_instance: str,
    env: str,
) -> tuple[Dict[str, str], Dict[str, Set[str]]]:
    sql = (
        "select urn, metadata "
        "from metadata_aspect_v2 "
        "where aspect='structuredProperties' "
        "and version=0 "
        "and urn like 'urn:li:dataset:(urn:li:dataPlatform:hive,"
        f"{platform_instance}.%,{env})'"
    )
    etl_by_urn: Dict[str, str] = {}
    flags_by_urn: Dict[str, Set[str]] = {}
    for urn, metadata in _parse_tab_rows(_run_mysql_query(sql), 2):
        etl_script = _extract_etl_script(metadata)
        if etl_script:
            etl_by_urn[urn] = etl_script
        flags = _extract_data_availability_flags(metadata)
        if flags:
            flags_by_urn[urn] = flags
    return etl_by_urn, flags_by_urn


def fetch_documentation_from_mysql(platform_instance: str, env: str) -> Dict[str, str]:
    sql = (
        "select urn, metadata "
        "from metadata_aspect_v2 "
        "where aspect='editableDatasetProperties' "
        "and version=0 "
        "and urn like 'urn:li:dataset:(urn:li:dataPlatform:hive,"
        f"{platform_instance}.%,{env})'"
    )
    out: Dict[str, str] = {}
    for urn, metadata in _parse_tab_rows(_run_mysql_query(sql), 2):
        description = _extract_editable_description(metadata)
        if description:
            out[urn] = description
    return out


def fetch_view_logic_from_mysql(platform_instance: str, env: str) -> Dict[str, str]:
    sql = (
        "select urn, metadata "
        "from metadata_aspect_v2 "
        "where aspect='viewProperties' "
        "and version=0 "
        "and urn like 'urn:li:dataset:(urn:li:dataPlatform:hive,"
        f"{platform_instance}.%,{env})'"
    )
    out: Dict[str, str] = {}
    for urn, metadata in _parse_tab_rows(_run_mysql_query(sql), 2):
        out[urn] = _extract_view_logic(metadata)
    return out


def fetch_schema_field_counts_from_mysql(platform_instance: str, env: str) -> Dict[str, int]:
    sql = (
        "select urn, metadata "
        "from metadata_aspect_v2 "
        "where aspect='schemaMetadata' "
        "and version=0 "
        "and urn like 'urn:li:dataset:(urn:li:dataPlatform:hive,"
        f"{platform_instance}.%,{env})'"
    )
    out: Dict[str, int] = {}
    for urn, metadata in _parse_tab_rows(_run_mysql_query(sql), 2):
        out[urn] = _extract_schema_field_count(metadata)
    return out


def _make_graph(gms_url: str, token: Optional[str]) -> Any:
    try:
        from datahub.ingestion.graph.client import DataHubGraph, DatahubClientConfig
    except ImportError as exc:
        raise RuntimeError("需要安装 acryl-datahub: pip install acryl-datahub") from exc

    return DataHubGraph(DatahubClientConfig(server=gms_url.rstrip("/"), token=token))


def fetch_dataset_urns(
    graph: Any,
    *,
    platform_instance: str,
    env: str,
    query: str,
    batch_size: int,
    max_datasets: int,
) -> Set[str]:
    urns: Set[str] = set()
    for urn in graph.get_urns_by_filter(
        entity_types=["dataset"],
        platform="hive",
        platform_instance=platform_instance,
        env=env,
        query=query,
        batch_size=batch_size,
    ):
        if isinstance(urn, str) and urn:
            urns.add(urn)
        if max_datasets > 0 and len(urns) >= max_datasets:
            break
    return urns


def fetch_upstreams_by_target(graph: Any, dataset_urns: Iterable[str]) -> Dict[str, List[str]]:
    try:
        from datahub.metadata.schema_classes import UpstreamLineageClass
    except ImportError as exc:
        raise RuntimeError("需要安装 acryl-datahub: pip install acryl-datahub") from exc

    out: Dict[str, List[str]] = {}
    for idx, target_urn in enumerate(sorted(dataset_urns), start=1):
        if idx % 500 == 0:
            print(f"[INFO] 已读取 upstreamLineage: {idx} 个 dataset", file=sys.stderr)
        try:
            lineage = graph.get_aspect(
                entity_urn=target_urn,
                aspect_type=UpstreamLineageClass,
            )
        except Exception as exc:
            print(f"[WARN] 读取 upstreamLineage 失败: {target_urn} err={exc}", file=sys.stderr)
            continue
        upstreams = []
        for upstream in getattr(lineage, "upstreams", []) or []:
            upstream_urn = getattr(upstream, "dataset", None)
            if isinstance(upstream_urn, str) and upstream_urn:
                upstreams.append(upstream_urn)
        if upstreams:
            out[target_urn] = upstreams
    return out


def write_jsonl(path: Path, result: AuditResult) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as fh:
        fh.write(json.dumps({"summary": result.summary}, ensure_ascii=False) + "\n")
        for issue in result.issues:
            fh.write(json.dumps(asdict(issue), ensure_ascii=False) + "\n")


def write_xlsx(path: Path, result: AuditResult) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise SystemExit("请先安装: pip install openpyxl") from exc

    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws_summary = wb.active
    ws_summary.title = "summary"
    ws_summary.append(["指标", "值"])
    for key, value in result.summary.items():
        if isinstance(value, dict):
            value = json.dumps(value, ensure_ascii=False, sort_keys=True)
        ws_summary.append([key, value])

    ws_issues = wb.create_sheet("issues")
    headers = [
        "优先级",
        "问题类型",
        "目标表",
        "上游表",
        "原因",
        "明细",
        "修复建议",
        "证据",
        "目标URN",
        "上游URN",
    ]
    ws_issues.append(headers)
    for issue in result.issues:
        ws_issues.append(
            [
                issue.severity,
                issue.issue_type,
                issue.target_table,
                issue.upstream_table,
                issue.reason,
                issue.detail,
                issue.fix_hint,
                issue.evidence,
                issue.target_urn,
                issue.upstream_urn,
            ]
        )

    for ws in (ws_summary, ws_issues):
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D9E1F2")
        for col_idx, _ in enumerate(ws[1], start=1):
            col_letter = get_column_letter(col_idx)
            max_len = 10
            for cell in ws[col_letter]:
                if cell.value:
                    max_len = min(max(max_len, len(str(cell.value))), 80)
            ws.column_dimensions[col_letter].width = max_len + 2

    wb.save(path)


def run(
    *,
    gms_url: str,
    token: Optional[str],
    platform_instance: str,
    env: str,
    query: str,
    batch_size: int,
    max_datasets: int,
    hive_chunk_size: int,
    include_no_upstream: bool,
    jsonl_path: Path,
    xlsx_path: Path,
) -> AuditResult:
    graph = _make_graph(gms_url, token)
    print(
        f"[INFO] 扫描 DataHub datasets: platform=hive platform_instance={platform_instance} env={env} query={query!r}"
    )
    dataset_urns = fetch_dataset_urns(
        graph,
        platform_instance=platform_instance,
        env=env,
        query=query,
        batch_size=batch_size,
        max_datasets=max_datasets,
    )
    print(f"[INFO] DataHub dataset 数量: {len(dataset_urns)}")

    upstreams_by_target = fetch_upstreams_by_target(graph, dataset_urns)
    edge_count = sum(len(v) for v in upstreams_by_target.values())
    print(f"[INFO] 有表级上游血缘的目标表: {len(upstreams_by_target)} edge_count={edge_count}")

    try:
        view_dataset_urns = fetch_view_dataset_urns_from_mysql(platform_instance, env)
        etl_script_by_urn, data_availability_flags_by_urn = fetch_structured_properties_from_mysql(
            platform_instance,
            env,
        )
        etl_script_dataset_urns = {
            urn for urn, etl_script in etl_script_by_urn.items() if _has_meaningful_etl_script_value(etl_script)
        }
        documentation_by_urn = fetch_documentation_from_mysql(platform_instance, env)
        view_logic_by_urn = fetch_view_logic_from_mysql(platform_instance, env)
        schema_field_count_by_urn = fetch_schema_field_counts_from_mysql(platform_instance, env)
        print(
            f"[INFO] MySQL aspect 检查: views={len(view_dataset_urns)} "
            f"datasets_with_etl_script={len(etl_script_dataset_urns)} "
            f"documentation={len(documentation_by_urn)} flags={len(data_availability_flags_by_urn)} "
            f"view_logic={len(view_logic_by_urn)} schema={len(schema_field_count_by_urn)}"
        )
    except Exception as exc:
        print(f"[WARN] MySQL aspect 检查准备失败，跳过增强检查: {exc}", file=sys.stderr)
        view_dataset_urns = None
        etl_script_dataset_urns = None
        etl_script_by_urn = None
        data_availability_flags_by_urn = None
        documentation_by_urn = None
        view_logic_by_urn = None
        schema_field_count_by_urn = None

    needed_hive_tables = _collect_needed_hive_tables(
        dataset_urns,
        upstreams_by_target,
        platform_instance,
        env,
    )
    if documentation_by_urn is not None:
        needed_hive_tables.update(_collect_documented_hive_tables(documentation_by_urn))
    print(f"[INFO] Hive information_schema 待校验表数量: {len(needed_hive_tables)}")
    hive_existing = query_hive_existing_fqtns(needed_hive_tables, chunk_size=hive_chunk_size)
    print(f"[INFO] Hive information_schema 存在表数量: {len(hive_existing)}")

    result = evaluate_lineage_quality(
        dataset_urns=dataset_urns,
        upstreams_by_target=upstreams_by_target,
        hive_existing_fqtns=hive_existing,
        platform_instance=platform_instance,
        env=env,
        include_no_upstream=include_no_upstream,
        check_datahub_entity_existence=max_datasets <= 0,
        view_dataset_urns=view_dataset_urns,
        etl_script_dataset_urns=etl_script_dataset_urns,
        documentation_by_urn=documentation_by_urn,
        data_availability_flags_by_urn=data_availability_flags_by_urn,
        view_logic_by_urn=view_logic_by_urn,
        schema_field_count_by_urn=schema_field_count_by_urn,
        etl_script_by_urn=etl_script_by_urn,
    )
    write_jsonl(jsonl_path, result)
    write_xlsx(xlsx_path, result)

    print(f"[INFO] issue_count={result.summary['issue_count']}")
    for issue_type, count in result.summary["issue_counts_by_type"].items():
        print(f"[INFO]   {issue_type}: {count}")
    print(f"[DONE] jsonl: {jsonl_path}")
    print(f"[DONE] xlsx: {xlsx_path}")
    return result


def _parse_bool_env(name: str, default: bool = False) -> bool:
    raw = os.environ.get(name, "").strip().lower()
    if not raw:
        return default
    return raw in ("1", "true", "yes", "on")


def parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--gms-url", default=os.getenv("DATAHUB_GMS_URL", "http://localhost:8080"))
    p.add_argument("--token", default=os.getenv("DATAHUB_GMS_TOKEN"))
    p.add_argument(
        "--platform-instance",
        default=os.getenv("BLF_DATAHUB_PLATFORM_INSTANCE", DEFAULT_PLATFORM_INSTANCE),
    )
    p.add_argument("--env", default=os.getenv("DATAHUB_ENV", DEFAULT_ENV))
    p.add_argument("--query", default=os.getenv("LINEAGE_QUALITY_QUERY", "*"))
    p.add_argument("--batch-size", type=int, default=int(os.getenv("BATCH_SIZE", "2000")))
    p.add_argument("--max-datasets", type=int, default=int(os.getenv("MAX_DATASETS", "0")))
    p.add_argument(
        "--hive-chunk-size",
        type=int,
        default=int(os.getenv("HIVE_CHUNK_SIZE", "1000")),
        help="Hive information_schema IN 查询分批大小，默认 1000",
    )
    p.add_argument(
        "--include-no-upstream",
        action="store_true",
        default=_parse_bool_env("CHECK_NO_UPSTREAM", False),
        help="额外检查非 ODS/源层表无上游血缘；可能产生较多结果",
    )
    p.add_argument("--jsonl", required=True, type=Path)
    p.add_argument("--xlsx", required=True, type=Path)
    return p.parse_args()


def main() -> int:
    args = parse_args()
    try:
        run(
            gms_url=args.gms_url,
            token=args.token,
            platform_instance=args.platform_instance,
            env=args.env,
            query=args.query,
            batch_size=args.batch_size,
            max_datasets=args.max_datasets,
            hive_chunk_size=args.hive_chunk_size,
            include_no_upstream=args.include_no_upstream,
            jsonl_path=args.jsonl,
            xlsx_path=args.xlsx,
        )
    except Exception as exc:
        print(f"[ERROR] {exc}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
