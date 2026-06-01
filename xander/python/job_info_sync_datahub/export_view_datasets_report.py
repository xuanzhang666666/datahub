#!/usr/bin/env python3
"""Export Hive view datasets from DataHub (MySQL metadata_aspect_v2).

Output columns:
  - view name (db.table)
  - upstream table count
  - data_availability_flag (structured property)
  - whether View Definition (viewProperties.viewLogic) has meaningful content (是/否)
"""

from __future__ import annotations

import argparse
import csv
import json
import os
import sys
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any, Optional

from .check_dataset_availability import is_meaningful_text
from .query_upstream_lineage import extract_data_availability_flag
from .table_documentation_full_discovery import (
    _parse_tab_rows,
    _run_mysql_query,
    parse_hive_dataset_urn,
    table_name_matches_prefix,
)

DEFAULT_PLATFORM_INSTANCE = "blf-prod-hive"
DEFAULT_ENV = "PROD"


@dataclass(frozen=True)
class ViewDatasetExportRow:
    view_name: str
    upstream_count: int
    data_availability_flag: str
    view_definition_has_content: str


def _hive_urn_like(platform_instance: str, env: str) -> str:
    return (
        f"'urn:li:dataset:(urn:li:dataPlatform:hive,"
        f"{platform_instance}.%,{env})'"
    )


def _load_view_properties_rows(platform_instance: str, env: str) -> list[tuple[str, str]]:
    sql = (
        "select urn, metadata "
        "from metadata_aspect_v2 "
        "where aspect='viewProperties' "
        "and version=0 "
        f"and urn like {_hive_urn_like(platform_instance, env)} "
        "order by urn"
    )
    return [(urn, metadata) for urn, metadata in _parse_tab_rows(_run_mysql_query(sql), 2)]


def _load_structured_properties_by_urn(platform_instance: str, env: str) -> dict[str, str]:
    sql = (
        "select urn, metadata "
        "from metadata_aspect_v2 "
        "where aspect='structuredProperties' "
        "and version=0 "
        f"and urn like {_hive_urn_like(platform_instance, env)}"
    )
    out: dict[str, str] = {}
    for urn, metadata in _parse_tab_rows(_run_mysql_query(sql), 2):
        out[urn] = metadata
    return out


def _load_upstream_lineage_by_urn(platform_instance: str, env: str) -> dict[str, str]:
    sql = (
        "select urn, metadata "
        "from metadata_aspect_v2 "
        "where aspect='upstreamLineage' "
        "and version=0 "
        f"and urn like {_hive_urn_like(platform_instance, env)}"
    )
    out: dict[str, str] = {}
    for urn, metadata in _parse_tab_rows(_run_mysql_query(sql), 2):
        out[urn] = metadata
    return out


def _unwrap_aspect_dict(raw: Any) -> dict[str, Any]:
    current: Any = raw
    if isinstance(current, str):
        try:
            current = json.loads(current)
        except json.JSONDecodeError:
            return {}
    if not isinstance(current, dict):
        return {}
    for key in ("value",):
        nested = current.get(key)
        if isinstance(nested, dict) and not any(
            field in current for field in ("viewLogic", "upstreams", "properties")
        ):
            current = nested
    return current


def view_logic_from_view_properties_metadata(metadata: str) -> str:
    aspect = _unwrap_aspect_dict(metadata)
    view_logic = aspect.get("viewLogic")
    return view_logic if isinstance(view_logic, str) else ""


def upstream_count_from_lineage_metadata(metadata: str) -> int:
    aspect = _unwrap_aspect_dict(metadata)
    upstreams = aspect.get("upstreams")
    if not isinstance(upstreams, list):
        return 0
    count = 0
    for item in upstreams:
        if isinstance(item, dict) and item.get("dataset"):
            count += 1
    return count


def data_availability_flag_from_structured_metadata(metadata: str) -> str:
    aspect = _unwrap_aspect_dict(metadata)
    payload = {"structuredProperties": {"value": aspect}}
    flag = extract_data_availability_flag(payload)
    return flag if flag else "-"


def build_view_export_rows(
    *,
    platform_instance: str,
    env: str,
    table_prefix: str = "",
) -> list[ViewDatasetExportRow]:
    view_rows = _load_view_properties_rows(platform_instance, env)
    structured_by_urn = _load_structured_properties_by_urn(platform_instance, env)
    lineage_by_urn = _load_upstream_lineage_by_urn(platform_instance, env)

    rows: list[ViewDatasetExportRow] = []
    for urn, view_metadata in view_rows:
        table_name = parse_hive_dataset_urn(urn, platform_instance=platform_instance, env=env)
        if not table_name:
            continue
        if table_prefix and not table_name_matches_prefix(table_name, table_prefix):
            continue
        view_logic = view_logic_from_view_properties_metadata(view_metadata)
        structured_metadata = structured_by_urn.get(urn, "")
        lineage_metadata = lineage_by_urn.get(urn, "")
        rows.append(
            ViewDatasetExportRow(
                view_name=table_name,
                upstream_count=upstream_count_from_lineage_metadata(lineage_metadata),
                data_availability_flag=data_availability_flag_from_structured_metadata(
                    structured_metadata
                ),
                view_definition_has_content=(
                    "是" if is_meaningful_text(view_logic) else "否"
                ),
            )
        )
    rows.sort(key=lambda row: row.view_name)
    return rows


def write_csv(path: Path, rows: list[ViewDatasetExportRow]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=[
                "view_name",
                "upstream_count",
                "data_availability_flag",
                "view_definition_has_content",
            ],
        )
        writer.writeheader()
        for row in rows:
            writer.writerow(asdict(row))


def write_xlsx(path: Path, rows: list[ViewDatasetExportRow]) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise RuntimeError("需要 openpyxl: pip install openpyxl") from exc

    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "views"
    headers = [
        "view_name",
        "upstream_count",
        "data_availability_flag",
        "view_definition_has_content",
    ]
    ws.append(headers)
    for row in rows:
        ws.append(
            [
                row.view_name,
                row.upstream_count,
                row.data_availability_flag,
                row.view_definition_has_content,
            ]
        )
    for col_idx, _ in enumerate(headers, start=1):
        col_letter = get_column_letter(col_idx)
        max_len = len(headers[col_idx - 1])
        for cell in ws[col_letter]:
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max_len + 2, 80)
    wb.save(path)


def parse_args(argv: Optional[list[str]] = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--output-csv",
        required=True,
        help="输出 CSV 路径（UTF-8 BOM，便于 Excel 打开）",
    )
    parser.add_argument("--output-xlsx", help="可选 XLSX 路径")
    parser.add_argument(
        "--platform-instance",
        default=os.getenv("BLF_DATAHUB_PLATFORM_INSTANCE", DEFAULT_PLATFORM_INSTANCE),
    )
    parser.add_argument("--env", default=os.getenv("DATAHUB_ENV", DEFAULT_ENV))
    parser.add_argument(
        "--table-prefix",
        default=os.getenv("TABLE_PRE", "").strip(),
        help="只导出表名（点号后一段）以此前缀开头的 view",
    )
    return parser.parse_args(argv)


def main(argv: Optional[list[str]] = None) -> int:
    args = parse_args(argv)
    rows = build_view_export_rows(
        platform_instance=args.platform_instance,
        env=args.env,
        table_prefix=args.table_prefix,
    )
    csv_path = Path(args.output_csv)
    write_csv(csv_path, rows)
    print(f"[DONE] view_count={len(rows)} csv={csv_path}")
    if args.output_xlsx:
        xlsx_path = Path(args.output_xlsx)
        write_xlsx(xlsx_path, rows)
        print(f"[DONE] xlsx={xlsx_path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
