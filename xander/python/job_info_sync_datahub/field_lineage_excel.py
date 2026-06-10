"""Excel review files for field-level lineage candidates."""

from __future__ import annotations

import re
from dataclasses import replace
from datetime import datetime, timezone
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Set, Tuple

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .field_lineage_models import (
    FieldLineageCandidate,
    FieldLineageInput,
    FieldLineageReviewStatus,
    SourceTableValidation,
    UnresolvedField,
)
from .field_lineage_constants import (
    constant_expression_from_reason,
    has_explained_transform,
    is_constant_transform_expression,
    is_placeholder_transform_text,
)
from .field_lineage_policy import (
    ephemeral_source_fold_note,
    has_residual_sql_alias_in_source_table,
    has_unsafe_source_field_pattern,
    is_ephemeral_source_table,
    is_partition_field,
    is_self_dependency,
    normalize_source_field_name,
    normalize_source_table_name,
    resolve_ephemeral_source_table,
)

CANDIDATE_HEADERS = [
    "review_status",
    "target_table",
    "target_field",
    "source_table",
    "source_field",
    "transform_expression",
    "transform_explanation",
    "evidence_sql",
    "confidence",
    "llm_notes",
    "reviewer_notes",
    "import_error",
]

UNRESOLVED_HEADERS = ["target_field", "reason"]
CONTEXT_HEADERS = ["key", "value"]
SOURCE_VALIDATION_HEADERS = [
    "source_table",
    "dataset_exists",
    "in_target_upstreams",
    "validation_error",
]
DEFAULT_IMPORT_STATUSES = {
    FieldLineageReviewStatus.APPROVED,
    FieldLineageReviewStatus.AUTO_APPROVED,
}
_CREATE_TABLE_LIKE_RE = re.compile(
    r"\bcreate\s+(?:external\s+)?table\s+(?:if\s+not\s+exists\s+)?"
    r"(?P<target>`?[\w.]+`?)\s+like\s+(?P<source>`?[\w.]+`?)",
    re.I,
)


def _style_sheet(ws) -> None:
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9E1F2")
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for idx, _ in enumerate(ws[1], start=1):
        letter = get_column_letter(idx)
        max_len = 12
        for cell in ws[letter]:
            if cell.value:
                max_len = min(max(max_len, len(str(cell.value))), 80)
        ws.column_dimensions[letter].width = max_len + 2
    ws.freeze_panes = "A2"


def _effective_review_status(candidate: FieldLineageCandidate) -> FieldLineageReviewStatus:
    if candidate.import_error.strip():
        return FieldLineageReviewStatus.NEEDS_REVIEW
    if candidate.review_status in {
        FieldLineageReviewStatus.APPROVED,
        FieldLineageReviewStatus.REJECTED,
        FieldLineageReviewStatus.NEEDS_REVIEW,
        FieldLineageReviewStatus.NEEDS_FIX,
    }:
        return candidate.review_status
    if candidate.review_status in {
        FieldLineageReviewStatus.PENDING,
        FieldLineageReviewStatus.AUTO_APPROVED,
    }:
        if _is_safe_for_auto_approval(candidate):
            return FieldLineageReviewStatus.AUTO_APPROVED
        return FieldLineageReviewStatus.NEEDS_REVIEW
    return candidate.review_status


def _is_safe_for_auto_approval(candidate: FieldLineageCandidate) -> bool:
    required_values = [
        candidate.target_table,
        candidate.target_field,
        candidate.transform_expression,
    ]
    if any(not value.strip() for value in required_values):
        return False
    if any(
        is_placeholder_transform_text(value)
        for value in (
            candidate.transform_expression,
            candidate.transform_explanation,
            candidate.evidence_sql,
        )
    ):
        return False
    has_source_table = bool(candidate.source_table.strip())
    has_source_field = bool(candidate.source_field.strip())
    has_source = has_source_table and has_source_field
    has_partial_source = has_source_table != has_source_field
    is_constant = is_constant_transform_expression(candidate.transform_expression)
    if has_partial_source:
        return False
    if has_unsafe_source_field_pattern(candidate.source_field):
        return False
    if has_residual_sql_alias_in_source_table(candidate.source_table):
        return False
    if has_source_table and is_ephemeral_source_table(candidate.source_table):
        return False
    if has_source or is_constant:
        return True
    return has_explained_transform(
        candidate.transform_expression,
        candidate.transform_explanation,
    )


def _candidate_to_row(candidate: FieldLineageCandidate) -> List[str]:
    return [
        _effective_review_status(candidate).value,
        candidate.target_table,
        candidate.target_field,
        candidate.source_table,
        candidate.source_field,
        candidate.transform_expression,
        candidate.transform_explanation,
        candidate.evidence_sql,
        candidate.confidence,
        candidate.llm_notes,
        candidate.reviewer_notes,
        candidate.import_error,
    ]


def _canonical_candidate(
    source_input: FieldLineageInput,
    candidate: FieldLineageCandidate,
) -> FieldLineageCandidate:
    target_table = candidate.target_table.strip().lower()
    aliases = {alias.strip().lower() for alias in source_input.target_table_aliases}
    if target_table in aliases:
        target_table = source_input.table_name
    return FieldLineageCandidate(
        target_table=target_table,
        target_field=candidate.target_field,
        source_table=normalize_source_table_name(candidate.source_table),
        source_field=normalize_source_field_name(candidate.source_field),
        transform_expression=candidate.transform_expression,
        transform_explanation=candidate.transform_explanation,
        evidence_sql=candidate.evidence_sql,
        confidence=candidate.confidence,
        llm_notes=candidate.llm_notes,
        reviewer_notes=candidate.reviewer_notes,
        import_error=candidate.import_error,
        review_status=candidate.review_status,
    )


def _source_validation_error(validation: SourceTableValidation) -> str:
    if not validation.dataset_exists:
        return f"SOURCE_DATASET_NOT_FOUND: {validation.source_table}"
    if not validation.in_target_upstreams:
        return f"SOURCE_NOT_IN_TABLE_UPSTREAMS: {validation.source_table}"
    return ""


def _apply_source_table_validation(
    candidate: FieldLineageCandidate,
    source_table_validations: Optional[Dict[str, SourceTableValidation]],
) -> FieldLineageCandidate:
    if source_table_validations is None or not candidate.source_table.strip():
        return candidate
    validation = source_table_validations.get(candidate.source_table)
    validation_error = (
        _source_validation_error(validation)
        if validation is not None
        else f"SOURCE_VALIDATION_MISSING: {candidate.source_table}"
    )
    if not validation_error:
        return candidate
    errors = [value for value in (candidate.import_error, validation_error) if value]
    return replace(
        candidate,
        import_error="；".join(errors),
        review_status=FieldLineageReviewStatus.NEEDS_REVIEW,
    )


def _expand_placeholder_transform_text(
    candidate: FieldLineageCandidate,
    previous_by_target: Dict[Tuple[str, str], FieldLineageCandidate],
) -> FieldLineageCandidate:
    """Expand LLM shorthand such as "同上" using a prior row for the same target field."""
    key = (
        candidate.target_table.strip().lower(),
        candidate.target_field.strip().lower(),
    )
    previous = previous_by_target.get(key)

    def expand(value: str, previous_value: str) -> str:
        if is_placeholder_transform_text(value):
            return previous_value if not is_placeholder_transform_text(previous_value) else ""
        return value

    expanded = replace(
        candidate,
        transform_expression=expand(
            candidate.transform_expression,
            previous.transform_expression if previous else "",
        ),
        transform_explanation=expand(
            candidate.transform_explanation,
            previous.transform_explanation if previous else "",
        ),
        evidence_sql=expand(
            candidate.evidence_sql,
            previous.evidence_sql if previous else "",
        ),
    )
    if any(
        value.strip() and not is_placeholder_transform_text(value)
        for value in (
            expanded.transform_expression,
            expanded.transform_explanation,
            expanded.evidence_sql,
        )
    ):
        previous_by_target[key] = expanded
    return expanded


def _unresolved_to_candidate_row(
    source_input: FieldLineageInput,
    unresolved: UnresolvedField,
) -> List[str]:
    return [
        FieldLineageReviewStatus.NEEDS_REVIEW.value,
        source_input.table_name,
        unresolved.target_field,
        "",
        "",
        "",
        "",
        "",
        "",
        f"UNRESOLVED: {unresolved.reason}",
        "",
        "",
    ]


def _unresolved_constant_candidate(
    source_input: FieldLineageInput,
    unresolved: UnresolvedField,
) -> Optional[FieldLineageCandidate]:
    expression, evidence_sql = constant_expression_from_reason(
        unresolved.reason,
        unresolved.target_field,
    )
    if not expression:
        return None
    return FieldLineageCandidate(
        target_table=source_input.table_name,
        target_field=unresolved.target_field,
        source_table="",
        source_field="",
        transform_expression=expression,
        transform_explanation=(
            f"目标字段 {unresolved.target_field} 由常量 {expression} 写入，"
            "无上游来源字段。"
        ),
        evidence_sql=evidence_sql,
        confidence="HIGH",
        llm_notes=f"UNRESOLVED converted to constant field: {unresolved.reason}",
        review_status=FieldLineageReviewStatus.AUTO_APPROVED,
    )


def _is_target_partition_field(source_input: FieldLineageInput, field_name: str) -> bool:
    normalized = field_name.strip().lower()
    dynamic_partitions = {
        field.strip().lower()
        for field in source_input.target_partition_fields
        if field.strip()
    }
    return is_partition_field(normalized) or normalized in dynamic_partitions


def _target_schema_field_set(source_input: FieldLineageInput) -> Set[str]:
    return {
        field.strip().lower()
        for field in source_input.target_schema_fields
        if field.strip() and not _is_target_partition_field(source_input, field)
    }


def _is_known_target_field(target_field: str, schema_fields: Set[str]) -> bool:
    if not schema_fields:
        return True
    return target_field.strip().lower() in schema_fields


def _missing_schema_field_row(source_input: FieldLineageInput, target_field: str) -> List[str]:
    return [
        FieldLineageReviewStatus.NEEDS_REVIEW.value,
        source_input.table_name,
        target_field,
        "",
        "",
        "",
        "",
        "",
        "",
        "UNRESOLVED: DataHub DDL 字段未在 LLM 输出中找到，请按 Hive insert select 顺序确认来源",
        "",
        "",
    ]


def _normalize_table_ref(table_name: str, default_db: str) -> str:
    cleaned = table_name.strip().strip("`").lower()
    if "." in cleaned:
        return cleaned
    return f"{default_db}.{cleaned}"


def _create_like_source_table(source_input: FieldLineageInput) -> str:
    target_db = (
        source_input.table_name.rsplit(".", 1)[0]
        if "." in source_input.table_name
        else "default"
    )
    targets = {source_input.table_name.strip().lower()}
    targets.update(alias.strip().lower() for alias in source_input.target_table_aliases)
    for match in _CREATE_TABLE_LIKE_RE.finditer(source_input.etl_script or ""):
        target = _normalize_table_ref(match.group("target"), target_db)
        if target in targets:
            return _normalize_table_ref(match.group("source"), target_db)
    return ""


def _create_like_candidate(
    source_input: FieldLineageInput,
    source_table: str,
    target_field: str,
) -> FieldLineageCandidate:
    return FieldLineageCandidate(
        target_table=source_input.table_name,
        target_field=target_field,
        source_table=source_table,
        source_field=target_field,
        transform_expression=target_field,
        transform_explanation=(
            f"目标表通过 CREATE TABLE LIKE 复制 {source_table} 的表结构，"
            f"字段 {target_field} 与来源表同名字段一一对应。"
        ),
        evidence_sql=f"create table {source_input.table_name} like {source_table}",
        confidence="HIGH",
        review_status=FieldLineageReviewStatus.AUTO_APPROVED,
    )


def write_candidate_workbook(
    path: Path,
    *,
    source_input: FieldLineageInput,
    candidates: Iterable[FieldLineageCandidate],
    unresolved_fields: Iterable[UnresolvedField],
    llm_model: str,
    debug_dir: Optional[Path] = None,
    source_table_validations: Optional[Dict[str, SourceTableValidation]] = None,
) -> None:
    """Write reviewable field-lineage candidates to an Excel workbook."""
    wb = Workbook()
    ws = wb.active
    ws.title = "candidate_lineage"
    ws.append(CANDIDATE_HEADERS)
    schema_fields = _target_schema_field_set(source_input)
    emitted_target_fields: Set[str] = set()
    previous_by_target: Dict[Tuple[str, str], FieldLineageCandidate] = {}
    create_like_source_table = _create_like_source_table(source_input)
    for candidate in candidates:
        candidate = _canonical_candidate(source_input, candidate)
        candidate = _apply_source_table_validation(candidate, source_table_validations)
        if _is_target_partition_field(source_input, candidate.target_field):
            continue
        if is_self_dependency(candidate.target_table, candidate.source_table):
            continue
        if not _is_known_target_field(candidate.target_field, schema_fields):
            continue
        candidate = _expand_placeholder_transform_text(candidate, previous_by_target)
        emitted_target_fields.add(candidate.target_field.strip().lower())
        ws.append(_candidate_to_row(candidate))
    remaining_unresolved_fields: List[UnresolvedField] = []
    for unresolved in unresolved_fields:
        if _is_target_partition_field(source_input, unresolved.target_field):
            continue
        if not _is_known_target_field(unresolved.target_field, schema_fields):
            continue
        constant_candidate = _unresolved_constant_candidate(source_input, unresolved)
        if constant_candidate is not None:
            emitted_target_fields.add(unresolved.target_field.strip().lower())
            ws.append(_candidate_to_row(constant_candidate))
            continue
        remaining_unresolved_fields.append(unresolved)
        emitted_target_fields.add(unresolved.target_field.strip().lower())
        ws.append(_unresolved_to_candidate_row(source_input, unresolved))
    for target_field in source_input.target_schema_fields:
        normalized = target_field.strip().lower()
        if (
            not normalized
            or _is_target_partition_field(source_input, normalized)
            or normalized in emitted_target_fields
        ):
            continue
        if create_like_source_table:
            ws.append(
                _candidate_to_row(
                    _create_like_candidate(
                        source_input,
                        create_like_source_table,
                        normalized,
                    )
                )
            )
        else:
            ws.append(_missing_schema_field_row(source_input, normalized))
        emitted_target_fields.add(normalized)
    _style_sheet(ws)

    unresolved_ws = wb.create_sheet("unresolved_fields")
    unresolved_ws.append(UNRESOLVED_HEADERS)
    for unresolved in remaining_unresolved_fields:
        if _is_target_partition_field(source_input, unresolved.target_field):
            continue
        if not _is_known_target_field(unresolved.target_field, schema_fields):
            continue
        unresolved_ws.append([unresolved.target_field, unresolved.reason])
    _style_sheet(unresolved_ws)

    context_ws = wb.create_sheet("source_context")
    context_ws.append(CONTEXT_HEADERS)
    context_ws.append(["table_name", source_input.table_name])
    context_ws.append(["dataset_urn", source_input.dataset_urn])
    context_ws.append(["etl_script_chars", str(len(source_input.etl_script))])
    context_ws.append(["execute_shell_chars", str(len(source_input.execute_shell))])
    context_ws.append(["target_schema_field_count", str(len(source_input.target_schema_fields))])
    context_ws.append(["target_schema_fields", ",".join(source_input.target_schema_fields)])
    context_ws.append(["target_partition_fields", ",".join(source_input.target_partition_fields)])
    context_ws.append(["target_table_aliases", ",".join(source_input.target_table_aliases)])
    context_ws.append(["create_like_source_table", create_like_source_table])
    context_ws.append(["execute_shell", source_input.execute_shell])
    context_ws.append(["llm_model", llm_model])
    if debug_dir is not None:
        context_ws.append(["debug_artifacts_dir", str(debug_dir)])
    context_ws.append(["generated_at", datetime.now(timezone.utc).isoformat()])
    _style_sheet(context_ws)

    if source_table_validations is not None:
        validation_ws = wb.create_sheet("source_validation")
        validation_ws.append(SOURCE_VALIDATION_HEADERS)
        for source_table, validation in sorted(source_table_validations.items()):
            validation_ws.append(
                [
                    source_table,
                    validation.dataset_exists,
                    validation.in_target_upstreams,
                    _source_validation_error(validation),
                ]
            )
        _style_sheet(validation_ws)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def _row_dict(headers: List[str], values: List[object]) -> Dict[str, str]:
    out: Dict[str, str] = {}
    for idx, header in enumerate(headers):
        value = values[idx] if idx < len(values) else ""
        out[header] = "" if value is None else str(value).strip()
    return out


def sanitize_field_lineage_candidate(
    candidate: FieldLineageCandidate,
) -> FieldLineageCandidate:
    return replace(
        candidate,
        source_table=normalize_source_table_name(candidate.source_table),
        source_field=normalize_source_field_name(candidate.source_field),
    )


def fold_ephemeral_source_candidates(
    candidates: List[FieldLineageCandidate],
    upstream_tables: Set[str],
) -> List[FieldLineageCandidate]:
    """Fold tmp_/not_verified_ source_table names toward direct upstream tables."""
    folded: List[FieldLineageCandidate] = []
    for candidate in candidates:
        sanitized = sanitize_field_lineage_candidate(candidate)
        if not sanitized.source_table.strip():
            folded.append(sanitized)
            continue
        resolved, reason = resolve_ephemeral_source_table(
            sanitized.source_table,
            upstream_tables,
        )
        if reason.startswith("folded_"):
            note = ephemeral_source_fold_note(sanitized.source_table, resolved, reason)
            notes = "; ".join(item for item in (sanitized.llm_notes, note) if item)
            if is_ephemeral_source_table(resolved):
                err = (
                    f"EPHEMERAL_SOURCE_MUST_FOLD: {sanitized.source_table}"
                    f" (resolved={resolved}, {reason})"
                )
                errors = "; ".join(item for item in (sanitized.import_error, err) if item)
                folded.append(
                    replace(
                        sanitized,
                        source_table=resolved,
                        llm_notes=notes,
                        import_error=errors,
                        review_status=FieldLineageReviewStatus.NEEDS_REVIEW,
                    )
                )
            else:
                folded.append(replace(sanitized, source_table=resolved, llm_notes=notes))
            continue
        if is_ephemeral_source_table(sanitized.source_table):
            err = f"EPHEMERAL_SOURCE_MUST_FOLD: {sanitized.source_table}"
            errors = "; ".join(item for item in (sanitized.import_error, err) if item)
            folded.append(
                replace(
                    sanitized,
                    import_error=errors,
                    review_status=FieldLineageReviewStatus.NEEDS_REVIEW,
                )
            )
            continue
        folded.append(sanitized)
    return folded


def validate_import_candidates(
    candidates: List[FieldLineageCandidate],
    source_table_validations: Dict[str, SourceTableValidation],
) -> Tuple[List[FieldLineageCandidate], List[str]]:
    """Sanitize and reject import rows that are still unsafe or fail source validation."""
    accepted: List[FieldLineageCandidate] = []
    errors: List[str] = []
    for candidate in candidates:
        sanitized = sanitize_field_lineage_candidate(candidate)
        issues: List[str] = []
        has_source_table = bool(sanitized.source_table.strip())
        has_source_field = bool(sanitized.source_field.strip())
        if has_source_table != has_source_field:
            issues.append("SOURCE_TABLE_FIELD_MISMATCH")
        if has_unsafe_source_field_pattern(sanitized.source_field):
            issues.append(f"SOURCE_FIELD_UNSAFE: {sanitized.source_field}")
        if has_residual_sql_alias_in_source_table(sanitized.source_table):
            issues.append(f"SOURCE_TABLE_ALIAS: {sanitized.source_table}")
        if has_source_table and is_ephemeral_source_table(sanitized.source_table):
            issues.append(f"EPHEMERAL_SOURCE_TABLE: {sanitized.source_table}")
        if has_source_table:
            validation = source_table_validations.get(sanitized.source_table)
            validation_error = (
                _source_validation_error(validation)
                if validation is not None
                else f"SOURCE_VALIDATION_MISSING: {sanitized.source_table}"
            )
            if validation_error:
                issues.append(validation_error)
        if issues:
            errors.append(
                f"{sanitized.target_table}.{sanitized.target_field}: {'; '.join(issues)}"
            )
            continue
        accepted.append(sanitized)
    return accepted, errors


def _expand_multi_source_tables(rec: Dict[str, str]) -> List[Dict[str, str]]:
    """Expand a row whose source_table contains comma-separated table names.

    LLM sometimes writes multiple upstream tables in a single cell, e.g.
    "default.table_a, default.table_b".  We fan-out each table name into its
    own record so the importer can generate one FineGrainedLineage entry per
    upstream table.
    """
    raw = rec.get("source_table", "")
    tables = [t.strip() for t in raw.split(",") if t.strip()]
    if len(tables) <= 1:
        return [rec]
    expanded = []
    for tbl in tables:
        copy = dict(rec)
        copy["source_table"] = tbl
        expanded.append(copy)
    return expanded


def load_approved_review_rows(
    path: Path,
    *,
    import_statuses: Optional[Set[FieldLineageReviewStatus]] = None,
) -> List[FieldLineageCandidate]:
    """Read importable rows from a human-reviewed workbook.

    Rows whose ``source_table`` cell contains comma-separated table names are
    automatically expanded into one candidate per source table.
    """
    statuses = import_statuses or DEFAULT_IMPORT_STATUSES
    wb = load_workbook(path)
    if "candidate_lineage" not in wb.sheetnames:
        raise RuntimeError("Excel 缺少 candidate_lineage 工作表")
    ws = wb["candidate_lineage"]
    headers = [str(cell.value or "").strip() for cell in ws[1]]
    approved: List[FieldLineageCandidate] = []
    previous_by_target: Dict[Tuple[str, str], FieldLineageCandidate] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        rec = _row_dict(headers, list(row))
        status_value = rec.get("review_status", "").upper()
        if is_partition_field(rec.get("target_field", "")):
            continue
        for expanded_rec in _expand_multi_source_tables(rec):
            if is_self_dependency(
                expanded_rec.get("target_table", ""),
                expanded_rec.get("source_table", ""),
            ):
                continue
            candidate = sanitize_field_lineage_candidate(
                _expand_placeholder_transform_text(
                    FieldLineageCandidate(
                        target_table=expanded_rec.get("target_table", "").lower(),
                        target_field=expanded_rec.get("target_field", "").lower(),
                        source_table=expanded_rec.get("source_table", ""),
                        source_field=expanded_rec.get("source_field", ""),
                        transform_expression=expanded_rec.get("transform_expression", ""),
                        transform_explanation=expanded_rec.get("transform_explanation", ""),
                        evidence_sql=expanded_rec.get("evidence_sql", ""),
                        confidence=expanded_rec.get("confidence", "").upper(),
                        llm_notes=expanded_rec.get("llm_notes", ""),
                        reviewer_notes=expanded_rec.get("reviewer_notes", ""),
                        import_error=expanded_rec.get("import_error", ""),
                        review_status=FieldLineageReviewStatus.APPROVED,
                    ),
                    previous_by_target,
                )
            )
            if status_value not in {status.value for status in statuses}:
                continue
            approved.append(candidate)
    return approved
