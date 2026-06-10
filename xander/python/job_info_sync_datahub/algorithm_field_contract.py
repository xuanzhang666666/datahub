"""Build A-side algorithm field contracts from the DataHub field-lineage graph.

The CLI reads schemaMetadata, upstreamLineage, fineGrainedLineages, and the
Data Availability Flag. It recursively traces all non-partition input fields
without parsing ETL scripts, calling an LLM, or writing metadata back to DataHub.

Legacy ETL parsing helpers remain in this module for compatibility with older
callers and tests; the formal CLI path uses DataHub graph aspects only.
"""

from __future__ import annotations

import argparse
import json
import os
import re
import shlex
import urllib.error
import urllib.parse
import urllib.request
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any, Callable, Dict, Iterable, List, Optional, Sequence, Set, Tuple

from openpyxl import Workbook

from .field_lineage_datahub_reader import (
    extract_data_availability_flags,
    extract_schema_field_names,
    extract_schema_partition_field_names,
    fetch_structured_properties,
    make_hive_dataset_urn,
    strip_markdown_code_fence,
)
from .lineage_parser import build_lineage_summary, parse_block_lineage
from .models import FieldLineage, ParseStatus, SqlBlock, TableLineage
from .query_upstream_lineage import normalize_table_name, urn_to_table_name
from .sql_extractor import compute_format_date_vars, extract_sql_blocks
from .structured_properties import (
    URN_ETL_SCRIPT,
    URN_EXECUTE_SHELL,
    URN_SCHEDULE_URL,
)

OK = "OK"
MISSING_ETL_SCRIPT = "MISSING_ETL_SCRIPT"
NO_RUNTIME_VARS = "NO_RUNTIME_VARS"
SQL_PARSE_FAILED = "SQL_PARSE_FAILED"

COMPLETE = "COMPLETE"
MISSING_FIELD_LINEAGE = "MISSING_FIELD_LINEAGE"
INVALID_SCHEMA_FIELD_URN = "INVALID_SCHEMA_FIELD_URN"
DATASET_NOT_FOUND = "DATASET_NOT_FOUND"
CYCLE_DETECTED = "CYCLE_DETECTED"
MAX_DEPTH_REACHED = "MAX_DEPTH_REACHED"

CONFIRMED = "CONFIRMED"
RISK_UNCONFIRMED_NODE = "RISK_UNCONFIRMED_NODE"
INCOMPLETE = "INCOMPLETE"


@dataclass(frozen=True)
class DatasetJobProperties:
    table_name: str = ""
    dataset_urn: str = ""
    etl_script: str = ""
    execute_shell: str = ""
    schedule_url: str = ""
    property_urns: Dict[str, str] = field(default_factory=dict)


@dataclass(frozen=True)
class RuntimeVariableContext:
    raw_execute_shell: str
    variables: Dict[str, str]
    positionals: List[str]
    unparsed_tokens: List[str] = field(default_factory=list)


@dataclass(frozen=True)
class ResolvedEtlScript:
    original_script: str
    resolved_script: str
    unresolved_variables: List[str]


@dataclass
class TableContractAnalysis:
    table_name: str
    dataset_urn: str
    status: str
    properties: DatasetJobProperties
    runtime_context: RuntimeVariableContext
    resolved_etl_script: Optional[str] = None
    unresolved_variables: List[str] = field(default_factory=list)
    target_fields: List[str] = field(default_factory=list)
    unresolved_target_fields: List[str] = field(default_factory=list)
    sql_blocks: List[SqlBlock] = field(default_factory=list)
    table_lineages: List[TableLineage] = field(default_factory=list)
    field_lineages: List[FieldLineage] = field(default_factory=list)
    open_questions: List[str] = field(default_factory=list)


def _iter_property_assignments(payload: Dict[str, Any]) -> Iterable[Dict[str, Any]]:
    current: Any = payload
    for key in ("structuredProperties", "value"):
        if isinstance(current, dict) and key in current:
            current = current[key]
    if isinstance(current, dict) and isinstance(current.get("properties"), list):
        for item in current["properties"]:
            if isinstance(item, dict):
                yield item


def _first_string_value(assignment: Dict[str, Any]) -> str:
    values = assignment.get("values")
    if not isinstance(values, list):
        return ""
    for value in values:
        if isinstance(value, dict) and isinstance(value.get("string"), str):
            return value["string"]
        if isinstance(value, str):
            return value
    return ""


def _property_key(assignment: Dict[str, Any]) -> str:
    property_urn = assignment.get("propertyUrn")
    if isinstance(property_urn, str):
        return property_urn
    prop = assignment.get("property")
    if isinstance(prop, dict):
        urn = prop.get("urn")
        if isinstance(urn, str):
            return urn
        display_name = prop.get("displayName") or prop.get("name")
        if isinstance(display_name, str):
            return display_name
    display_name = assignment.get("displayName") or assignment.get("name")
    return display_name if isinstance(display_name, str) else ""


def extract_dataset_job_properties(
    payload: Dict[str, Any],
    *,
    table_name: str = "",
    dataset_urn: str = "",
) -> DatasetJobProperties:
    """Extract schedule_url, etl_script, and execute_shell from structuredProperties."""
    values: Dict[str, str] = {}
    for assignment in _iter_property_assignments(payload):
        key = _property_key(assignment)
        if key:
            values[key] = _first_string_value(assignment)

    def by_any(*keys: str) -> str:
        for key in keys:
            value = values.get(key, "")
            if value:
                return value
        return ""

    etl_script = strip_markdown_code_fence(
        by_any(URN_ETL_SCRIPT, "Etl Script", "表的etl 脚本", "表的etl脚本")
    )
    execute_shell = strip_markdown_code_fence(
        by_any(URN_EXECUTE_SHELL, "Execute Shell", "调度系统上配置的启动命令")
    )
    schedule_url = by_any(URN_SCHEDULE_URL, "Schedule URL", "调度系统任务链接").strip()

    return DatasetJobProperties(
        table_name=table_name,
        dataset_urn=dataset_urn,
        etl_script=etl_script,
        execute_shell=execute_shell,
        schedule_url=schedule_url,
        property_urns={
            "etl_script": URN_ETL_SCRIPT,
            "execute_shell": URN_EXECUTE_SHELL,
            "schedule_url": URN_SCHEDULE_URL,
        },
    )


def _strip_shell_quotes(value: str) -> str:
    value = value.strip()
    if len(value) >= 2 and value[0] == value[-1] and value[0] in {"'", '"'}:
        return value[1:-1]
    return value


def _parse_assignment_token(token: str) -> Optional[tuple[str, str]]:
    if "=" not in token or token.startswith("--"):
        return None
    key, value = token.split("=", 1)
    if not re.match(r"^[A-Za-z_][A-Za-z0-9_]*$", key):
        return None
    return key, _strip_shell_quotes(value)


def _add_date_derivatives(variables: Dict[str, str]) -> None:
    for key in ("DATE", "date", "dt", "bizdate", "BIZDATE"):
        value = variables.get(key)
        if value and re.match(r"^\d{8}$", value):
            for k, v in compute_format_date_vars(value).items():
                variables.setdefault(k, v)
            return


def parse_execute_shell_variables(execute_shell: str) -> RuntimeVariableContext:
    """Parse runtime variables from Execute Shell.

    Supports ``VAR=value``, ``export VAR=value``, ``--key=value``,
    ``--key value``, and preserves positional tokens for evidence.
    """
    raw = strip_markdown_code_fence(execute_shell or "")
    variables: Dict[str, str] = {}
    positionals: List[str] = []
    unparsed: List[str] = []

    try:
        tokens = shlex.split(raw, posix=True)
    except ValueError:
        tokens = raw.split()
        unparsed.append("SHLEX_PARSE_FAILED")

    i = 0
    while i < len(tokens):
        token = tokens[i]
        if token == "export":
            if i + 1 < len(tokens):
                parsed = _parse_assignment_token(tokens[i + 1])
                if parsed:
                    variables[parsed[0]] = parsed[1]
                    i += 2
                    continue
            i += 1
            continue

        parsed = _parse_assignment_token(token)
        if parsed:
            variables[parsed[0]] = parsed[1]
            i += 1
            continue

        if token.startswith("--"):
            option = token[2:]
            if "=" in option:
                key, value = option.split("=", 1)
                if key:
                    variables[key] = _strip_shell_quotes(value)
            else:
                key = option
                if key and i + 1 < len(tokens) and not tokens[i + 1].startswith("--"):
                    variables[key] = _strip_shell_quotes(tokens[i + 1])
                    i += 1
                elif key:
                    variables[key] = ""
            i += 1
            continue

        positionals.append(token)
        i += 1

    _add_date_derivatives(variables)
    return RuntimeVariableContext(
        raw_execute_shell=raw,
        variables=variables,
        positionals=positionals,
        unparsed_tokens=unparsed,
    )


_BRACED_VAR_RE = re.compile(r"\$\{([A-Za-z_][A-Za-z0-9_]*)\}")
_DOLLAR_VAR_RE = re.compile(r"(?<!\$)\$([A-Za-z_][A-Za-z0-9_]*)\b")
_JINJA_VAR_RE = re.compile(r"\{\{\s*([A-Za-z_][A-Za-z0-9_]*)\s*\}\}")
_IGNORED_UNRESOLVED_VARS = {
    "HIVE",
    "hive",
    "HDFS",
    "hdfs",
    "PYTHON",
    "python",
    "CHECK",
}


def resolve_etl_script(
    etl_script: str,
    runtime_context: RuntimeVariableContext,
) -> ResolvedEtlScript:
    """Replace ETL script placeholders using runtime variables from Execute Shell."""
    variables = runtime_context.variables
    unresolved: Set[str] = set()

    def replace_var(match: re.Match[str]) -> str:
        name = match.group(1)
        if name in variables:
            return variables[name]
        if name not in _IGNORED_UNRESOLVED_VARS:
            unresolved.add(name)
        return match.group(0)

    resolved = _BRACED_VAR_RE.sub(replace_var, etl_script or "")
    resolved = _DOLLAR_VAR_RE.sub(replace_var, resolved)
    resolved = _JINJA_VAR_RE.sub(replace_var, resolved)
    return ResolvedEtlScript(
        original_script=etl_script or "",
        resolved_script=resolved,
        unresolved_variables=sorted(unresolved),
    )


def _date_for_sql_extractor(runtime_context: RuntimeVariableContext) -> Optional[str]:
    for key in ("DATE", "date", "dt", "bizdate", "BIZDATE"):
        value = runtime_context.variables.get(key)
        if value and re.match(r"^\d{8}$", value):
            return value
    return None


def _job_file_name(table_name: str, etl_script: str) -> str:
    stripped = etl_script.lstrip()
    if stripped.startswith("#!") or re.search(r"\bpython\b|spark\.sql|cursor\.execute", stripped[:500], re.I):
        return f"{table_name}.structured_property.py"
    return f"{table_name}.structured_property.job"


def _filter_lineages_for_target(
    table_lineages: List[TableLineage],
    field_lineages: List[FieldLineage],
    table_name: str,
) -> tuple[List[TableLineage], List[FieldLineage]]:
    target = normalize_table_name(table_name)
    return (
        [lineage for lineage in table_lineages if lineage.target.full_name.lower() == target],
        [lineage for lineage in field_lineages if lineage.target_table.full_name.lower() == target],
    )


def _unresolved_target_fields(
    target_fields: Sequence[str],
    field_lineages: Sequence[FieldLineage],
) -> List[str]:
    parsed_fields = {lineage.target_field.lower() for lineage in field_lineages}
    unresolved: List[str] = []
    for field_name in target_fields:
        normalized = field_name.strip()
        if normalized and normalized.lower() not in parsed_fields:
            unresolved.append(normalized)
    return unresolved


def analyze_table_payload(
    table_name: str,
    dataset_urn: str,
    payload: Dict[str, Any],
    *,
    target_fields: Optional[Sequence[str]] = None,
) -> TableContractAnalysis:
    """Analyze one target table from a DataHub structuredProperties payload."""
    normalized_table = normalize_table_name(table_name)
    props = extract_dataset_job_properties(
        payload,
        table_name=normalized_table,
        dataset_urn=dataset_urn,
    )

    if not props.etl_script.strip():
        return TableContractAnalysis(
            table_name=normalized_table,
            dataset_urn=dataset_urn,
            status=MISSING_ETL_SCRIPT,
            properties=props,
            runtime_context=RuntimeVariableContext(props.execute_shell, {}, []),
            target_fields=list(target_fields or []),
            unresolved_target_fields=list(target_fields or []),
            open_questions=[f"{normalized_table}: missing {URN_ETL_SCRIPT}"],
        )

    runtime = parse_execute_shell_variables(props.execute_shell)
    resolved = resolve_etl_script(props.etl_script, runtime)
    job_file_name = _job_file_name(normalized_table, props.etl_script)
    blocks = extract_sql_blocks(
        resolved.resolved_script,
        job_file_name=job_file_name,
        date_str=_date_for_sql_extractor(runtime),
    )
    parsed_blocks = [parse_block_lineage(block, job_display_name=normalized_table) for block in blocks]
    table_lineages, field_lineages = build_lineage_summary(parsed_blocks)
    table_lineages, field_lineages = _filter_lineages_for_target(
        table_lineages,
        field_lineages,
        normalized_table,
    )

    open_questions: List[str] = []
    if not props.execute_shell.strip():
        open_questions.append(f"{normalized_table}: missing {URN_EXECUTE_SHELL}; parsed static ETL only")
    for variable in resolved.unresolved_variables:
        open_questions.append(f"{normalized_table}: unresolved variable {variable}")
    failed_blocks = [block for block in parsed_blocks if block.status == ParseStatus.SQL_PARSE_FAILED]
    for block in failed_blocks:
        open_questions.append(f"{normalized_table}: SQL block #{block.index} parse failed: {block.error_detail}")

    target_fields_list = list(target_fields or [])
    unresolved_target_fields = _unresolved_target_fields(target_fields_list, field_lineages)
    for field_name in unresolved_target_fields:
        open_questions.append(f"{normalized_table}.{field_name}: target field has no parsed upstream field lineage")

    status = OK
    if not props.execute_shell.strip():
        status = NO_RUNTIME_VARS
    elif failed_blocks and not table_lineages:
        status = SQL_PARSE_FAILED

    return TableContractAnalysis(
        table_name=normalized_table,
        dataset_urn=dataset_urn,
        status=status,
        properties=props,
        runtime_context=runtime,
        resolved_etl_script=resolved.resolved_script,
        unresolved_variables=resolved.unresolved_variables,
        target_fields=target_fields_list,
        unresolved_target_fields=unresolved_target_fields,
        sql_blocks=parsed_blocks,
        table_lineages=table_lineages,
        field_lineages=field_lineages,
        open_questions=open_questions,
    )


StructuredPayloadFetcher = Callable[[str, str], Dict[str, Any]]


def _dataset_aspect_url(gms_url: str, dataset_urn: str, aspect_name: str) -> str:
    encoded = urllib.parse.quote(dataset_urn, safe="")
    return f"{gms_url.rstrip('/')}/openapi/v3/entity/dataset/{encoded}/{aspect_name}"


def fetch_schema_fields(
    gms_url: str,
    dataset_urn: str,
    token: Optional[str] = None,
    timeout_sec: int = 60,
) -> List[str]:
    """Fetch field paths from DataHub schemaMetadata for one dataset."""
    req = urllib.request.Request(
        _dataset_aspect_url(gms_url, dataset_urn, "schemaMetadata"),
        method="GET",
        headers={"Accept": "application/json"},
    )
    if token:
        req.add_header("Authorization", f"Bearer {token}")
    try:
        with urllib.request.urlopen(req, timeout=timeout_sec) as resp:
            payload = json.loads(resp.read().decode("utf-8"))
    except urllib.error.HTTPError as exc:
        if exc.code == 404:
            return []
        raise

    current: Any = payload
    for key in ("schemaMetadata", "value"):
        if isinstance(current, dict) and key in current:
            current = current[key]
    fields = current.get("fields") if isinstance(current, dict) else None
    if not isinstance(fields, list):
        return []
    result: List[str] = []
    for field in fields:
        if isinstance(field, dict) and isinstance(field.get("fieldPath"), str):
            result.append(field["fieldPath"])
    return result


def analyze_tables(
    table_names: Sequence[str],
    *,
    gms_url: str,
    token: Optional[str] = None,
    platform_instance: str = "blf-prod-hive",
    env: str = "PROD",
    max_depth: int = 3,
    fetcher: Optional[StructuredPayloadFetcher] = None,
) -> List[TableContractAnalysis]:
    """Analyze target tables and recursively analyze parsed upstream tables."""
    if fetcher is None:
        def fetcher(_: str, urn: str) -> Dict[str, Any]:
            return fetch_structured_properties(gms_url, urn, token=token)

    results: List[TableContractAnalysis] = []
    seen: Set[str] = set()
    queue: List[tuple[str, int]] = [
        (normalize_table_name(table_name), 0) for table_name in table_names
    ]

    while queue:
        table_name, depth = queue.pop(0)
        if table_name in seen or depth > max_depth:
            continue
        seen.add(table_name)

        dataset_urn = make_hive_dataset_urn(table_name, platform_instance, env)
        try:
            payload = fetcher(table_name, dataset_urn)
        except Exception as exc:
            props = DatasetJobProperties(table_name=table_name, dataset_urn=dataset_urn)
            results.append(
                TableContractAnalysis(
                    table_name=table_name,
                    dataset_urn=dataset_urn,
                    status="FETCH_STRUCTURED_PROPERTIES_FAILED",
                    properties=props,
                    runtime_context=RuntimeVariableContext("", {}, []),
                    open_questions=[f"{table_name}: fetch structuredProperties failed: {exc}"],
                )
            )
            continue

        try:
            target_fields = fetch_schema_fields(gms_url, dataset_urn, token=token)
        except Exception:
            target_fields = []

        result = analyze_table_payload(
            table_name,
            dataset_urn,
            payload,
            target_fields=target_fields,
        )
        results.append(result)

        if depth >= max_depth:
            continue
        for lineage in result.table_lineages:
            for upstream in lineage.upstreams:
                upstream_name = normalize_table_name(upstream.full_name)
                if upstream_name not in seen:
                    queue.append((upstream_name, depth + 1))

    return results


def _field_mapping_to_dict(mapping: Any) -> Dict[str, Any]:
    return {
        "source_table": mapping.source_table,
        "source_field": mapping.source_field,
        "expression": mapping.expression,
        "confidence": mapping.confidence.value,
    }


def _table_result_to_json(result: TableContractAnalysis) -> Dict[str, Any]:
    return {
        "table_name": result.table_name,
        "dataset_urn": result.dataset_urn,
        "status": result.status,
        "schedule_url": result.properties.schedule_url,
        "unresolved_variables": result.unresolved_variables,
        "target_fields": result.target_fields,
        "unresolved_target_fields": result.unresolved_target_fields,
        "table_lineages": [
            {
                "target_table": lineage.target.full_name,
                "upstream_tables": [upstream.full_name for upstream in lineage.upstreams],
                "source_block_indices": lineage.source_block_indices,
            }
            for lineage in result.table_lineages
        ],
        "field_lineages": [
            {
                "target_table": lineage.target_table.full_name,
                "target_field": lineage.target_field,
                "confidence": lineage.confidence.value,
                "mappings": [_field_mapping_to_dict(mapping) for mapping in lineage.mappings],
            }
            for lineage in result.field_lineages
        ],
        "open_questions": result.open_questions,
    }


def write_contract_outputs(
    results: Sequence[TableContractAnalysis],
    output_dir: str | Path,
) -> None:
    """Write JSON, Markdown, and Excel artifacts for contract review."""
    out = Path(output_dir)
    out.mkdir(parents=True, exist_ok=True)

    (out / "field_lineage.json").write_text(
        json.dumps({"tables": [_table_result_to_json(result) for result in results]}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    (out / "runtime_context.json").write_text(
        json.dumps(
            {
                result.table_name: {
                    **asdict(result.runtime_context),
                    "resolved_etl_script_path": f"resolved_sql/{result.table_name.replace('.', '_')}.sql",
                }
                for result in results
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )

    resolved_dir = out / "resolved_sql"
    resolved_dir.mkdir(exist_ok=True)
    for result in results:
        if result.resolved_etl_script is not None:
            (resolved_dir / f"{result.table_name.replace('.', '_')}.sql").write_text(
                result.resolved_etl_script,
                encoding="utf-8",
            )

    _write_report(results, out / "lineage_report.md")
    _write_open_questions(results, out / "open_questions.md")
    _write_excel(results, out / "field_contract.xlsx")


def _write_report(results: Sequence[TableContractAnalysis], path: Path) -> None:
    lines = ["# A-side algorithm field contract", ""]
    for result in results:
        lines.extend(
            [
                f"## {result.table_name}",
                "",
                f"- status: `{result.status}`",
                f"- dataset_urn: `{result.dataset_urn}`",
                f"- schedule_url: {result.properties.schedule_url or '-'}",
        f"- unresolved_variables: {', '.join(result.unresolved_variables) or '-'}",
                f"- unresolved_target_fields: {', '.join(result.unresolved_target_fields) or '-'}",
                "",
                "| target_field | source_table | source_field | expression | confidence |",
                "| --- | --- | --- | --- | --- |",
            ]
        )
        for lineage in result.field_lineages:
            for mapping in lineage.mappings:
                lines.append(
                    "| {target} | {source_table} | {source_field} | {expr} | {confidence} |".format(
                        target=lineage.target_field,
                        source_table=mapping.source_table or "-",
                        source_field=mapping.source_field or "-",
                        expr=(mapping.expression or "-").replace("|", "\\|"),
                        confidence=mapping.confidence.value,
                    )
                )
        if not result.field_lineages:
            lines.append("| - | - | - | - | - |")
        lines.append("")
    path.write_text("\n".join(lines), encoding="utf-8")


def _write_open_questions(results: Sequence[TableContractAnalysis], path: Path) -> None:
    lines = ["# Open Questions", ""]
    has_questions = False
    for result in results:
        for question in result.open_questions:
            has_questions = True
            lines.append(f"- {question}")
    if not has_questions:
        lines.append("- 无")
    path.write_text("\n".join(lines), encoding="utf-8")


def _write_excel(results: Sequence[TableContractAnalysis], path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "field_contract"
    ws.append(
        [
            "table_name",
            "dataset_urn",
            "status",
            "target_field",
            "source_table",
            "source_field",
            "expression",
            "confidence",
            "schedule_url",
            "unresolved_variables",
            "unresolved_target_field",
        ]
    )
    for result in results:
        if not result.field_lineages:
            ws.append(
                [
                    result.table_name,
                    result.dataset_urn,
                    result.status,
                    "",
                    "",
                    "",
                    "",
                    "",
                    result.properties.schedule_url,
                    ", ".join(result.unresolved_variables),
                    ", ".join(result.unresolved_target_fields),
                ]
            )
            for field_name in result.unresolved_target_fields:
                ws.append(
                    [
                        result.table_name,
                        result.dataset_urn,
                        result.status,
                        field_name,
                        "",
                        "",
                        "",
                        "UNRESOLVED",
                        result.properties.schedule_url,
                        ", ".join(result.unresolved_variables),
                        field_name,
                    ]
                )
            continue
        for lineage in result.field_lineages:
            for mapping in lineage.mappings:
                ws.append(
                    [
                        result.table_name,
                        result.dataset_urn,
                        result.status,
                        lineage.target_field,
                        mapping.source_table or "",
                        mapping.source_field or "",
                        mapping.expression,
                        mapping.confidence.value,
                        result.properties.schedule_url,
                        ", ".join(result.unresolved_variables),
                        "",
                    ]
                )
        for field_name in result.unresolved_target_fields:
            ws.append(
                [
                    result.table_name,
                    result.dataset_urn,
                    result.status,
                    field_name,
                    "",
                    "",
                    "",
                    "UNRESOLVED",
                    result.properties.schedule_url,
                    ", ".join(result.unresolved_variables),
                    field_name,
                ]
            )
    wb.save(path)


@dataclass
class DatasetGraphSnapshot:
    table_name: str
    dataset_urn: str
    schema_fields: List[str]
    partition_fields: List[str]
    field_descriptions: Dict[str, str]
    table_upstreams: List[str]
    field_edges: Dict[str, List[Tuple[str, str]]]
    field_lineage_confirmed: bool
    raw_aspects: Dict[str, Dict[str, Any]]


@dataclass
class FieldTracePath:
    target_table: str
    target_field: str
    target_description: str
    nodes: List[str]
    transform_operations: List[str]
    path_status: str
    trust_status: str
    unconfirmed_nodes: List[str] = field(default_factory=list)
    question: str = ""

    @property
    def direct_upstream(self) -> str:
        return self.nodes[1] if len(self.nodes) > 1 else ""

    @property
    def final_source(self) -> str:
        return self.nodes[-1] if self.nodes else ""

    @property
    def path_length(self) -> int:
        return max(0, len(self.nodes) - 1)


@dataclass
class FieldContractSummary:
    target_table: str
    target_field: str
    target_description: str
    path_count: int
    path_status: str
    trust_status: str
    direct_upstreams: List[str]
    final_sources: List[str]
    transform_chains: List[str]
    unconfirmed_nodes: List[str]
    questions: List[str]


@dataclass
class FieldContractGraphResult:
    input_tables: List[str]
    max_depth: int
    fields: List[FieldContractSummary]
    paths: List[FieldTracePath]
    open_questions: List[str]
    snapshots: Dict[str, DatasetGraphSnapshot]


GraphSnapshotLoader = Callable[[str], DatasetGraphSnapshot]


def _split_field_node(node: str) -> Optional[Tuple[str, str]]:
    normalized = node.strip().lower()
    if normalized.startswith("__invalid__:"):
        return None
    parts = normalized.split(".")
    if len(parts) < 3 or not parts[-1]:
        return None
    return ".".join(parts[:-1]), parts[-1]


def _path_trust(status: str, unconfirmed_nodes: Sequence[str]) -> str:
    if status != COMPLETE:
        return INCOMPLETE
    if unconfirmed_nodes:
        return RISK_UNCONFIRMED_NODE
    return CONFIRMED


def trace_field_contract(
    table_names: Sequence[str],
    *,
    snapshot_loader: GraphSnapshotLoader,
    max_depth: int = 20,
) -> FieldContractGraphResult:
    """Trace all non-partition input fields through DataHub fine-grained lineage."""
    normalized_inputs = list(dict.fromkeys(normalize_table_name(name) for name in table_names))
    snapshots: Dict[str, DatasetGraphSnapshot] = {}
    paths: List[FieldTracePath] = []

    def load(table_name: str) -> DatasetGraphSnapshot:
        normalized = normalize_table_name(table_name)
        if normalized not in snapshots:
            snapshots[normalized] = snapshot_loader(normalized)
        return snapshots[normalized]

    def finish(
        *,
        root_table: str,
        root_field: str,
        description: str,
        nodes: List[str],
        operations: List[str],
        status: str,
        unconfirmed: List[str],
        question: str = "",
    ) -> None:
        paths.append(
            FieldTracePath(
                target_table=root_table,
                target_field=root_field,
                target_description=description,
                nodes=nodes,
                transform_operations=operations,
                path_status=status,
                trust_status=_path_trust(status, unconfirmed),
                unconfirmed_nodes=list(dict.fromkeys(unconfirmed)),
                question=question,
            )
        )

    def walk(
        *,
        root_table: str,
        root_field: str,
        description: str,
        node: str,
        nodes: List[str],
        operations: List[str],
        visited: Set[str],
        unconfirmed: List[str],
        depth: int,
    ) -> None:
        parsed = _split_field_node(node)
        if parsed is None:
            finish(
                root_table=root_table,
                root_field=root_field,
                description=description,
                nodes=nodes,
                operations=operations,
                status=INVALID_SCHEMA_FIELD_URN,
                unconfirmed=unconfirmed,
                question=f"{root_table}.{root_field}: invalid schemaField URN/node {node}",
            )
            return
        table_name, field_name = parsed
        try:
            snapshot = load(table_name)
        except Exception as exc:
            finish(
                root_table=root_table,
                root_field=root_field,
                description=description,
                nodes=nodes,
                operations=operations,
                status=DATASET_NOT_FOUND,
                unconfirmed=unconfirmed,
                question=f"{root_table}.{root_field}: dataset/schema unavailable at {node}: {exc}",
            )
            return

        current_unconfirmed = list(unconfirmed)
        if not snapshot.field_lineage_confirmed:
            current_unconfirmed.append(node)
        if field_name.lower() not in {field.lower() for field in snapshot.schema_fields}:
            finish(
                root_table=root_table,
                root_field=root_field,
                description=description,
                nodes=nodes,
                operations=operations,
                status=INVALID_SCHEMA_FIELD_URN,
                unconfirmed=current_unconfirmed,
                question=f"{root_table}.{root_field}: field {node} not found in schemaMetadata",
            )
            return
        edges = snapshot.field_edges.get(field_name.lower(), [])
        if not edges:
            if snapshot.table_upstreams:
                finish(
                    root_table=root_table,
                    root_field=root_field,
                    description=description,
                    nodes=nodes,
                    operations=operations,
                    status=MISSING_FIELD_LINEAGE,
                    unconfirmed=current_unconfirmed,
                    question=(
                        f"{root_table}.{root_field}: {node} has table upstreams but no field lineage"
                    ),
                )
            else:
                finish(
                    root_table=root_table,
                    root_field=root_field,
                    description=description,
                    nodes=nodes,
                    operations=operations,
                    status=COMPLETE,
                    unconfirmed=current_unconfirmed,
                )
            return
        if depth >= max_depth:
            finish(
                root_table=root_table,
                root_field=root_field,
                description=description,
                nodes=nodes,
                operations=operations,
                status=MAX_DEPTH_REACHED,
                unconfirmed=current_unconfirmed,
                question=f"{root_table}.{root_field}: max depth {max_depth} reached at {node}",
            )
            return

        for upstream_node, operation in edges:
            if upstream_node in visited:
                finish(
                    root_table=root_table,
                    root_field=root_field,
                    description=description,
                    nodes=[*nodes, upstream_node],
                    operations=[*operations, operation],
                    status=CYCLE_DETECTED,
                    unconfirmed=current_unconfirmed,
                    question=f"{root_table}.{root_field}: cycle detected at {upstream_node}",
                )
                continue
            walk(
                root_table=root_table,
                root_field=root_field,
                description=description,
                node=upstream_node,
                nodes=[*nodes, upstream_node],
                operations=[*operations, operation],
                visited={*visited, upstream_node},
                unconfirmed=current_unconfirmed,
                depth=depth + 1,
            )

    for table_name in normalized_inputs:
        try:
            root_snapshot = load(table_name)
        except Exception as exc:
            finish(
                root_table=table_name,
                root_field="*",
                description="",
                nodes=[f"{table_name}.*"],
                operations=[],
                status=DATASET_NOT_FOUND,
                unconfirmed=[],
                question=f"{table_name}: dataset/schema unavailable: {exc}",
            )
            continue
        partition_fields = {field.lower() for field in root_snapshot.partition_fields}
        for field_name in root_snapshot.schema_fields:
            normalized_field = field_name.lower()
            if normalized_field in partition_fields:
                continue
            node = f"{table_name}.{normalized_field}"
            walk(
                root_table=table_name,
                root_field=normalized_field,
                description=root_snapshot.field_descriptions.get(normalized_field, ""),
                node=node,
                nodes=[node],
                operations=[],
                visited={node},
                unconfirmed=[],
                depth=0,
            )

    field_summaries: List[FieldContractSummary] = []
    grouped: Dict[Tuple[str, str], List[FieldTracePath]] = {}
    for path in paths:
        grouped.setdefault((path.target_table, path.target_field), []).append(path)
    for (table_name, field_name), field_paths in sorted(grouped.items()):
        statuses = {path.path_status for path in field_paths}
        trust_statuses = {path.trust_status for path in field_paths}
        summary_status = COMPLETE if statuses == {COMPLETE} else INCOMPLETE
        if INCOMPLETE in trust_statuses:
            summary_trust = INCOMPLETE
        elif RISK_UNCONFIRMED_NODE in trust_statuses:
            summary_trust = RISK_UNCONFIRMED_NODE
        else:
            summary_trust = CONFIRMED
        field_summaries.append(
            FieldContractSummary(
                target_table=table_name,
                target_field=field_name,
                target_description=field_paths[0].target_description,
                path_count=len(field_paths),
                path_status=summary_status,
                trust_status=summary_trust,
                direct_upstreams=sorted({p.direct_upstream for p in field_paths if p.direct_upstream}),
                final_sources=sorted({p.final_source for p in field_paths if p.final_source}),
                transform_chains=[
                    " -> ".join(operation for operation in p.transform_operations if operation)
                    for p in field_paths
                ],
                unconfirmed_nodes=sorted(
                    {node for p in field_paths for node in p.unconfirmed_nodes}
                ),
                questions=[p.question for p in field_paths if p.question],
            )
        )

    questions = list(
        dict.fromkeys(
            [
                *(path.question for path in paths if path.question),
                *(
                    f"{path.target_table}.{path.target_field}: unconfirmed field lineage node {node}"
                    for path in paths
                    for node in path.unconfirmed_nodes
                ),
            ]
        )
    )
    return FieldContractGraphResult(
        input_tables=normalized_inputs,
        max_depth=max_depth,
        fields=field_summaries,
        paths=paths,
        open_questions=questions,
        snapshots=snapshots,
    )


def _unwrap_openapi_aspect(payload: Dict[str, Any], aspect_name: str) -> Dict[str, Any]:
    current: Any = payload
    for key in (aspect_name, "value"):
        if isinstance(current, dict) and key in current:
            current = current[key]
    return current if isinstance(current, dict) else {}


def _fetch_openapi_aspect(
    gms_url: str,
    dataset_urn: str,
    aspect_name: str,
    token: Optional[str],
    *,
    required: bool = False,
) -> Dict[str, Any]:
    req = urllib.request.Request(
        _dataset_aspect_url(gms_url, dataset_urn, aspect_name),
        method="GET",
        headers={"Accept": "application/json"},
    )
    if token:
        req.add_header("Authorization", f"Bearer {token}")
    try:
        with urllib.request.urlopen(req, timeout=60) as resp:
            return json.loads(resp.read().decode("utf-8"))
    except urllib.error.HTTPError as exc:
        if exc.code == 404 and not required:
            return {}
        detail = exc.read().decode("utf-8", errors="replace")
        raise RuntimeError(f"GET {aspect_name} HTTP {exc.code}: {detail}") from exc


def _schema_field_descriptions(payload: Dict[str, Any]) -> Dict[str, str]:
    aspect = _unwrap_openapi_aspect(payload, "schemaMetadata")
    fields = aspect.get("fields")
    descriptions: Dict[str, str] = {}
    if not isinstance(fields, list):
        return descriptions
    for item in fields:
        if not isinstance(item, dict):
            continue
        name = item.get("fieldPath")
        if not isinstance(name, str) or not name.strip():
            continue
        description = item.get("description")
        descriptions[name.strip().lower()] = description if isinstance(description, str) else ""
    return descriptions


def _schema_field_urn_to_node(
    urn: str,
    platform_instance: str,
) -> Optional[str]:
    prefix = "urn:li:schemaField:("
    if not isinstance(urn, str) or not urn.startswith(prefix) or not urn.endswith(")"):
        return None
    inner = urn[len(prefix) : -1]
    if "," not in inner:
        return None
    dataset_urn, field_name = inner.rsplit(",", 1)
    table_name = urn_to_table_name(dataset_urn.strip(), platform_instance)
    if table_name == dataset_urn.strip() or "." not in table_name:
        return None
    field_name = urllib.parse.unquote(field_name.strip())
    if not field_name:
        return None
    return f"{normalize_table_name(table_name)}.{field_name.lower()}"


def load_dataset_graph_snapshot(
    table_name: str,
    *,
    gms_url: str,
    token: Optional[str],
    platform_instance: str,
    env: str,
) -> DatasetGraphSnapshot:
    normalized = normalize_table_name(table_name)
    dataset_urn = make_hive_dataset_urn(normalized, platform_instance, env)
    schema_payload = _fetch_openapi_aspect(
        gms_url, dataset_urn, "schemaMetadata", token, required=True
    )
    lineage_payload = _fetch_openapi_aspect(gms_url, dataset_urn, "upstreamLineage", token)
    structured_payload = _fetch_openapi_aspect(
        gms_url, dataset_urn, "structuredProperties", token
    )
    schema_fields = extract_schema_field_names(schema_payload)
    if not schema_fields:
        raise RuntimeError("schemaMetadata fields empty")

    lineage_aspect = _unwrap_openapi_aspect(lineage_payload, "upstreamLineage")
    table_upstreams: List[str] = []
    for item in lineage_aspect.get("upstreams") or []:
        if not isinstance(item, dict) or not isinstance(item.get("dataset"), str):
            continue
        table_upstreams.append(
            normalize_table_name(urn_to_table_name(item["dataset"], platform_instance))
        )
    field_edges: Dict[str, List[Tuple[str, str]]] = {}
    for entry in lineage_aspect.get("fineGrainedLineages") or []:
        if not isinstance(entry, dict):
            continue
        operation = entry.get("transformOperation")
        operation_text = operation if isinstance(operation, str) else ""
        upstream_nodes: List[str] = []
        for upstream_urn in entry.get("upstreams") or []:
            node = _schema_field_urn_to_node(upstream_urn, platform_instance)
            upstream_nodes.append(node or f"__invalid__:{upstream_urn}")
        for downstream_urn in entry.get("downstreams") or []:
            downstream_node = _schema_field_urn_to_node(downstream_urn, platform_instance)
            parsed = _split_field_node(downstream_node or "")
            if parsed is None or parsed[0] != normalized:
                continue
            field_edges.setdefault(parsed[1], []).extend(
                (upstream_node, operation_text) for upstream_node in upstream_nodes
            )

    return DatasetGraphSnapshot(
        table_name=normalized,
        dataset_urn=dataset_urn,
        schema_fields=schema_fields,
        partition_fields=extract_schema_partition_field_names(schema_payload),
        field_descriptions=_schema_field_descriptions(schema_payload),
        table_upstreams=sorted(set(table_upstreams)),
        field_edges=field_edges,
        field_lineage_confirmed="字段血缘" in extract_data_availability_flags(structured_payload),
        raw_aspects={
            "schemaMetadata": schema_payload,
            "upstreamLineage": lineage_payload,
            "structuredProperties": structured_payload,
        },
    )


def _write_graph_excel(result: FieldContractGraphResult, path: Path) -> None:
    def tables(nodes: Sequence[str]) -> str:
        return "\n".join(
            sorted(
                {
                    parsed[0]
                    for node in nodes
                    if (parsed := _split_field_node(node)) is not None
                }
            )
        )

    def fields(nodes: Sequence[str]) -> str:
        return "\n".join(
            sorted(
                {
                    parsed[1]
                    for node in nodes
                    if (parsed := _split_field_node(node)) is not None
                }
            )
        )

    wb = Workbook()
    contract = wb.active
    contract.title = "field_contract"
    contract.append(
        [
            "算法目标表", "算法目标字段", "字段描述", "直接上游表", "直接上游字段",
            "最终源头表", "最终源头字段",
            "加工逻辑链", "路径数量", "路径状态", "信任状态", "未确认节点", "待确认问题",
            "是否必需", "B侧候选表", "B侧候选字段", "适配状态",
        ]
    )
    for item in result.fields:
        contract.append(
            [
                item.target_table, item.target_field, item.target_description,
                tables(item.direct_upstreams), fields(item.direct_upstreams),
                tables(item.final_sources), fields(item.final_sources),
                "\n".join(item.transform_chains), item.path_count, item.path_status,
                item.trust_status, "\n".join(item.unconfirmed_nodes), "\n".join(item.questions),
                "", "", "", "",
            ]
        )
    trace = wb.create_sheet("trace_paths")
    trace.append(
        [
            "算法目标表", "算法目标字段", "完整字段路径", "直接上游", "最终源头",
            "加工逻辑链", "路径长度", "路径状态", "信任状态", "未确认节点", "待确认问题",
        ]
    )
    for item in result.paths:
        trace.append(
            [
                item.target_table, item.target_field, " -> ".join(item.nodes),
                item.direct_upstream, item.final_source, "\n".join(item.transform_operations),
                item.path_length, item.path_status, item.trust_status,
                "\n".join(item.unconfirmed_nodes), item.question,
            ]
        )
    questions = wb.create_sheet("open_questions")
    questions.append(["待确认问题"])
    for question in result.open_questions:
        questions.append([question])
    summary = wb.create_sheet("run_summary")
    confirmed = sum(1 for item in result.fields if item.trust_status == CONFIRMED)
    complete = sum(1 for item in result.fields if item.path_status == COMPLETE)
    field_count = len(result.fields)
    for row in [
        ("input_table_count", len(result.input_tables)),
        ("target_field_count", field_count),
        ("trace_path_count", len(result.paths)),
        ("complete_field_count", complete),
        ("field_coverage_percent", round(complete * 100 / field_count, 2) if field_count else 0.0),
        ("confirmed_field_count", confirmed),
        ("confirmed_field_percent", round(confirmed * 100 / field_count, 2) if field_count else 0.0),
        ("open_question_count", len(result.open_questions)),
        ("max_depth", result.max_depth),
    ]:
        summary.append(row)
    wb.save(path)


def write_graph_contract_outputs(
    result: FieldContractGraphResult,
    output_dir: str | Path,
) -> None:
    out = Path(output_dir)
    out.mkdir(parents=True, exist_ok=True)
    payload = {
        "input_tables": result.input_tables,
        "max_depth": result.max_depth,
        "fields": [asdict(item) for item in result.fields],
        "paths": [asdict(item) for item in result.paths],
        "open_questions": result.open_questions,
    }
    (out / "field_lineage.json").write_text(
        json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    report_lines = ["# A-side algorithm field contract", ""]
    for table_name in result.input_tables:
        fields = [field for field in result.fields if field.target_table == table_name]
        report_lines.extend(
            [
                f"## {table_name}",
                "",
                f"- target_fields: {len(fields)}",
                f"- complete_fields: {sum(1 for field in fields if field.path_status == COMPLETE)}",
                f"- confirmed_fields: {sum(1 for field in fields if field.trust_status == CONFIRMED)}",
                "",
            ]
        )
    (out / "lineage_report.md").write_text("\n".join(report_lines), encoding="utf-8")
    question_lines = ["# Open Questions", ""]
    question_lines.extend(f"- {question}" for question in result.open_questions)
    if not result.open_questions:
        question_lines.append("- 无")
    (out / "open_questions.md").write_text("\n".join(question_lines), encoding="utf-8")
    raw_dir = out / "raw_aspects"
    for table_name, snapshot in result.snapshots.items():
        table_dir = raw_dir / table_name
        table_dir.mkdir(parents=True, exist_ok=True)
        for aspect_name, aspect_payload in snapshot.raw_aspects.items():
            (table_dir / f"{aspect_name}.json").write_text(
                json.dumps(aspect_payload, ensure_ascii=False, indent=2),
                encoding="utf-8",
            )
    _write_graph_excel(result, out / "field_contract.xlsx")


def _read_table_names(args: argparse.Namespace) -> List[str]:
    names: List[str] = []
    for value in args.table or []:
        for token in re.split(r"[,，\n]+", value):
            if token.strip():
                names.append(token.strip())
    if args.table_file:
        for line in Path(args.table_file).read_text(encoding="utf-8").splitlines():
            stripped = line.strip()
            if stripped and not stripped.startswith("#"):
                names.append(stripped)
    if not names:
        names.extend(table_names_from_env())
    return names


def table_names_from_env(env_var: str = "TABLE_NAMES") -> List[str]:
    """Read Jenkins TABLE_NAMES multi-line parameter.

    Supports newline, comma, and Chinese comma separators. Comment-only lines
    beginning with ``#`` are ignored.
    """
    raw = os.getenv(env_var, "")
    names: List[str] = []
    for line in raw.replace("\r", "\n").splitlines():
        stripped = line.strip()
        if not stripped or stripped.startswith("#"):
            continue
        for token in re.split(r"[,，]+", stripped):
            table_name = token.strip()
            if table_name and not table_name.startswith("#"):
                names.append(table_name)
    return names


def main(argv: Optional[Sequence[str]] = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--table", action="append", help="Target table name; repeatable or comma-separated")
    parser.add_argument("--table-file", help="File containing target table names")
    parser.add_argument("--gms-url", default=os.getenv("DATAHUB_GMS_URL", "http://localhost:8080"))
    parser.add_argument("--token", default=os.getenv("DATAHUB_GMS_TOKEN"))
    parser.add_argument("--platform-instance", default=os.getenv("BLF_DATAHUB_PLATFORM_INSTANCE", "blf-prod-hive"))
    parser.add_argument("--env", default=os.getenv("DATAHUB_ENV", "PROD"))
    parser.add_argument("--max-depth", type=int, default=int(os.getenv("FIELD_CONTRACT_MAX_DEPTH", "20")))
    parser.add_argument("--output-dir", default=os.getenv("FIELD_CONTRACT_OUTPUT_DIR", "tmp/algorithm_field_contract"))
    args = parser.parse_args(argv)

    table_names = _read_table_names(args)
    if not table_names:
        parser.error("at least one --table, --table-file, or TABLE_NAMES env value is required")

    result = trace_field_contract(
        table_names,
        snapshot_loader=lambda table_name: load_dataset_graph_snapshot(
            table_name,
            gms_url=args.gms_url,
            token=args.token,
            platform_instance=args.platform_instance,
            env=args.env,
        ),
        max_depth=args.max_depth,
    )
    write_graph_contract_outputs(result, args.output_dir)
    complete_fields = sum(1 for item in result.fields if item.path_status == COMPLETE)
    confirmed_fields = sum(1 for item in result.fields if item.trust_status == CONFIRMED)
    print(
        "field contract summary: "
        f"tables={len(result.input_tables)} fields={len(result.fields)} "
        f"paths={len(result.paths)} complete_fields={complete_fields} "
        f"confirmed_fields={confirmed_fields} open_questions={len(result.open_questions)}"
    )
    print(f"wrote contract outputs to {args.output_dir}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
