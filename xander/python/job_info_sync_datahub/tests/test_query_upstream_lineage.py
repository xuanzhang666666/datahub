"""query_upstream_lineage view-table structured property checks."""

from __future__ import annotations

import json
import urllib.error
import urllib.request

from openpyxl import load_workbook

from job_info_sync_datahub import query_upstream_lineage as q
from job_info_sync_datahub.structured_properties import (
    URN_DATA_AVAILABILITY_FLAG,
    URN_ETL_SCRIPT,
    URN_EXECUTE_SHELL,
    URN_OTHER_REMARK,
    URN_SCHEDULE_URL,
)


class _FakeResponse:
    def __init__(self, payload: dict) -> None:
        self._payload = payload

    def __enter__(self) -> "_FakeResponse":
        return self

    def __exit__(self, *args: object) -> None:
        return None

    def read(self) -> bytes:
        return json.dumps(self._payload).encode("utf-8")


def test_run_allows_view_upstreams_with_empty_etl_properties(monkeypatch) -> None:
    upstream = q.make_hive_dataset_urn("view_db.some_view")

    monkeypatch.setattr(q, "fetch_all_upstream_urns_with_counts", lambda *args, **kwargs: ({upstream}, {upstream: 0}))
    monkeypatch.setattr(q, "read_upstream_structured_status", lambda *args, **kwargs: (False, False, False, "", ""))
    monkeypatch.setattr(
        urllib.request,
        "urlopen",
        lambda *args, **kwargs: _FakeResponse(
            {
                "viewProperties": {
                    "value": {
                        "materialized": False,
                        "viewLogic": "select 1",
                        "viewLanguage": "SQL",
                    }
                }
            }
        ),
    )

    assert q.run(["target_db.target_table"], "http://gms") == 0


def test_run_groups_missing_properties_and_returns_success(monkeypatch, capsys) -> None:
    upstream_etl_missing = q.make_hive_dataset_urn("table_db.etl_missing")
    upstream_shell_missing = q.make_hive_dataset_urn("table_db.shell_missing")
    upstream_both_missing = q.make_hive_dataset_urn("table_db.both_missing")

    monkeypatch.setattr(
        q,
        "fetch_all_upstream_urns_with_counts",
        lambda *args, **kwargs: (
            {
                upstream_etl_missing,
                upstream_shell_missing,
                upstream_both_missing,
            },
            {
                upstream_etl_missing: 0,
                upstream_shell_missing: 0,
                upstream_both_missing: 0,
            },
        ),
    )

    def _read_upstream_structured_status(*args, **kwargs):
        urn = args[2]
        table_name = q.urn_to_table_name(urn)
        if table_name == "table_db.etl_missing":
            return False, True, True, "", ""
        if table_name == "table_db.shell_missing":
            return True, True, False, "", ""
        return False, True, False, "", ""

    monkeypatch.setattr(q, "read_upstream_structured_status", _read_upstream_structured_status)

    def _raise_404(*args, **kwargs):
        raise urllib.error.HTTPError(
            url="http://gms/openapi/v3/entity/dataset/urn/viewProperties",
            code=404,
            msg="Not Found",
            hdrs=None,
            fp=None,
        )

    monkeypatch.setattr(urllib.request, "urlopen", _raise_404)

    assert q.run(["target_db.target_table"], "http://gms") == 0

    output = capsys.readouterr().out
    assert "缺少: Etl Script 的如下：" in output
    assert "  table_db.etl_missing" in output
    assert "  table_db.both_missing" in output
    assert "缺少: Execute Shell 的如下：" in output
    assert "  table_db.shell_missing" in output
    assert "  table_db.both_missing" in output


def test_run_prints_data_availability_flag_for_sorted_upstream_tables(monkeypatch, capsys) -> None:
    upstream_a = q.make_hive_dataset_urn("table_db.a_upstream")
    upstream_b = q.make_hive_dataset_urn("table_db.b_upstream")

    monkeypatch.setattr(
        q,
        "fetch_all_upstream_urns_with_counts",
        lambda *args, **kwargs: ({upstream_b, upstream_a}, {upstream_a: 1, upstream_b: 0}),
    )
    monkeypatch.setattr(q, "read_dataset_type", lambda *args, **kwargs: "table")

    def _fetch_structured_properties(*args, **kwargs):
        urn = args[1]
        table_name = q.urn_to_table_name(urn)
        flag = "表血缘, DDL" if table_name == "table_db.a_upstream" else ""
        return {
            "structuredProperties": {
                "value": {
                    "properties": [
                        {
                            "propertyUrn": q.URN_DATA_AVAILABILITY_FLAG,
                            "values": [{"string": value} for value in flag.split(", ") if value],
                        }
                    ]
                }
            }
        }

    monkeypatch.setattr(q, "fetch_structured_properties", _fetch_structured_properties)

    assert q.run(["target_db.target_table"], "http://gms", skip_check_props=True) == 0

    output = capsys.readouterr().out
    first = output.index("  table_db.a_upstream\tdata_availability_flag=DDL, 表血缘")
    second = output.index("  table_db.b_upstream\tdata_availability_flag=-")
    assert first < second


def test_extract_data_availability_flag_reads_all_values() -> None:
    payload = {
        "structuredProperties": {
            "value": {
                "properties": [
                    {
                        "propertyUrn": q.URN_DATA_AVAILABILITY_FLAG,
                        "values": [{"string": "DDL"}, {"string": "表血缘"}],
                    }
                ]
            }
        }
    }

    assert q.extract_data_availability_flag(payload) == "DDL, 表血缘"


def test_extract_data_availability_flag_sorts_values() -> None:
    payload = {
        "structuredProperties": {
            "value": {
                "properties": [
                    {
                        "propertyUrn": q.URN_DATA_AVAILABILITY_FLAG,
                        "values": [{"string": "表血缘"}, {"string": "DDL"}],
                    }
                ]
            }
        }
    }

    assert q.extract_data_availability_flag(payload) == "DDL, 表血缘"


def test_extract_other_remark_reads_value() -> None:
    payload = {
        "structuredProperties": {
            "value": {
                "properties": [
                    {
                        "propertyUrn": q.URN_OTHER_REMARK,
                        "values": [{"string": "已废弃，勿使用"}],
                    }
                ]
            }
        }
    }

    assert q.extract_other_remark(payload) == "已废弃，勿使用"


def test_field_lineage_coverage_deduplicates_and_excludes_partition_fields() -> None:
    schema_payload = {
        "schemaMetadata": {
            "value": {
                "fields": [
                    {"fieldPath": "store_code"},
                    {"fieldPath": "sku_code"},
                    {
                        "fieldPath": "version",
                        "type": {"type": {"com.linkedin.schema.StringType": {}}},
                        "nativeDataType": "Partition Key",
                    },
                ]
            }
        }
    }
    upstream_lineage_payload = {
        "upstreamLineage": {
            "value": {
                "fineGrainedLineages": [
                    {
                        "downstreams": [
                            "urn:li:schemaField:(urn:li:dataset:(urn:li:dataPlatform:hive,"
                            "blf-prod-hive.default.target,PROD),store_code)"
                        ]
                    },
                    {
                        "downstreams": [
                            "urn:li:schemaField:(urn:li:dataset:(urn:li:dataPlatform:hive,"
                            "blf-prod-hive.default.target,PROD),store_code)",
                            "urn:li:schemaField:(urn:li:dataset:(urn:li:dataPlatform:hive,"
                            "blf-prod-hive.default.target,PROD),version)",
                        ]
                    },
                ]
            }
        }
    }

    coverage = q.calculate_field_lineage_coverage(
        schema_payload,
        upstream_lineage_payload,
    )

    assert coverage == {
        "field_lineage_coverage_percent": 50.0,
        "field_lineage_covered_field_count": 1,
        "ddl_non_partition_field_count": 2,
        "partition_fields": ["version"],
    }


def test_run_writes_deduplicated_upstream_detail_excel(monkeypatch, tmp_path) -> None:
    upstream_a = q.make_hive_dataset_urn("data_smartorder.dw_sku_display_snap")
    upstream_b = q.make_hive_dataset_urn("default.dim_store_info")
    output_xlsx = tmp_path / "upstreams.xlsx"

    monkeypatch.setattr(
        q,
        "fetch_all_upstream_urns_with_counts",
        lambda *args, **kwargs: (
            {upstream_a, upstream_b},
            {upstream_a: 2, upstream_b: 0},
        ),
    )
    monkeypatch.setattr(
        q,
        "read_dataset_type",
        lambda _gms_url, _token, urn: "view" if urn == upstream_b else "table",
    )
    monkeypatch.setattr(q, "is_deprecated_dataset", lambda _gms_url, _token, urn: urn == upstream_b)
    monkeypatch.setattr(q, "is_llm_generated_documentation", lambda _gms_url, _token, urn: urn == upstream_a)
    monkeypatch.setattr(
        q,
        "read_field_lineage_coverage",
        lambda _gms_url, _token, urn: {
            "field_lineage_coverage_percent": 50.0 if urn == upstream_a else 0.0,
            "field_lineage_covered_field_count": 1 if urn == upstream_a else 0,
            "ddl_non_partition_field_count": 2 if urn == upstream_a else 0,
            "partition_fields": ["dt"] if urn == upstream_a else [],
        },
    )

    def _fetch_structured_properties(*args, **kwargs):
        urn = args[1]
        table_name = q.urn_to_table_name(urn)
        values = {
            URN_ETL_SCRIPT: "select 1" if table_name == "data_smartorder.dw_sku_display_snap" else "",
            URN_SCHEDULE_URL: "https://schedule/job/x" if table_name == "data_smartorder.dw_sku_display_snap" else "",
            URN_EXECUTE_SHELL: "sh run.sh" if table_name == "data_smartorder.dw_sku_display_snap" else "",
            URN_DATA_AVAILABILITY_FLAG: "DDL, 表血缘",
            URN_OTHER_REMARK: "view 无调度" if table_name == "default.dim_store_info" else "",
        }
        return {
            "structuredProperties": {
                "value": {
                    "properties": [
                        {"propertyUrn": urn, "values": [{"string": value}]}
                        for urn, value in values.items()
                    ]
                }
            }
        }

    monkeypatch.setattr(q, "fetch_structured_properties", _fetch_structured_properties)

    assert q.run(
        ["target_db.target_table"],
        "http://gms",
        skip_check_props=True,
        output_xlsx=str(output_xlsx),
    ) == 0

    wb = load_workbook(output_xlsx)
    ws = wb["upstream_lineage"]
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0] == (
        "库名",
        "表名",
        "表名前辍",
        "完整表表",
        "表类型",
        "标记废弃",
        "Documentation生成",
        "etl_script",
        "schedule_url",
        "execute_shell",
        "data_availability_flag",
        "other_remark",
        "上游表数量",
        "字段血缘覆盖率",
        "有血缘字段数量(去重,已过滤分区字段)",
        "DDL字段数量(已过滤分区字段)",
        "分区字段",
    )
    by_full_name = {row[3]: row for row in rows[1:]}
    assert by_full_name["data_smartorder.dw_sku_display_snap"] == (
        "data_smartorder",
        "dw_sku_display_snap",
        "dw",
        "data_smartorder.dw_sku_display_snap",
        "table",
        "-",
        "是",
        "是",
        "是",
        "是",
        "DDL, 表血缘",
        "-",
        2,
        "50.00%",
        1,
        2,
        "dt",
    )
    assert by_full_name["default.dim_store_info"] == (
        "default",
        "dim_store_info",
        "dim",
        "default.dim_store_info",
        "view",
        "是",
        "-",
        "-",
        "-",
        "-",
        "DDL, 表血缘",
        "view 无调度",
        0,
        "0.00%",
        0,
        0,
        "-",
    )


def test_table_name_prefix_uses_first_underscore_segment() -> None:
    assert q.table_name_prefix("ods_logs_v1") == "ods"
    assert q.table_name_prefix("dim") == "dim"


def test_is_deprecated_dataset_reads_deprecation_aspect(monkeypatch) -> None:
    monkeypatch.setattr(
        q,
        "_fetch_dataset_aspect",
        lambda *args, **kwargs: {"deprecation": {"value": {"deprecated": True}}},
    )

    assert q.is_deprecated_dataset("http://gms", None, "urn:dataset") is True


def test_is_llm_generated_description_requires_section_four_heading() -> None:
    assert q.is_llm_generated_description("") is False
    assert q.is_llm_generated_description("## 表加工逻辑说明") is False
    assert q.is_llm_generated_description("## 4. 数据来源") is False
    assert (
        q.is_llm_generated_description(
            "人工说明\n### 4. 数据来源\n| 上游表 | 用途 |\n| --- | --- |\n| `default.ods_a` | x |"
        )
        is True
    )
    assert (
        q.is_llm_generated_description(
            "### 4\\. 数据来源\n| 上游表 | 用途 |\n| --- | --- |\n| `data_md.dm_md_dim_a` | x |"
        )
        is True
    )


def test_is_llm_generated_documentation_detects_section_four(monkeypatch) -> None:
    monkeypatch.setattr(
        q,
        "_fetch_dataset_aspect",
        lambda *args, **kwargs: {
            "editableDatasetProperties": {
                "value": {
                    "description": "### 4. 数据来源\n| `db.t` | 用途 |",
                }
            }
        },
    )

    assert q.is_llm_generated_documentation("http://gms", None, "urn:dataset") is True
