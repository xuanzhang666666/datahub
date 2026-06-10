from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook

from job_info_sync_datahub import algorithm_field_contract as module
from job_info_sync_datahub.algorithm_field_contract import (
    CONFIRMED,
    CYCLE_DETECTED,
    INCOMPLETE,
    INVALID_SCHEMA_FIELD_URN,
    MAX_DEPTH_REACHED,
    MISSING_FIELD_LINEAGE,
    MISSING_ETL_SCRIPT,
    NO_RUNTIME_VARS,
    OK,
    RISK_UNCONFIRMED_NODE,
    DatasetGraphSnapshot,
    analyze_table_payload,
    extract_dataset_job_properties,
    parse_execute_shell_variables,
    resolve_etl_script,
    table_names_from_env,
    trace_field_contract,
    load_dataset_graph_snapshot,
    write_contract_outputs,
    write_graph_contract_outputs,
)
from job_info_sync_datahub.structured_properties import (
    URN_ETL_SCRIPT,
    URN_EXECUTE_SHELL,
    URN_SCHEDULE_URL,
)


def _payload(
    *,
    etl_script: str = "",
    execute_shell: str = "",
    schedule_url: str = "",
) -> dict:
    return {
        "structuredProperties": {
            "value": {
                "properties": [
                    {"propertyUrn": URN_ETL_SCRIPT, "values": [{"string": etl_script}]},
                    {"propertyUrn": URN_EXECUTE_SHELL, "values": [{"string": execute_shell}]},
                    {"propertyUrn": URN_SCHEDULE_URL, "values": [{"string": schedule_url}]},
                ]
            }
        }
    }


def test_extract_dataset_job_properties_reads_three_structured_properties() -> None:
    props = extract_dataset_job_properties(
        _payload(
            etl_script="```sql\nselect 1\n```",
            execute_shell="```shell\nDATE=20260603 sh run.sh\n```",
            schedule_url="https://schedule/job/a",
        )
    )

    assert props.etl_script == "select 1"
    assert props.execute_shell == "DATE=20260603 sh run.sh"
    assert props.schedule_url == "https://schedule/job/a"


def test_parse_execute_shell_variables_supports_assignments_exports_flags_and_positionals() -> None:
    ctx = parse_execute_shell_variables(
        "DATE=20260603 export DATABASE=default /bin/w-run-task.sh foo/bar "
        "--hive_table_name=dw.target --bizdate 20260602 prod"
    )

    assert ctx.variables["DATE"] == "20260603"
    assert ctx.variables["DATABASE"] == "default"
    assert ctx.variables["hive_table_name"] == "dw.target"
    assert ctx.variables["bizdate"] == "20260602"
    assert ctx.positionals[-1] == "prod"
    assert ctx.variables["DATE_SUB1DAY"] == "20260602"
    assert ctx.variables["FDATE_SUB1DAY"] == "2026-06-02"


def test_resolve_etl_script_uses_execute_shell_variables_and_reports_unresolved() -> None:
    runtime = parse_execute_shell_variables(
        "DATE=20260603 --hive_table_name=dw.target"
    )

    resolved = resolve_etl_script(
        "insert overwrite table ${hive_table_name} "
        "select * from ods.source where dt='${DATE_SUB1DAY}' and shop=$SHOP;"
        "select '{{ DATE }}';",
        runtime,
    )

    assert "dw.target" in resolved.resolved_script
    assert "20260602" in resolved.resolved_script
    assert "20260603" in resolved.resolved_script
    assert "$SHOP" in resolved.resolved_script
    assert "SHOP" in resolved.unresolved_variables


def test_analyze_table_payload_marks_missing_etl_and_no_runtime_vars() -> None:
    missing = analyze_table_payload("dw.target", "urn:dataset", _payload())
    assert missing.status == MISSING_ETL_SCRIPT

    no_runtime = analyze_table_payload(
        "dw.target",
        "urn:dataset",
        _payload(
            etl_script=(
                "$HIVE <<EOF\n"
                "insert overwrite table dw.target select id from ods.source;\n"
                "EOF"
            )
        ),
    )
    assert no_runtime.status == NO_RUNTIME_VARS
    assert no_runtime.table_lineages[0].upstreams[0].full_name == "ods.source"


def test_analyze_table_payload_parses_resolved_etl_field_lineage() -> None:
    result = analyze_table_payload(
        "dw.target",
        "urn:dataset",
        _payload(
            schedule_url="https://schedule/job/target",
            execute_shell="DATE=20260603 TARGET_TABLE=dw.target",
            etl_script=(
                "$HIVE <<EOF\n"
                "insert overwrite table ${TARGET_TABLE} "
                "select s.id as sku_id, s.qty from ods.source s "
                "where s.dt='${DATE_SUB1DAY}';\n"
                "EOF"
            ),
        ),
    )

    assert result.status == OK
    assert result.unresolved_variables == []
    assert result.resolved_etl_script is not None
    assert "20260602" in result.resolved_etl_script
    assert result.table_lineages[0].target.full_name == "dw.target"
    assert result.table_lineages[0].upstreams[0].full_name == "ods.source"
    by_field = {lineage.target_field: lineage for lineage in result.field_lineages}
    assert by_field["sku_id"].mappings[0].source_field == "id"
    assert by_field["qty"].mappings[0].source_field == "qty"


def test_analyze_table_payload_marks_target_fields_without_lineage_as_open_questions() -> None:
    result = analyze_table_payload(
        "dw.target",
        "urn:dataset",
        _payload(
            execute_shell="TARGET_TABLE=dw.target",
            etl_script=(
                "$HIVE <<EOF\n"
                "insert overwrite table ${TARGET_TABLE} select id from ods.source;\n"
                "EOF"
            ),
        ),
        target_fields=["id", "missing_field"],
    )

    assert "missing_field" in result.unresolved_target_fields
    assert any("missing_field" in question for question in result.open_questions)


def test_write_contract_outputs_creates_json_markdown_and_excel(tmp_path: Path) -> None:
    result = analyze_table_payload(
        "dw.target",
        "urn:dataset",
        _payload(
            schedule_url="https://schedule/job/target",
            execute_shell="TARGET_TABLE=dw.target",
            etl_script="$HIVE <<EOF\ninsert overwrite table ${TARGET_TABLE} select id from ods.source;\nEOF",
        ),
    )

    write_contract_outputs([result], tmp_path)

    assert (tmp_path / "field_lineage.json").exists()
    assert (tmp_path / "runtime_context.json").exists()
    assert (tmp_path / "lineage_report.md").exists()
    assert (tmp_path / "open_questions.md").exists()
    assert (tmp_path / "field_contract.xlsx").exists()

    payload = json.loads((tmp_path / "field_lineage.json").read_text())
    assert payload["tables"][0]["table_name"] == "dw.target"
    assert payload["tables"][0]["field_lineages"][0]["target_field"] == "id"


def test_table_names_from_env_accepts_jenkins_multiline_and_comma_separated(monkeypatch) -> None:
    monkeypatch.setenv(
        "TABLE_NAMES",
        "default.table_a\n table_b,ods.table_c，dim.table_d \n# comment\n",
    )

    assert table_names_from_env() == [
        "default.table_a",
        "table_b",
        "ods.table_c",
        "dim.table_d",
    ]


def _snapshot(
    table: str,
    *,
    fields: list[str],
    partitions: list[str] | None = None,
    table_upstreams: list[str] | None = None,
    field_edges: dict[str, list[tuple[str, str]]] | None = None,
    confirmed: bool = True,
) -> DatasetGraphSnapshot:
    return DatasetGraphSnapshot(
        table_name=table,
        dataset_urn=f"urn:dataset:{table}",
        schema_fields=fields,
        partition_fields=partitions or [],
        field_descriptions={field: f"{field} description" for field in fields},
        table_upstreams=table_upstreams or [],
        field_edges=field_edges or {},
        field_lineage_confirmed=confirmed,
        raw_aspects={},
    )


def test_trace_field_contract_uses_all_non_partition_fields_and_recurses() -> None:
    snapshots = {
        "dw.target": _snapshot(
            "dw.target",
            fields=["sku_id", "qty", "dt"],
            partitions=["dt"],
            table_upstreams=["dim.sku", "ods.qty"],
            field_edges={
                "sku_id": [("dim.sku.id", "cast(id as string)")],
                "qty": [("ods.qty.qty", "coalesce(qty, 0)")],
            },
        ),
        "dim.sku": _snapshot(
            "dim.sku",
            fields=["id"],
            table_upstreams=["ods.sku"],
            field_edges={"id": [("ods.sku.id", "id")]},
        ),
        "ods.sku": _snapshot("ods.sku", fields=["id"]),
        "ods.qty": _snapshot("ods.qty", fields=["qty"]),
    }

    result = trace_field_contract(
        ["dw.target"],
        snapshot_loader=lambda table: snapshots[table],
        max_depth=20,
    )

    assert {field.target_field for field in result.fields} == {"sku_id", "qty"}
    sku_path = next(path for path in result.paths if path.target_field == "sku_id")
    assert sku_path.nodes == ["dw.target.sku_id", "dim.sku.id", "ods.sku.id"]
    assert sku_path.transform_operations == ["cast(id as string)", "id"]
    assert sku_path.final_source == "ods.sku.id"
    assert sku_path.path_status == "COMPLETE"
    assert sku_path.trust_status == CONFIRMED


def test_trace_field_contract_splits_multi_source_and_marks_unconfirmed_risk() -> None:
    snapshots = {
        "dw.target": _snapshot(
            "dw.target",
            fields=["name"],
            table_upstreams=["dim.a", "dim.b"],
            field_edges={
                "name": [
                    ("dim.a.name", "coalesce(a.name, b.name)"),
                    ("dim.b.name", "coalesce(a.name, b.name)"),
                ]
            },
        ),
        "dim.a": _snapshot("dim.a", fields=["name"], confirmed=False),
        "dim.b": _snapshot("dim.b", fields=["name"]),
    }

    result = trace_field_contract(
        ["dw.target"],
        snapshot_loader=lambda table: snapshots[table],
        max_depth=20,
    )

    assert len(result.paths) == 2
    risky = next(path for path in result.paths if path.final_source == "dim.a.name")
    assert risky.trust_status == RISK_UNCONFIRMED_NODE
    assert risky.unconfirmed_nodes == ["dim.a.name"]
    assert any("dim.a.name" in question for question in result.open_questions)


def test_trace_field_contract_reports_missing_cycle_and_max_depth() -> None:
    missing_snapshots = {
        "dw.target": _snapshot(
            "dw.target",
            fields=["id"],
            table_upstreams=["ods.source"],
        )
    }
    missing = trace_field_contract(
        ["dw.target"],
        snapshot_loader=lambda table: missing_snapshots[table],
        max_depth=20,
    )
    assert missing.paths[0].path_status == MISSING_FIELD_LINEAGE
    assert missing.paths[0].trust_status == INCOMPLETE

    cycle_snapshots = {
        "dw.a": _snapshot(
            "dw.a",
            fields=["id"],
            table_upstreams=["dw.b"],
            field_edges={"id": [("dw.b.id", "id")]},
        ),
        "dw.b": _snapshot(
            "dw.b",
            fields=["id"],
            table_upstreams=["dw.a"],
            field_edges={"id": [("dw.a.id", "id")]},
        ),
    }
    cycle = trace_field_contract(
        ["dw.a"],
        snapshot_loader=lambda table: cycle_snapshots[table],
        max_depth=20,
    )
    assert cycle.paths[0].path_status == CYCLE_DETECTED

    depth = trace_field_contract(
        ["dw.a"],
        snapshot_loader=lambda table: cycle_snapshots[table],
        max_depth=1,
    )
    assert depth.paths[0].path_status == MAX_DEPTH_REACHED


def test_trace_field_contract_reports_invalid_field_node_and_missing_dataset() -> None:
    invalid_snapshots = {
        "dw.target": _snapshot(
            "dw.target",
            fields=["id"],
            table_upstreams=["ods.source"],
            field_edges={"id": [("__invalid__:bad-urn", "")]},
        )
    }
    invalid = trace_field_contract(
        ["dw.target"],
        snapshot_loader=lambda table: invalid_snapshots[table],
        max_depth=20,
    )
    assert invalid.paths[0].path_status == INVALID_SCHEMA_FIELD_URN

    unknown_field = trace_field_contract(
        ["dw.target"],
        snapshot_loader=lambda table: (
            _snapshot(
                "dw.target",
                fields=["id"],
                table_upstreams=["ods.source"],
                field_edges={"id": [("ods.source.unknown", "unknown")]},
            )
            if table == "dw.target"
            else _snapshot("ods.source", fields=["id"])
        ),
        max_depth=20,
    )
    assert unknown_field.paths[0].path_status == INVALID_SCHEMA_FIELD_URN

    missing = trace_field_contract(
        ["dw.target"],
        snapshot_loader=lambda table: (
            _snapshot(
                "dw.target",
                fields=["id"],
                table_upstreams=["ods.missing"],
                field_edges={"id": [("ods.missing.id", "id")]},
            )
            if table == "dw.target"
            else (_ for _ in ()).throw(KeyError(table))
        ),
        max_depth=20,
    )
    assert missing.paths[0].path_status == "DATASET_NOT_FOUND"


def test_load_dataset_graph_snapshot_parses_datahub_aspects(monkeypatch) -> None:
    target_urn = (
        "urn:li:dataset:(urn:li:dataPlatform:hive,"
        "blf-prod-hive.dw.target,PROD)"
    )
    source_urn = (
        "urn:li:dataset:(urn:li:dataPlatform:hive,"
        "blf-prod-hive.ods.source,PROD)"
    )
    payloads = {
        "schemaMetadata": {
            "schemaMetadata": {
                "value": {
                    "fields": [
                        {"fieldPath": "id", "description": "identifier"},
                        {
                            "fieldPath": "dt",
                            "nativeDataType": "Partition Key",
                        },
                    ]
                }
            }
        },
        "upstreamLineage": {
            "upstreamLineage": {
                "value": {
                    "upstreams": [{"dataset": source_urn}],
                    "fineGrainedLineages": [
                        {
                            "upstreams": [f"urn:li:schemaField:({source_urn},id)"],
                            "downstreams": [f"urn:li:schemaField:({target_urn},id)"],
                            "transformOperation": "cast(id as bigint)",
                        }
                    ],
                }
            }
        },
        "structuredProperties": {
            "structuredProperties": {
                "value": {
                    "properties": [
                        {
                            "propertyUrn": (
                                "urn:li:structuredProperty:"
                                "blf.data.warehouse.data_availability_flag"
                            ),
                            "values": [{"string": "字段血缘"}],
                        }
                    ]
                }
            }
        },
    }
    monkeypatch.setattr(
        "job_info_sync_datahub.algorithm_field_contract._fetch_openapi_aspect",
        lambda _gms, _urn, aspect, _token, required=False: payloads[aspect],
    )

    snapshot = load_dataset_graph_snapshot(
        "dw.target",
        gms_url="http://localhost:8080",
        token=None,
        platform_instance="blf-prod-hive",
        env="PROD",
    )

    assert snapshot.schema_fields == ["id", "dt"]
    assert snapshot.partition_fields == ["dt"]
    assert snapshot.field_descriptions == {"id": "identifier", "dt": ""}
    assert snapshot.table_upstreams == ["ods.source"]
    assert snapshot.field_edges == {"id": [("ods.source.id", "cast(id as bigint)")]}
    assert snapshot.field_lineage_confirmed is True


def test_write_graph_contract_outputs_creates_required_sheets_and_raw_aspects(
    tmp_path: Path,
) -> None:
    snapshot = _snapshot("ods.source", fields=["id"])
    snapshot.raw_aspects = {
        "schemaMetadata": {"schemaMetadata": {"value": {"fields": [{"fieldPath": "id"}]}}},
        "upstreamLineage": {"upstreamLineage": {"value": {}}},
        "structuredProperties": {"structuredProperties": {"value": {"properties": []}}},
    }
    result = trace_field_contract(
        ["ods.source"],
        snapshot_loader=lambda _table: snapshot,
        max_depth=20,
    )

    write_graph_contract_outputs(result, tmp_path)

    workbook = load_workbook(tmp_path / "field_contract.xlsx", read_only=True)
    assert workbook.sheetnames == [
        "field_contract",
        "trace_paths",
        "open_questions",
        "run_summary",
    ]
    assert (tmp_path / "field_lineage.json").exists()
    assert (tmp_path / "lineage_report.md").exists()
    assert (tmp_path / "open_questions.md").exists()
    assert (tmp_path / "raw_aspects" / "ods.source" / "schemaMetadata.json").exists()


def test_main_uses_datahub_graph_path_without_legacy_etl_analysis(
    monkeypatch,
    tmp_path: Path,
) -> None:
    result = trace_field_contract(
        ["ods.source"],
        snapshot_loader=lambda _table: _snapshot("ods.source", fields=["id"]),
        max_depth=20,
    )
    captured: dict[str, object] = {}
    monkeypatch.setattr(
        module,
        "analyze_tables",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(
            AssertionError("legacy ETL analysis must not run")
        ),
    )
    monkeypatch.setattr(module, "trace_field_contract", lambda *_args, **_kwargs: result)
    monkeypatch.setattr(
        module,
        "write_graph_contract_outputs",
        lambda graph_result, output_dir: captured.update(
            {"result": graph_result, "output_dir": output_dir}
        ),
    )

    exit_code = module.main(
        [
            "--table",
            "ods.source",
            "--output-dir",
            str(tmp_path),
        ]
    )

    assert exit_code == 0
    assert captured == {"result": result, "output_dir": str(tmp_path)}
