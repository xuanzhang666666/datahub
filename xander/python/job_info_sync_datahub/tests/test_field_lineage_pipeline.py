from __future__ import annotations

import importlib.util
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from job_info_sync_datahub.field_lineage_datahub_reader import (
    build_target_table_aliases,
    extract_data_availability_flags,
    extract_field_lineage_input,
    extract_field_lineage_input_with_debug,
    extract_schema_field_names,
    extract_schema_partition_field_names,
    extract_target_partition_fields_from_script,
    has_confirmed_field_lineage,
    make_hive_dataset_urn,
    missing_field_lineage_source_reason,
    strip_commented_sql_for_llm,
    strip_markdown_code_fence,
    write_field_lineage_debug_artifacts,
)
from job_info_sync_datahub.field_lineage_batch_summary import write_batch_summary_workbook
from job_info_sync_datahub.field_lineage_cli import main as field_lineage_cli_main
from job_info_sync_datahub.field_lineage_constants import (
    is_runtime_date_variable_expression,
)
from job_info_sync_datahub.field_lineage_excel import (
    load_approved_review_rows,
    write_candidate_workbook,
)
from job_info_sync_datahub.field_lineage_llm import (
    FIELD_LINEAGE_SYSTEM_PROMPT,
    build_field_lineage_user_message,
    build_field_lineage_request_debug_info,
    parse_field_lineage_payload,
)
from job_info_sync_datahub.field_lineage_models import (
    FieldLineageCandidate,
    FieldLineageInput,
    FieldLineageReviewStatus,
    UnresolvedField,
)
from job_info_sync_datahub.field_lineage_writer import (
    _field_name_from_schema_field_urn,
    _read_fine_grained_lineages_with_retry,
    build_fine_grained_lineage_class,
    build_transform_operation_for_ui,
    group_approved_rows,
    merge_fine_grained_lineages,
    verify_field_lineage_completeness,
)
from job_info_sync_datahub.structured_properties import URN_ETL_SCRIPT, URN_EXECUTE_SHELL
from job_info_sync_datahub.structured_properties import URN_DATA_AVAILABILITY_FLAG


def _structured_properties_payload(
    etl_script: str,
    execute_shell: str,
    availability_flags: list[str] | None = None,
) -> dict:
    properties = [
        {
            "propertyUrn": URN_ETL_SCRIPT,
            "values": [{"string": etl_script}],
        },
        {
            "propertyUrn": URN_EXECUTE_SHELL,
            "values": [{"string": execute_shell}],
        },
    ]
    if availability_flags is not None:
        properties.append(
            {
                "propertyUrn": URN_DATA_AVAILABILITY_FLAG,
                "values": [{"string": flag} for flag in availability_flags],
            }
        )
    return {
        "structuredProperties": {
            "value": {
                "properties": properties
            }
        }
    }


def test_make_hive_dataset_urn_uses_blf_defaults() -> None:
    assert make_hive_dataset_urn("default.dim_store_info") == (
        "urn:li:dataset:(urn:li:dataPlatform:hive,"
        "blf-prod-hive.default.dim_store_info,PROD)"
    )


def test_strip_markdown_code_fence_preserves_inner_script() -> None:
    assert strip_markdown_code_fence("```sql\nselect 1;\n```") == "select 1;"
    assert strip_markdown_code_fence("select 1;") == "select 1;"


def test_extract_field_lineage_input_from_structured_properties() -> None:
    payload = _structured_properties_payload(
        etl_script="```sql\ninsert overwrite table default.dim_store_info select 1;\n```",
        execute_shell="```shell\nsh run_dim_store_info.sh\n```",
    )

    result = extract_field_lineage_input(
        dataset_urn=make_hive_dataset_urn("default.dim_store_info"),
        table_name="default.dim_store_info",
        payload=payload,
    )

    assert result == FieldLineageInput(
        dataset_urn=make_hive_dataset_urn("default.dim_store_info"),
        table_name="default.dim_store_info",
        etl_script="insert overwrite table default.dim_store_info select 1;",
        execute_shell="sh run_dim_store_info.sh",
        target_table_aliases=["default.not_verified_dim_store_info"],
    )


def test_extract_field_lineage_input_resolves_execute_shell_variables() -> None:
    payload = _structured_properties_payload(
        etl_script=(
            "```sql\n"
            "insert overwrite table default.${TABLE_NAME} partition(dt='${DATE}') "
            "select * from default.${SOURCE_TABLE} where dt='${DATE_SUB1DAY}';\n"
            "```"
        ),
        execute_shell=(
            "```shell\n"
            "export TABLE_NAME=dim_store_info\n"
            "SOURCE_TABLE=mid_store_info_bach\n"
            "sh run.sh --DATE 20260508\n"
            "```"
        ),
    )

    result = extract_field_lineage_input(
        dataset_urn=make_hive_dataset_urn("default.dim_store_info"),
        table_name="default.dim_store_info",
        payload=payload,
    )

    assert "default.dim_store_info" in result.etl_script
    assert "default.mid_store_info_bach" in result.etl_script
    assert "dt='20260508'" in result.etl_script
    assert "dt='20260507'" in result.etl_script
    assert "${" not in result.etl_script


def test_extract_field_lineage_input_resolves_nested_not_verified_table_name() -> None:
    payload = _structured_properties_payload(
        etl_script=(
            "```shell\n"
            "source ${ETC}/format_date.cnf\n"
            "NOT_VERIFIED_TABLE_NAME=\"${TABLE_NAME}\"\n"
            "TABLE_NAME=\"dw_store_order_plan_sku\"\n"
            "function dw_store_order_plan_sku_run {\n"
            "  calculate\n"
            "}\n"
            "function calculate {\n"
            "  ${HIVE} -e << EOF \"\n"
            "    insert overwrite table ${NOT_VERIFIED_TABLE_NAME} partition(dt=${DATE})\n"
            "    select store_code, sku_code from dw_store_order_plan_sku_di where dt='${DATE}';\n"
            "  \"\n"
            "EOF\n"
            "}\n"
            "```"
        ),
        execute_shell="```shell\nDATE=20260603 sh run.sh\n```",
    )

    result = extract_field_lineage_input(
        dataset_urn=make_hive_dataset_urn("default.dw_store_order_plan_sku"),
        table_name="default.dw_store_order_plan_sku",
        payload=payload,
    )

    assert "insert overwrite table dw_store_order_plan_sku" in result.etl_script
    assert "partition(dt=20260603)" in result.etl_script
    assert "${NOT_VERIFIED_TABLE_NAME}" not in result.etl_script


def test_extract_data_availability_flags_detects_confirmed_field_lineage() -> None:
    payload = _structured_properties_payload(
        etl_script="select 1",
        execute_shell="sh run.sh",
        availability_flags=["DDL", "表血缘", "字段血缘"],
    )

    assert extract_data_availability_flags(payload) == ["DDL", "表血缘", "字段血缘"]
    assert has_confirmed_field_lineage(payload) is True


def test_extract_schema_partition_field_names_from_partition_key_type() -> None:
    payload = {
        "schemaMetadata": {
            "value": {
                "fields": [
                    {"fieldPath": "sku_code", "nativeDataType": "string"},
                    {"fieldPath": "kpt", "nativeDataType": "Partition Key"},
                    {"fieldPath": "biz_hour", "type": "Partition Key"},
                    {"fieldPath": "[version=2.0].[type=string].version", "isPartitioningKey": True},
                ]
            }
        }
    }

    assert extract_schema_field_names(payload) == ["sku_code", "kpt", "biz_hour", "version"]
    assert extract_schema_partition_field_names(payload) == ["kpt", "biz_hour", "version"]


def test_extract_target_partition_fields_from_script_for_target_write() -> None:
    script = """
    insert overwrite table data_build.dm_site_selection_project_location_type_sign_v1
    partition (dt='20260608', version='1.1.0')
    select project_id, location_type, store_code from tmp;
    analyze table data_build.dm_site_selection_project_location_type_sign_v1
    partition (dt='20260608') compute statistics;
    """

    assert extract_target_partition_fields_from_script(
        script,
        "data_build.dm_site_selection_project_location_type_sign_v1",
    ) == ["dt", "version"]


def test_extract_schema_field_names_ignores_nested_struct_subfields() -> None:
    payload = {
        "schemaMetadata": {
            "value": {
                "fields": [
                    {
                        "fieldPath": "[version=2.0].[type=string].document_id",
                        "nativeDataType": "string",
                    },
                    {
                        "fieldPath": "[version=2.0].[type=struct].sale_spec",
                        "nativeDataType": "struct",
                    },
                    {
                        "fieldPath": "[version=2.0].[type=struct].sale_spec.expression",
                        "nativeDataType": "string",
                    },
                    {
                        "fieldPath": "[version=2.0].[type=struct].sale_spec.unit",
                        "nativeDataType": "string",
                    },
                    {
                        "fieldPath": "[version=2.0].[type=struct].sale_spec.qty",
                        "nativeDataType": "string",
                    },
                    {
                        "fieldPath": "[version=2.0].[type=struct].ordering_spec",
                        "nativeDataType": "struct",
                    },
                ]
            }
        }
    }

    assert extract_schema_field_names(payload) == [
        "document_id",
        "sale_spec",
        "ordering_spec",
    ]


def test_strip_commented_sql_for_llm_removes_commented_sources() -> None:
    script = """
with active_data as (
    select store_code, sku_code -- active inline comment
    from data_smartorder.active_source_di
),
--store_contract_hurdle_data as--项目维度，合同
--(select
    --project_id,
    --business_estimate_value_hurdle-0 as hurdle--有数据只有600多条
--from dw_store_construction_contract_detail_v1--上游209单
--where dt ='20260603'
--),
/* commented_block as (
   select * from data_smartorder.block_commented_source_di
) */
final_data as (
    select * from active_data
)
"""

    cleaned = strip_commented_sql_for_llm(script)

    assert "data_smartorder.active_source_di" in cleaned
    assert "active inline comment" not in cleaned
    assert "dw_store_construction_contract_detail_v1" not in cleaned
    assert "data_smartorder.block_commented_source_di" not in cleaned


def test_extract_field_lineage_input_keeps_only_job_entry_reachable_functions() -> None:
    payload = _structured_properties_payload(
        etl_script=(
            "```sql\n"
            "TABLE_NAME=ods_bach_baseinfo_shop_company\n"
            "HDFS_DIR=/user/wstats/$TABLE_NAME\n"
            "\n"
            "function unused_debug_sql {\n"
            "  $HIVE -e \"select * from default.should_not_send\"\n"
            "}\n"
            "\n"
            "function calculate {\n"
            "  $HIVE << EOF\n"
            "insert overwrite table default.$TABLE_NAME select company_id from ods.company;\n"
            "--select bad_id from ods.commented_source;\n"
            "EOF\n"
            "}\n"
            "\n"
            "function do_check {\n"
            "  echo check\n"
            "}\n"
            "\n"
            "function ods_bach_baseinfo_shop_company_run {\n"
            "  rebuild_hdfs_dir ${HDFS_DIR} && calculate && do_check\n"
            "}\n"
            "```"
        ),
        execute_shell="```shell\n/home/w/analysis-jobs/bin/w-run-task.sh ods_bach_baseinfo_shop_company\n```",
    )

    result = extract_field_lineage_input(
        dataset_urn=make_hive_dataset_urn("default.ods_bach_baseinfo_shop_company"),
        table_name="default.ods_bach_baseinfo_shop_company",
        payload=payload,
    )

    assert "function ods_bach_baseinfo_shop_company_run" in result.etl_script
    assert "function calculate" in result.etl_script
    assert "function do_check" in result.etl_script
    assert "HDFS_DIR=/user/wstats/ods_bach_baseinfo_shop_company" in result.etl_script
    assert "unused_debug_sql" not in result.etl_script
    assert "should_not_send" not in result.etl_script
    assert "ods.commented_source" not in result.etl_script


def test_write_field_lineage_debug_artifacts_saves_intermediate_results(tmp_path: Path) -> None:
    payload = _structured_properties_payload(
        etl_script=(
            "```sql\n"
            "TABLE_NAME=ods_bach_baseinfo_shop_company\n"
            "UNIQ_KEY=company_id\n"
            "function calculate {\n"
            "  $HIVE -e \"select count(distinct ${UNIQ_KEY}) from ${TABLE_NAME} where dt='$DATE'\"\n"
            "}\n"
            "function unused_sql {\n"
            "  $HIVE -e \"select * from should_not_send\"\n"
            "}\n"
            "function ods_bach_baseinfo_shop_company_run {\n"
            "  calculate\n"
            "}\n"
            "```"
        ),
        execute_shell="```shell\nsh run.sh --DATE 20260508\n```",
    )
    source_input, debug = extract_field_lineage_input_with_debug(
        dataset_urn=make_hive_dataset_urn("default.ods_bach_baseinfo_shop_company"),
        table_name="default.ods_bach_baseinfo_shop_company",
        payload=payload,
    )

    written = write_field_lineage_debug_artifacts(tmp_path, source_input, debug)

    assert sorted(written) == [
        "execute_shell.txt",
        "llm_input_etl_script.sql",
        "original_etl_script.sql",
        "processing_summary.json",
        "resolved_etl_script.sql",
        "runtime_variables.json",
    ]
    assert "should_not_send" in (tmp_path / "resolved_etl_script.sql").read_text(
        encoding="utf-8"
    )
    llm_input = (tmp_path / "llm_input_etl_script.sql").read_text(encoding="utf-8")
    assert "count(distinct company_id)" in llm_input
    assert "from ods_bach_baseinfo_shop_company" in llm_input
    assert "should_not_send" not in llm_input
    summary = json.loads((tmp_path / "processing_summary.json").read_text(encoding="utf-8"))
    assert summary["entry_function"] == "ods_bach_baseinfo_shop_company_run"
    assert summary["reachable_functions"] == [
        "calculate",
        "ods_bach_baseinfo_shop_company_run",
    ]
    assert summary["removed_functions"] == ["unused_sql"]
    assert summary["unresolved_variables"] == []


def test_extract_field_lineage_input_keeps_python_script_complete() -> None:
    python_script = (
        "```python\n"
        "import os\n"
        "TABLE_NAME = 'dim_store_info'\n"
        "\n"
        "def build_sql():\n"
        "    return f\"insert overwrite table {TABLE_NAME} select * from ods.store\"\n"
        "\n"
        "def main():\n"
        "    spark.sql(build_sql())\n"
        "\n"
        "if __name__ == '__main__':\n"
        "    main()\n"
        "```"
    )
    payload = _structured_properties_payload(
        etl_script=python_script,
        execute_shell="```shell\npython dim_store_info.py\n```",
    )

    result = extract_field_lineage_input(
        dataset_urn=make_hive_dataset_urn("default.dim_store_info"),
        table_name="default.dim_store_info",
        payload=payload,
    )

    assert "import os" in result.etl_script
    assert "TABLE_NAME = 'dim_store_info'" in result.etl_script
    assert "def build_sql" in result.etl_script
    assert "def main" in result.etl_script
    assert "if __name__ == '__main__'" in result.etl_script


def test_missing_field_lineage_source_reason_when_etl_empty() -> None:
    payload = _structured_properties_payload(etl_script="", execute_shell="")
    reason = missing_field_lineage_source_reason(payload)
    assert reason is not None
    assert "Etl Script" in reason
    assert "Execute Shell" in reason


def test_missing_field_lineage_source_reason_none_when_etl_present() -> None:
    payload = _structured_properties_payload(
        etl_script="select 1",
        execute_shell="",
    )
    assert missing_field_lineage_source_reason(payload) is None


def test_parse_field_lineage_payload_high_confidence_is_auto_approved() -> None:
    payload = json.dumps(
        {
            "target_table": "default.dim_store_info",
            "mappings": [
                {
                    "target_field": "store_id",
                    "source_table": "ods.store_info",
                    "source_field": "id",
                    "transform_expression": "cast(id as bigint)",
                    "transform_explanation": "取 ods.store_info 表中的 id 字段，表示将门店 ID 转为 bigint 类型写入 store_id。",
                    "evidence_sql": "select cast(id as bigint) as store_id",
                    "confidence": "HIGH",
                    "notes": "direct mapping",
                }
            ],
            "unresolved_fields": [{"target_field": "store_name", "reason": "dynamic SQL"}],
        }
    )

    parsed = parse_field_lineage_payload(payload)

    assert parsed.target_table == "default.dim_store_info"
    assert parsed.mappings == [
        FieldLineageCandidate(
            target_table="default.dim_store_info",
            target_field="store_id",
            source_table="ods.store_info",
            source_field="id",
            transform_expression="cast(id as bigint)",
            transform_explanation="取 ods.store_info 表中的 id 字段，表示将门店 ID 转为 bigint 类型写入 store_id。",
            evidence_sql="select cast(id as bigint) as store_id",
            confidence="HIGH",
            llm_notes="direct mapping",
            review_status=FieldLineageReviewStatus.AUTO_APPROVED,
        )
    ]
    assert parsed.unresolved_fields[0].target_field == "store_name"


def test_parse_field_lineage_payload_non_high_confidence_stays_pending() -> None:
    payload = json.dumps(
        {
            "target_table": "default.dim_store_info",
            "mappings": [
                {
                    "target_field": "store_id",
                    "source_table": "ods.store_info",
                    "source_field": "id",
                    "confidence": "MEDIUM",
                }
            ],
            "unresolved_fields": [],
        }
    )

    parsed = parse_field_lineage_payload(payload)

    assert parsed.mappings[0].review_status == FieldLineageReviewStatus.PENDING


def test_field_lineage_prompt_requires_complex_map_field_mappings() -> None:
    assert "map/struct/json" in FIELD_LINEAGE_SYSTEM_PROMPT
    assert "map key 字符串常量不算来源字段" in FIELD_LINEAGE_SYSTEM_PROMPT
    assert "复杂类型本身不是 unresolved 的理由" in FIELD_LINEAGE_SYSTEM_PROMPT


def test_field_lineage_prompt_describes_auto_approved_and_constant_fields() -> None:
    assert "confidence 仅作为参考" in FIELD_LINEAGE_SYSTEM_PROMPT
    assert "source_table 和 source_field 留空" in FIELD_LINEAGE_SYSTEM_PROMPT
    assert "evidence_sql 必须回指到输入 ETL 脚本中的 SQL 片段" in FIELD_LINEAGE_SYSTEM_PROMPT
    assert "未写库名时，必须按 `default.<表名>` 输出 source_table" in FIELD_LINEAGE_SYSTEM_PROMPT
    assert "禁止在 transform_expression、transform_explanation、evidence_sql 中输出“同上”" in FIELD_LINEAGE_SYSTEM_PROMPT
    assert "调度运行时日期变量也属于常量" in FIELD_LINEAGE_SYSTEM_PROMPT
    assert "CSV、本地文件、手工维护文件、API 拉取" in FIELD_LINEAGE_SYSTEM_PROMPT


def test_runtime_date_variable_expressions_are_constants() -> None:
    assert is_runtime_date_variable_expression("${DATE}")
    assert is_runtime_date_variable_expression("$FORMAT_DATE")
    assert is_runtime_date_variable_expression("'${DATE_SUB7DAY}'")
    assert is_runtime_date_variable_expression("cast(${FDATE_ADD1MONTH} as string)")
    assert is_runtime_date_variable_expression("${MONTH_SUB1MONTH_FIRSTDAY}")
    assert is_runtime_date_variable_expression("{{ LAST_YEAR_MONTH }}")
    assert not is_runtime_date_variable_expression("${SOURCE_ID}")


def test_build_field_lineage_request_debug_info_counts_prompt_size() -> None:
    source = FieldLineageInput(
        dataset_urn=make_hive_dataset_urn("default.dim_store_info"),
        table_name="default.dim_store_info",
        etl_script="select store_id from ods.store_info",
        execute_shell="sh dim_store_info",
    )

    debug = build_field_lineage_request_debug_info(
        source_input=source,
        base_v1="https://api.deepseek.com/v1",
        model="deepseek-chat",
        timeout_sec=90,
    )

    assert debug["llm_base"] == "https://api.deepseek.com/v1"
    assert debug["llm_model"] == "deepseek-chat"
    assert debug["timeout_sec"] == 90
    assert debug["user_message_chars"] > len(source.etl_script)
    assert debug["system_prompt_chars"] > 0


def test_build_field_lineage_user_message_includes_partition_fields() -> None:
    source = FieldLineageInput(
        dataset_urn=make_hive_dataset_urn("default.dim_store_info"),
        table_name="default.dim_store_info",
        etl_script="select store_id from ods.store_info",
        execute_shell="sh dim_store_info",
        target_schema_fields=["store_id", "version"],
        target_partition_fields=["version"],
    )

    message = build_field_lineage_user_message(source)

    assert "目标表分区字段（不需要输出字段血缘）:" in message
    assert "- version" in message


def test_write_candidate_workbook_creates_review_sheets(tmp_path: Path) -> None:
    output = tmp_path / "default_dim_store_info_field_lineage.xlsx"
    source = FieldLineageInput(
        dataset_urn=make_hive_dataset_urn("default.dim_store_info"),
        table_name="default.dim_store_info",
        etl_script="insert overwrite table default.dim_store_info select 1;",
        execute_shell="sh run_dim_store_info.sh",
    )
    candidate = FieldLineageCandidate(
        target_table="default.dim_store_info",
        target_field="store_id",
        source_table="ods.store_info",
        source_field="id",
        transform_expression="cast(id as bigint)",
        transform_explanation="取 ods.store_info 表中的 id 字段，表示将门店 ID 转为 bigint 类型写入 store_id。",
        evidence_sql="select cast(id as bigint) as store_id",
        confidence="HIGH",
    )

    write_candidate_workbook(
        output,
        source_input=source,
        candidates=[candidate],
        unresolved_fields=[],
        llm_model="deepseek-test",
        debug_dir=tmp_path / "debug",
    )

    wb = load_workbook(output)
    assert wb.sheetnames == ["candidate_lineage", "unresolved_fields", "source_context"]
    ws = wb["candidate_lineage"]
    headers = [cell.value for cell in ws[1]]
    assert headers[:5] == [
        "review_status",
        "target_table",
        "target_field",
        "source_table",
        "source_field",
    ]
    assert "transform_explanation" in headers
    assert ws["A2"].value == "AUTO_APPROVED"
    assert ws["B2"].value == "default.dim_store_info"
    context_values = {
        row[0].value: row[1].value
        for row in wb["source_context"].iter_rows(min_row=2, max_col=2)
    }
    assert context_values["debug_artifacts_dir"] == str(tmp_path / "debug")


def test_write_candidate_workbook_defaults_unqualified_source_table_to_default(
    tmp_path: Path,
) -> None:
    output = tmp_path / "review.xlsx"
    source = FieldLineageInput(
        dataset_urn=make_hive_dataset_urn("data_build.target"),
        table_name="data_build.target",
        etl_script="select project_id from dm_source",
        execute_shell="sh run.sh",
    )
    write_candidate_workbook(
        output,
        source_input=source,
        candidates=[
            FieldLineageCandidate(
                target_table="data_build.target",
                target_field="project_id",
                source_table="dm_source",
                source_field="project_id",
                transform_expression="project_id",
                confidence="HIGH",
            )
        ],
        unresolved_fields=[],
        llm_model="deepseek-test",
    )

    wb = load_workbook(output)
    assert wb["candidate_lineage"]["D2"].value == "default.dm_source"
    approved = load_approved_review_rows(output)
    grouped = group_approved_rows(approved)
    assert grouped[0].sources == (("default.dm_source", "project_id"),)


def test_write_candidate_workbook_expands_same_as_above_for_same_target(
    tmp_path: Path,
) -> None:
    output = tmp_path / "review.xlsx"
    source = FieldLineageInput(
        dataset_urn=make_hive_dataset_urn("data_logistics.target"),
        table_name="data_logistics.target",
        etl_script="select length, width from default.dim_sku_info_spec",
        execute_shell="sh run.sh",
    )
    expression = "round(percentile_approx(vol, 0.50), 3)"
    explanation = (
        "基于 default.dim_sku_info_spec 表的 length, width, height 字段计算体积后，"
        "按小分类分组取 50 分位值。"
    )
    evidence = "round(percentile_approx(vol, 0.50),3) as section_vol_1b2"
    write_candidate_workbook(
        output,
        source_input=source,
        candidates=[
            FieldLineageCandidate(
                target_table="data_logistics.target",
                target_field="section_vol_1b2",
                source_table="default.dim_sku_info_spec",
                source_field="length",
                transform_expression=expression,
                transform_explanation=explanation,
                evidence_sql=evidence,
                confidence="MEDIUM",
            ),
            FieldLineageCandidate(
                target_table="data_logistics.target",
                target_field="section_vol_1b2",
                source_table="default.dim_sku_info_spec",
                source_field="width",
                transform_expression="同上",
                transform_explanation="同上",
                evidence_sql="同上",
                confidence="MEDIUM",
            ),
        ],
        unresolved_fields=[],
        llm_model="deepseek-test",
    )

    wb = load_workbook(output)
    ws = wb["candidate_lineage"]
    assert ws["F3"].value == expression
    assert ws["G3"].value == explanation
    assert ws["H3"].value == evidence


def test_auto_review_marks_same_as_above_without_prior_context_for_review(
    tmp_path: Path,
) -> None:
    output = tmp_path / "review.xlsx"
    source = FieldLineageInput(
        dataset_urn=make_hive_dataset_urn("data_logistics.target"),
        table_name="data_logistics.target",
        etl_script="select length from default.dim_sku_info_spec",
        execute_shell="sh run.sh",
    )
    write_candidate_workbook(
        output,
        source_input=source,
        candidates=[
            FieldLineageCandidate(
                target_table="data_logistics.target",
                target_field="section_vol_1b2",
                source_table="default.dim_sku_info_spec",
                source_field="length",
                transform_expression="同上",
                transform_explanation="同上",
                evidence_sql="同上",
                confidence="HIGH",
            )
        ],
        unresolved_fields=[],
        llm_model="deepseek-test",
    )

    wb = load_workbook(output)
    assert wb["candidate_lineage"]["A2"].value == "NEEDS_REVIEW"


def test_load_review_rows_expands_same_as_above_before_import(tmp_path: Path) -> None:
    output = tmp_path / "review.xlsx"
    source = FieldLineageInput(
        dataset_urn=make_hive_dataset_urn("data_logistics.target"),
        table_name="data_logistics.target",
        etl_script="select length, width from default.dim_sku_info_spec",
        execute_shell="sh run.sh",
    )
    expression = "round(percentile_approx(vol, 0.50), 3)"
    explanation = "基于 length, width, height 字段计算体积后取 50 分位值。"
    write_candidate_workbook(
        output,
        source_input=source,
        candidates=[
            FieldLineageCandidate(
                target_table="data_logistics.target",
                target_field="section_vol_1b2",
                source_table="default.dim_sku_info_spec",
                source_field="length",
                transform_expression=expression,
                transform_explanation=explanation,
                evidence_sql="select section_vol_1b2",
                confidence="HIGH",
            ),
            FieldLineageCandidate(
                target_table="data_logistics.target",
                target_field="section_vol_1b2",
                source_table="default.dim_sku_info_spec",
                source_field="width",
                transform_expression="同上",
                transform_explanation="同上",
                evidence_sql="同上",
                confidence="HIGH",
            ),
        ],
        unresolved_fields=[],
        llm_model="deepseek-test",
    )
    wb = load_workbook(output)
    ws = wb["candidate_lineage"]
    ws["F3"] = "同上"
    ws["G3"] = "同上"
    ws["H3"] = "同上"
    wb.save(output)

    approved = load_approved_review_rows(output)

    assert approved[1].transform_expression == expression
    assert approved[1].transform_explanation == explanation
    assert approved[1].evidence_sql == "select section_vol_1b2"


def test_group_approved_rows_filters_same_as_above_transform_text() -> None:
    grouped = group_approved_rows(
        [
            FieldLineageCandidate(
                target_table="data_logistics.target",
                target_field="section_vol_1b2",
                source_table="default.dim_sku_info_spec",
                source_field="length",
                transform_expression="同上",
                transform_explanation="同上",
                confidence="HIGH",
            )
        ]
    )

    assert grouped[0].transform_operation == ""
    assert grouped[0].transform_explanation == ""


def test_group_approved_rows_keeps_explained_offline_source_without_upstreams() -> None:
    grouped = group_approved_rows(
        [
            FieldLineageCandidate(
                target_table="data_factory.target",
                target_field="manual_store_name",
                source_table="",
                source_field="",
                transform_expression="pandas.read_csv('/data/manual/store.csv')['store_name']",
                transform_explanation=(
                    "字段来自本地 CSV 文件 /data/manual/store.csv 的 store_name 列，"
                    "无法映射到 Hive 物理字段。"
                ),
                confidence="MEDIUM",
            )
        ]
    )

    assert len(grouped) == 1
    assert grouped[0].sources == ()
    assert "read_csv" in grouped[0].transform_operation
    assert "本地 CSV" in grouped[0].transform_operation


def test_load_review_rows_returns_approved_and_auto_approved_by_default(tmp_path: Path) -> None:
    output = tmp_path / "review.xlsx"
    source = FieldLineageInput(
        dataset_urn=make_hive_dataset_urn("default.dim_store_info"),
        table_name="default.dim_store_info",
        etl_script="select 1",
        execute_shell="sh run.sh",
    )
    rows = [
        FieldLineageCandidate(
            target_table="default.dim_store_info",
            target_field="auto_approved_field",
            source_table="ods.store_info",
            source_field="id",
            confidence="HIGH",
            transform_expression="cast(id as bigint)",
        ),
        FieldLineageCandidate(
            target_table="default.dim_store_info",
            target_field="pending_field",
            source_table="ods.store_info",
            source_field="name",
            confidence="MEDIUM",
        ),
        FieldLineageCandidate(
            target_table="default.dim_store_info",
            target_field="manual_approved_field",
            source_table="ods.store_info",
            source_field="code",
            confidence="MEDIUM",
            review_status=FieldLineageReviewStatus.APPROVED,
        ),
    ]
    write_candidate_workbook(
        output,
        source_input=source,
        candidates=rows,
        unresolved_fields=[],
        llm_model="deepseek-test",
    )

    wb = load_workbook(output)
    ws = wb["candidate_lineage"]
    ws["A2"] = "AUTO_APPROVED"
    ws["A3"] = "PENDING"
    ws["A4"] = "APPROVED"
    wb.save(output)

    approved = load_approved_review_rows(output)

    assert [row.target_field for row in approved] == [
        "auto_approved_field",
        "manual_approved_field",
    ]


def test_load_review_rows_can_limit_import_statuses(tmp_path: Path) -> None:
    output = tmp_path / "review.xlsx"
    source = FieldLineageInput(
        dataset_urn=make_hive_dataset_urn("default.dim_store_info"),
        table_name="default.dim_store_info",
        etl_script="select 1",
        execute_shell="sh run.sh",
    )
    write_candidate_workbook(
        output,
        source_input=source,
        candidates=[
            FieldLineageCandidate(
                target_table="default.dim_store_info",
                target_field="auto_approved_field",
                source_table="ods.store_info",
                source_field="id",
                transform_expression="cast(id as bigint)",
                confidence="HIGH",
            ),
            FieldLineageCandidate(
                target_table="default.dim_store_info",
                target_field="manual_approved_field",
                source_table="ods.store_info",
                source_field="code",
                transform_expression="code",
                confidence="MEDIUM",
                review_status=FieldLineageReviewStatus.APPROVED,
            ),
        ],
        unresolved_fields=[],
        llm_model="deepseek-test",
    )

    rows = load_approved_review_rows(
        output,
        import_statuses={FieldLineageReviewStatus.APPROVED},
    )

    assert [row.target_field for row in rows] == ["manual_approved_field"]


def test_auto_review_marks_risky_high_confidence_rows_for_review(tmp_path: Path) -> None:
    output = tmp_path / "review.xlsx"
    source = FieldLineageInput(
        dataset_urn=make_hive_dataset_urn("default.dim_store_info"),
        table_name="default.dim_store_info",
        etl_script="select 1",
        execute_shell="sh run.sh",
    )
    write_candidate_workbook(
        output,
        source_input=source,
        candidates=[
            FieldLineageCandidate(
                target_table="default.dim_store_info",
                target_field="safe_field",
                source_table="ods.store_info",
                source_field="id",
                transform_expression="cast(id as bigint)",
                confidence="HIGH",
            ),
            FieldLineageCandidate(
                target_table="default.dim_store_info",
                target_field="bad_source_field",
                source_table="ods.store_info",
                source_field="id,name",
                transform_expression="concat(id, name)",
                confidence="HIGH",
            ),
            FieldLineageCandidate(
                target_table="default.dim_store_info",
                target_field="unresolved_var_field",
                source_table="ods.store_info",
                source_field="id",
                transform_expression="cast(${SOURCE_ID} as bigint)",
                confidence="HIGH",
            ),
            FieldLineageCandidate(
                target_table="default.dim_store_info",
                target_field="missing_source_field",
                source_table="ods.store_info",
                source_field="",
                transform_expression="if(id is null, name, id)",
                confidence="HIGH",
                review_status=FieldLineageReviewStatus.AUTO_APPROVED,
            ),
        ],
        unresolved_fields=[],
        llm_model="deepseek-test",
    )

    wb = load_workbook(output)
    rows = list(wb["candidate_lineage"].iter_rows(min_row=2, max_col=2, values_only=True))

    assert rows == [
        ("AUTO_APPROVED", "default.dim_store_info"),
        ("NEEDS_REVIEW", "default.dim_store_info"),
        ("AUTO_APPROVED", "default.dim_store_info"),
        ("NEEDS_REVIEW", "default.dim_store_info"),
    ]


def test_auto_review_allows_unresolved_variable_in_evidence_sql(tmp_path: Path) -> None:
    output = tmp_path / "review.xlsx"
    source = FieldLineageInput(
        dataset_urn=make_hive_dataset_urn("data_build.location_type"),
        table_name="data_build.location_type",
        etl_script="select t1.store_code from source t1 where dt='${DATE}'",
        execute_shell="sh run.sh",
    )
    write_candidate_workbook(
        output,
        source_input=source,
        candidates=[
            FieldLineageCandidate(
                target_table="data_build.location_type",
                target_field="store_code",
                source_table="data_build.dwd_store_construction_project_upload_info_v1",
                source_field="store_code",
                transform_expression="t1.store_code",
                evidence_sql="left join source t1 on t1.dt = '${DATE}' select t1.store_code",
                confidence="HIGH",
            )
        ],
        unresolved_fields=[],
        llm_model="deepseek-test",
    )

    wb = load_workbook(output)
    status = wb["candidate_lineage"]["A2"].value
    assert status == "AUTO_APPROVED"


def test_auto_review_approves_runtime_variable_string_constant(tmp_path: Path) -> None:
    output = tmp_path / "review.xlsx"
    source = FieldLineageInput(
        dataset_urn=make_hive_dataset_urn("data_factory.target"),
        table_name="data_factory.target",
        etl_script="select '${FORMAT_DATE}' as cal_date",
        execute_shell="sh run.sh",
    )
    write_candidate_workbook(
        output,
        source_input=source,
        candidates=[
            FieldLineageCandidate(
                target_table="data_factory.target",
                target_field="cal_date",
                source_table="",
                source_field="",
                transform_expression="'${FORMAT_DATE}'",
                transform_explanation=(
                    "取调度日期变量 FORMAT_DATE 作为业务日期，表示数据计算对应的日历日期。"
                ),
                evidence_sql="'${FORMAT_DATE}' as cal_date",
                confidence="HIGH",
            )
        ],
        unresolved_fields=[],
        llm_model="deepseek-test",
    )

    wb = load_workbook(output)
    assert wb["candidate_lineage"]["A2"].value == "AUTO_APPROVED"


def test_auto_review_approves_medium_confidence_hive_source(tmp_path: Path) -> None:
    output = tmp_path / "review.xlsx"
    source = FieldLineageInput(
        dataset_urn=make_hive_dataset_urn("data_factory.target"),
        table_name="data_factory.target",
        etl_script="select cast(order_id as string) as order_id from ods.order_detail",
        execute_shell="sh run.sh",
    )
    write_candidate_workbook(
        output,
        source_input=source,
        candidates=[
            FieldLineageCandidate(
                target_table="data_factory.target",
                target_field="order_id",
                source_table="ods.order_detail",
                source_field="order_id",
                transform_expression="cast(order_id as string)",
                transform_explanation="取 ods.order_detail.order_id 并转换为字符串。",
                confidence="MEDIUM",
            )
        ],
        unresolved_fields=[],
        llm_model="deepseek-test",
    )

    wb = load_workbook(output)
    assert wb["candidate_lineage"]["A2"].value == "AUTO_APPROVED"


def test_auto_review_approves_explained_offline_source(tmp_path: Path) -> None:
    output = tmp_path / "review.xlsx"
    source = FieldLineageInput(
        dataset_urn=make_hive_dataset_urn("data_factory.target"),
        table_name="data_factory.target",
        etl_script="df = pandas.read_csv('/data/manual/store.csv')",
        execute_shell="python job.py",
    )
    write_candidate_workbook(
        output,
        source_input=source,
        candidates=[
            FieldLineageCandidate(
                target_table="data_factory.target",
                target_field="manual_store_name",
                source_table="",
                source_field="",
                transform_expression="pandas.read_csv('/data/manual/store.csv')['store_name']",
                transform_explanation=(
                    "字段来自本地 CSV 文件 /data/manual/store.csv 的 store_name 列，"
                    "无法映射到 Hive 物理字段。"
                ),
                confidence="MEDIUM",
            )
        ],
        unresolved_fields=[],
        llm_model="deepseek-test",
    )

    wb = load_workbook(output)
    assert wb["candidate_lineage"]["A2"].value == "AUTO_APPROVED"


def test_write_candidate_workbook_excludes_partition_fields(tmp_path: Path) -> None:
    output = tmp_path / "review.xlsx"
    source = FieldLineageInput(
        dataset_urn=make_hive_dataset_urn("default.dim_store_info"),
        table_name="default.dim_store_info",
        etl_script="select 1",
        execute_shell="sh run.sh",
    )
    write_candidate_workbook(
        output,
        source_input=source,
        candidates=[
            FieldLineageCandidate(
                target_table="default.dim_store_info",
                target_field="store_id",
                source_table="ods.store_info",
                source_field="id",
                transform_expression="cast(id as bigint)",
                confidence="HIGH",
            ),
            FieldLineageCandidate(
                target_table="default.dim_store_info",
                target_field="dt",
                source_table="ods.store_info",
                source_field="dt",
                transform_expression="'20260604'",
                confidence="HIGH",
            ),
            FieldLineageCandidate(
                target_table="default.dim_store_info",
                target_field="hr",
                source_table="ods.store_info",
                source_field="hr",
                transform_expression="'11'",
                confidence="HIGH",
            ),
        ],
        unresolved_fields=[
            UnresolvedField(target_field="dt", reason="partition field"),
            UnresolvedField(target_field="hr", reason="partition field"),
            UnresolvedField(target_field="unknown_field", reason="dynamic SQL"),
        ],
        llm_model="deepseek-test",
    )

    wb = load_workbook(output)
    fields = [
        row[0]
        for row in wb["candidate_lineage"].iter_rows(
            min_row=2,
            min_col=3,
            max_col=3,
            values_only=True,
        )
    ]
    assert fields == ["store_id", "unknown_field"]
    statuses = [
        row[0]
        for row in wb["candidate_lineage"].iter_rows(
            min_row=2,
            min_col=1,
            max_col=1,
            values_only=True,
        )
    ]
    assert statuses == ["AUTO_APPROVED", "NEEDS_REVIEW"]
    unresolved_fields = [
        row[0]
        for row in wb["unresolved_fields"].iter_rows(
            min_row=2,
            min_col=1,
            max_col=1,
            values_only=True,
        )
    ]
    assert unresolved_fields == ["unknown_field"]


def test_write_candidate_workbook_excludes_hr_partition_from_schema_fill(
    tmp_path: Path,
) -> None:
    workbook = tmp_path / "review.xlsx"
    source = FieldLineageInput(
        dataset_urn=make_hive_dataset_urn("default.product_shop_sku_hi"),
        table_name="default.product_shop_sku_hi",
        etl_script="select sku_code from ods.product_shop_sku",
        execute_shell="sh run.sh",
        target_schema_fields=["sku_code", "dt", "hr"],
    )

    write_candidate_workbook(
        workbook,
        source_input=source,
        candidates=[
            FieldLineageCandidate(
                target_table="default.product_shop_sku_hi",
                target_field="sku_code",
                source_table="ods.product_shop_sku",
                source_field="sku_code",
                transform_expression="sku_code",
                confidence="HIGH",
            )
        ],
        unresolved_fields=[],
        llm_model="deepseek-test",
    )

    wb = load_workbook(workbook)
    target_fields = [
        row[0]
        for row in wb["candidate_lineage"].iter_rows(
            min_row=2,
            min_col=3,
            max_col=3,
            values_only=True,
        )
    ]
    assert target_fields == ["sku_code"]


def test_write_candidate_workbook_excludes_version_partition_from_schema_fill(
    tmp_path: Path,
) -> None:
    workbook = tmp_path / "review.xlsx"
    source = FieldLineageInput(
        dataset_urn=make_hive_dataset_urn(
            "data_build.dm_site_selection_competitor_daily_sales_v1"
        ),
        table_name="data_build.dm_site_selection_competitor_daily_sales_v1",
        etl_script="select site_id from ods.site_selection",
        execute_shell="sh run.sh",
        target_schema_fields=["site_id", "version"],
        target_partition_fields=["version"],
    )

    write_candidate_workbook(
        workbook,
        source_input=source,
        candidates=[
            FieldLineageCandidate(
                target_table=source.table_name,
                target_field="site_id",
                source_table="ods.site_selection",
                source_field="site_id",
                transform_expression="site_id",
                confidence="HIGH",
            )
        ],
        unresolved_fields=[
            UnresolvedField(
                target_field="version",
                reason="DataHub DDL 字段未在 LLM 输出中找到，请按 Hive insert select 顺序确认来源",
            )
        ],
        llm_model="deepseek-test",
    )

    wb = load_workbook(workbook)
    target_fields = [
        row[0]
        for row in wb["candidate_lineage"].iter_rows(
            min_row=2,
            min_col=3,
            max_col=3,
            values_only=True,
        )
    ]
    assert target_fields == ["site_id"]
    assert list(wb["unresolved_fields"].iter_rows(min_row=2, values_only=True)) == []


def test_write_candidate_workbook_excludes_schema_partition_key_fields(
    tmp_path: Path,
) -> None:
    workbook = tmp_path / "review.xlsx"
    source = FieldLineageInput(
        dataset_urn=make_hive_dataset_urn("default.product_shop_sku_hi"),
        table_name="default.product_shop_sku_hi",
        etl_script="select sku_code from ods.product_shop_sku",
        execute_shell="sh run.sh",
        target_schema_fields=["sku_code", "kpt"],
        target_partition_fields=["kpt"],
    )

    write_candidate_workbook(
        workbook,
        source_input=source,
        candidates=[
            FieldLineageCandidate(
                target_table="default.product_shop_sku_hi",
                target_field="sku_code",
                source_table="ods.product_shop_sku",
                source_field="sku_code",
                transform_expression="sku_code",
                confidence="HIGH",
            ),
            FieldLineageCandidate(
                target_table="default.product_shop_sku_hi",
                target_field="kpt",
                source_table="ods.product_shop_sku",
                source_field="kpt",
                transform_expression="kpt",
                confidence="HIGH",
            ),
        ],
        unresolved_fields=[UnresolvedField(target_field="kpt", reason="partition key")],
        llm_model="deepseek-test",
    )

    wb = load_workbook(workbook)
    target_fields = [
        row[0]
        for row in wb["candidate_lineage"].iter_rows(
            min_row=2,
            min_col=3,
            max_col=3,
            values_only=True,
        )
    ]
    assert target_fields == ["sku_code"]
    assert list(wb["unresolved_fields"].iter_rows(min_row=2, values_only=True)) == []


def test_write_candidate_workbook_auto_approves_null_constant_field(tmp_path: Path) -> None:
    workbook = tmp_path / "review.xlsx"
    source = FieldLineageInput(
        dataset_urn=make_hive_dataset_urn("default.plan_price"),
        table_name="default.plan_price",
        etl_script="insert overwrite table default.plan_price select null as vendor_code from tmp1",
        execute_shell="sh run.sh",
        target_schema_fields=["plan_code", "vendor_code", "dt"],
    )

    write_candidate_workbook(
        workbook,
        source_input=source,
        candidates=[
            FieldLineageCandidate(
                target_table="default.plan_price",
                target_field="vendor_code",
                source_table="",
                source_field="",
                transform_expression="NULL",
                transform_explanation="目标字段 vendor_code 由常量 NULL 写入，表示业务系统废弃字段，无上游来源字段。",
                evidence_sql="null as vendor_code",
                confidence="HIGH",
            )
        ],
        unresolved_fields=[],
        llm_model="deepseek-test",
    )

    wb = load_workbook(workbook)
    headers = [cell.value for cell in wb["candidate_lineage"][1]]
    rows = [
        dict(zip(headers, row))
        for row in wb["candidate_lineage"].iter_rows(min_row=2, values_only=True)
    ]
    vendor_row = next(row for row in rows if row["target_field"] == "vendor_code")
    assert vendor_row["review_status"] == "AUTO_APPROVED"
    assert vendor_row["source_table"] is None
    assert vendor_row["source_field"] is None
    assert vendor_row["transform_expression"] == "NULL"


def test_write_candidate_workbook_auto_approves_aliased_null_constant_field(
    tmp_path: Path,
) -> None:
    workbook = tmp_path / "review.xlsx"
    source = FieldLineageInput(
        dataset_urn=make_hive_dataset_urn(
            "data_build.dm_site_selection_competitor_daily_sales_v1"
        ),
        table_name="data_build.dm_site_selection_competitor_daily_sales_v1",
        etl_script="select null as seasonal_coefficient",
        execute_shell="sh run.sh",
        target_schema_fields=["seasonal_coefficient"],
    )

    write_candidate_workbook(
        workbook,
        source_input=source,
        candidates=[
            FieldLineageCandidate(
                target_table=source.table_name,
                target_field="seasonal_coefficient",
                source_table="",
                source_field="",
                transform_expression="null as seasonal_coefficient",
                transform_explanation="该字段写入 NULL 常量，无上游来源字段，表示季节系数暂未计算。",
                evidence_sql="null as seasonal_coefficient",
                confidence="HIGH",
            )
        ],
        unresolved_fields=[],
        llm_model="deepseek-test",
    )

    wb = load_workbook(workbook)
    assert wb["candidate_lineage"]["A2"].value == "AUTO_APPROVED"

    approved = load_approved_review_rows(workbook)
    grouped = group_approved_rows(approved)
    assert len(grouped) == 1
    assert grouped[0].sources == ()
    assert "null as seasonal_coefficient" in grouped[0].transform_operation


def test_write_candidate_workbook_auto_approves_unresolved_null_constant_field(
    tmp_path: Path,
) -> None:
    workbook = tmp_path / "review.xlsx"
    source = FieldLineageInput(
        dataset_urn=make_hive_dataset_urn("default.pdw_bach_baseinfo_price_price_plan_detail"),
        table_name="default.pdw_bach_baseinfo_price_price_plan_detail",
        etl_script="insert overwrite table default.pdw_bach_baseinfo_price_price_plan_detail select null as vendor_code from tmp1",
        execute_shell="sh run.sh",
        target_schema_fields=["plan_code", "vendor_code", "dt"],
    )

    write_candidate_workbook(
        workbook,
        source_input=source,
        candidates=[],
        unresolved_fields=[
            UnresolvedField(
                target_field="vendor_code",
                reason=(
                    "SQL 第3列表达式为 `null as vendor_code`，写入 NULL 常量，"
                    "没有可确认的来源表字段。"
                ),
            )
        ],
        llm_model="deepseek-test",
    )

    wb = load_workbook(workbook)
    headers = [cell.value for cell in wb["candidate_lineage"][1]]
    rows = [
        dict(zip(headers, row))
        for row in wb["candidate_lineage"].iter_rows(min_row=2, values_only=True)
    ]
    vendor_row = next(row for row in rows if row["target_field"] == "vendor_code")
    assert vendor_row["review_status"] == "AUTO_APPROVED"
    assert vendor_row["source_table"] is None
    assert vendor_row["source_field"] is None
    assert vendor_row["transform_expression"] == "NULL"
    assert vendor_row["evidence_sql"] == "null as vendor_code"
    assert list(wb["unresolved_fields"].iter_rows(min_row=2, values_only=True)) == []


def test_write_candidate_workbook_auto_approves_hardcoded_null_unresolved_field(
    tmp_path: Path,
) -> None:
    workbook = tmp_path / "review.xlsx"
    source = FieldLineageInput(
        dataset_urn=make_hive_dataset_urn(
            "data_build.dm_site_selection_competitor_daily_sales_v1"
        ),
        table_name="data_build.dm_site_selection_competitor_daily_sales_v1",
        etl_script="select null as seasonal_coefficient",
        execute_shell="sh run.sh",
        target_schema_fields=["seasonal_coefficient"],
    )

    write_candidate_workbook(
        workbook,
        source_input=source,
        candidates=[],
        unresolved_fields=[
            UnresolvedField(
                target_field="seasonal_coefficient",
                reason="硬编码为 null，无来源字段",
            )
        ],
        llm_model="deepseek-test",
    )

    wb = load_workbook(workbook)
    headers = [cell.value for cell in wb["candidate_lineage"][1]]
    row = dict(zip(headers, next(wb["candidate_lineage"].iter_rows(min_row=2, values_only=True))))
    assert row["review_status"] == "AUTO_APPROVED"
    assert row["target_field"] == "seasonal_coefficient"
    assert row["source_table"] is None
    assert row["source_field"] is None
    assert row["transform_expression"] == "NULL"
    assert row["transform_explanation"] == (
        "目标字段 seasonal_coefficient 由常量 NULL 写入，无上游来源字段。"
    )
    assert row["evidence_sql"] == "NULL AS seasonal_coefficient"
    assert list(wb["unresolved_fields"].iter_rows(min_row=2, values_only=True)) == []


def test_write_candidate_workbook_auto_approves_runtime_constant_field(
    tmp_path: Path,
) -> None:
    workbook = tmp_path / "review.xlsx"
    source = FieldLineageInput(
        dataset_urn=make_hive_dataset_urn("default.price_plan"),
        table_name="default.price_plan",
        etl_script="select cast(date_format(current_date(),'yyyy-MM-dd') as string) as update_time",
        execute_shell="sh run.sh",
        target_schema_fields=["update_time"],
    )

    write_candidate_workbook(
        workbook,
        source_input=source,
        candidates=[],
        unresolved_fields=[
            UnresolvedField(
                target_field="update_time",
                reason=(
                    "新增数据分支第 11 个表达式为 "
                    "`cast(date_format(current_date(),'yyyy-MM-dd') as string) as update_time`，"
                    "是运行时当前日期生成值，不来自任何源表字段；因此无法给出源表字段级血缘。"
                ),
            )
        ],
        llm_model="deepseek-test",
    )

    wb = load_workbook(workbook)
    headers = [cell.value for cell in wb["candidate_lineage"][1]]
    row = dict(zip(headers, next(wb["candidate_lineage"].iter_rows(min_row=2, values_only=True))))
    assert row["review_status"] == "AUTO_APPROVED"
    assert row["source_table"] is None
    assert row["source_field"] is None
    assert row["target_field"] == "update_time"
    assert row["transform_expression"] == "cast(date_format(current_date(),'yyyy-MM-dd') as string)"
    assert row["evidence_sql"] == "cast(date_format(current_date(),'yyyy-MM-dd') as string) as update_time"
    assert list(wb["unresolved_fields"].iter_rows(min_row=2, values_only=True)) == []


def test_write_candidate_workbook_filters_to_schema_and_fills_missing_fields(
    tmp_path: Path,
) -> None:
    output = tmp_path / "review.xlsx"
    source = FieldLineageInput(
        dataset_urn=make_hive_dataset_urn(
            "default.pdw_order_store_90_order_detail_booking_main_di"
        ),
        table_name="default.pdw_order_store_90_order_detail_booking_main_di",
        etl_script=(
            "insert overwrite table target select col1, "
            "get_json_object(data, '$.bookingExtendInfo.bookingJsonInfo') from ods.source"
        ),
        execute_shell="sh run.sh",
        target_schema_fields=["order_id", "booking_json_info", "dt"],
    )

    write_candidate_workbook(
        output,
        source_input=source,
        candidates=[
            FieldLineageCandidate(
                target_table=source.table_name,
                target_field="order_id",
                source_table="ods.order_detail",
                source_field="order_id",
                transform_expression="order_id",
                confidence="HIGH",
            )
        ],
        unresolved_fields=[
            UnresolvedField(
                target_field="bookingextendinfo_bookingjsoninfo",
                reason="LLM used JSON path as target field name",
            )
        ],
        llm_model="deepseek-test",
    )

    wb = load_workbook(output)
    headers = [cell.value for cell in wb["candidate_lineage"][1]]
    rows = [
        dict(zip(headers, row))
        for row in wb["candidate_lineage"].iter_rows(min_row=2, values_only=True)
    ]

    assert [row["target_field"] for row in rows] == [
        "order_id",
        "booking_json_info",
    ]
    assert rows[0]["review_status"] == "AUTO_APPROVED"
    assert rows[1]["review_status"] == "NEEDS_REVIEW"
    assert rows[1]["source_table"] is None
    assert "DataHub DDL 字段" in rows[1]["llm_notes"]
    assert list(wb["unresolved_fields"].iter_rows(min_row=2, values_only=True)) == []


def test_build_field_lineage_user_message_includes_schema_order_rule() -> None:
    source = FieldLineageInput(
        dataset_urn=make_hive_dataset_urn("default.ods_table1"),
        table_name="default.ods_table1",
        etl_script=(
            "insert overwrite table ods_table1 "
            "select user_name as col1, col2, user_age from ods_table3"
        ),
        execute_shell="sh run.sh",
        target_schema_fields=["col1", "col2", "col3", "dt"],
    )

    message = build_field_lineage_user_message(source)

    assert "目标表 DataHub DDL 字段顺序" in message
    assert "1. col1" in message
    assert "3. col3" in message
    assert "按上面字段顺序与 SELECT 表达式位置对齐" in message


def test_build_field_lineage_user_message_includes_not_verified_alias() -> None:
    source = FieldLineageInput(
        dataset_urn=make_hive_dataset_urn(
            "data_smartorder.dw_ordering_opportunity_loss_actual_sold_out_time_store_sku_main_sku"
        ),
        table_name=(
            "data_smartorder."
            "dw_ordering_opportunity_loss_actual_sold_out_time_store_sku_main_sku"
        ),
        etl_script=(
            "insert overwrite table data_smartorder."
            "not_verified_dw_ordering_opportunity_loss_actual_sold_out_time_store_sku_main_sku "
            "select store_code, sku_main_code from ods.source"
        ),
        execute_shell="sh run.sh",
        target_table_aliases=build_target_table_aliases(
            "data_smartorder."
            "dw_ordering_opportunity_loss_actual_sold_out_time_store_sku_main_sku"
        ),
    )

    message = build_field_lineage_user_message(source)

    assert (
        "data_smartorder."
        "not_verified_dw_ordering_opportunity_loss_actual_sold_out_time_store_sku_main_sku"
        in message
    )
    assert "等同于" in message
    assert "不要判定为未写入目标表" in message


def test_write_candidate_workbook_canonicalizes_not_verified_target_table(
    tmp_path: Path,
) -> None:
    output = tmp_path / "review.xlsx"
    target = (
        "data_smartorder."
        "dw_ordering_opportunity_loss_actual_sold_out_time_store_sku_main_sku"
    )
    alias = (
        "data_smartorder."
        "not_verified_dw_ordering_opportunity_loss_actual_sold_out_time_store_sku_main_sku"
    )
    source = FieldLineageInput(
        dataset_urn=make_hive_dataset_urn(target),
        table_name=target,
        etl_script="insert overwrite table alias select store_code from ods.source",
        execute_shell="sh run.sh",
        target_schema_fields=["store_code"],
        target_table_aliases=[alias],
    )

    write_candidate_workbook(
        output,
        source_input=source,
        candidates=[
            FieldLineageCandidate(
                target_table=alias,
                target_field="store_code",
                source_table="ods.source",
                source_field="store_code",
                transform_expression="store_code",
                confidence="HIGH",
            )
        ],
        unresolved_fields=[],
        llm_model="deepseek-test",
    )

    wb = load_workbook(output)
    headers = [cell.value for cell in wb["candidate_lineage"][1]]
    row = dict(zip(headers, next(wb["candidate_lineage"].iter_rows(min_row=2, values_only=True))))
    assert row["target_table"] == target
    assert row["target_field"] == "store_code"


def test_write_candidate_workbook_maps_create_table_like_fields(
    tmp_path: Path,
) -> None:
    output = tmp_path / "review.xlsx"
    source = FieldLineageInput(
        dataset_urn=make_hive_dataset_urn("default.dw_order_sku_v1"),
        table_name="default.dw_order_sku_v1",
        etl_script=(
            "create table if not exists default.dw_order_sku_v1\n"
            "like default.dw_order_sku_v1_archive;"
        ),
        execute_shell="sh run.sh",
        target_schema_fields=["order_id", "sku_code", "dt"],
    )

    write_candidate_workbook(
        output,
        source_input=source,
        candidates=[],
        unresolved_fields=[],
        llm_model="deepseek-test",
    )

    wb = load_workbook(output)
    headers = [cell.value for cell in wb["candidate_lineage"][1]]
    rows = [
        dict(zip(headers, row))
        for row in wb["candidate_lineage"].iter_rows(min_row=2, values_only=True)
    ]

    assert [row["target_field"] for row in rows] == ["order_id", "sku_code"]
    assert [row["source_table"] for row in rows] == [
        "default.dw_order_sku_v1_archive",
        "default.dw_order_sku_v1_archive",
    ]
    assert [row["source_field"] for row in rows] == ["order_id", "sku_code"]
    assert [row["review_status"] for row in rows] == ["AUTO_APPROVED", "AUTO_APPROVED"]
    assert "CREATE TABLE LIKE" in rows[0]["transform_explanation"]


def test_extract_schema_field_names_from_openapi_payload() -> None:
    payload = {
        "value": {
            "fields": [
                {"fieldPath": "[version=2.0].[type=string].order_id"},
                {"fieldPath": "[version=2.0].[type=string].booking_json_info"},
                {"fieldPath": "[version=2.0].[type=string].dt"},
            ]
        }
    }

    assert extract_schema_field_names(payload) == [
        "order_id",
        "booking_json_info",
        "dt",
    ]


def test_write_batch_summary_workbook_counts_review_statuses(tmp_path: Path) -> None:
    batch_dir = tmp_path / "batch"
    batch_dir.mkdir()
    source = FieldLineageInput(
        dataset_urn=make_hive_dataset_urn("default.dim_store_info"),
        table_name="default.dim_store_info",
        etl_script="select 1",
        execute_shell="sh run.sh",
    )
    workbook = batch_dir / "202606041414_default.dim_store_info.xlsx"
    write_candidate_workbook(
        workbook,
        source_input=source,
        candidates=[
            FieldLineageCandidate(
                target_table="default.dim_store_info",
                target_field="safe_field",
                source_table="ods.store_info",
                source_field="id",
                transform_expression="cast(id as bigint)",
                confidence="HIGH",
            ),
            FieldLineageCandidate(
                target_table="default.dim_store_info",
                target_field="review_field",
                source_table="ods.store_info",
                source_field="id,name",
                transform_expression="concat(id, name)",
                confidence="HIGH",
            ),
        ],
        unresolved_fields=[UnresolvedField(target_field="missing_field", reason="dynamic SQL")],
        llm_model="deepseek-test",
    )

    summary = tmp_path / "summary.xlsx"
    write_batch_summary_workbook(batch_dir, summary)

    wb = load_workbook(summary)
    ws = wb["table_summary"]
    headers = [cell.value for cell in ws[1]]
    row = {headers[idx]: value for idx, value in enumerate(next(ws.iter_rows(min_row=2, values_only=True)))}
    assert row["table_name"] == "default.dim_store_info"
    assert row["total_candidate_count"] == 3
    assert row["auto_approved_count"] == 1
    assert row["auto_approved_percent"] == 1 / 3
    assert row["needs_review_count"] == 2
    assert row["unresolved_field_count"] == 1


def test_batch_summary_ignores_blank_unresolved_rows(tmp_path: Path) -> None:
    batch_dir = tmp_path / "batch"
    batch_dir.mkdir()
    source = FieldLineageInput(
        dataset_urn=make_hive_dataset_urn("default.dim_store_info"),
        table_name="default.dim_store_info",
        etl_script="select 1",
        execute_shell="sh run.sh",
    )
    workbook = batch_dir / "202606041414_default.dim_store_info.xlsx"
    write_candidate_workbook(
        workbook,
        source_input=source,
        candidates=[
            FieldLineageCandidate(
                target_table="default.dim_store_info",
                target_field="safe_field",
                source_table="ods.store_info",
                source_field="id",
                transform_expression="cast(id as bigint)",
                confidence="HIGH",
            )
        ],
        unresolved_fields=[],
        llm_model="deepseek-test",
    )
    wb = load_workbook(workbook)
    wb["unresolved_fields"]["A2"] = None
    wb["unresolved_fields"]["B2"] = None
    wb.save(workbook)

    summary = tmp_path / "summary.xlsx"
    write_batch_summary_workbook(batch_dir, summary)

    wb = load_workbook(summary)
    ws = wb["table_summary"]
    headers = [cell.value for cell in ws[1]]
    row = {headers[idx]: value for idx, value in enumerate(next(ws.iter_rows(min_row=2, values_only=True)))}
    assert row["unresolved_field_count"] == 0


def test_batch_summary_counts_candidate_lineage_rows_for_percent(tmp_path: Path) -> None:
    batch_dir = tmp_path / "batch"
    batch_dir.mkdir()
    source = FieldLineageInput(
        dataset_urn=make_hive_dataset_urn("default.dim_store_info"),
        table_name="default.dim_store_info",
        etl_script="select 1",
        execute_shell="sh run.sh",
    )
    workbook = batch_dir / "202606041414_default.dim_store_info.xlsx"
    write_candidate_workbook(
        workbook,
        source_input=source,
        candidates=[
            FieldLineageCandidate(
                target_table="default.dim_store_info",
                target_field="store_address",
                source_table="ods.store_a",
                source_field="address",
                transform_expression="coalesce(a.address, b.address)",
                confidence="HIGH",
            ),
            FieldLineageCandidate(
                target_table="default.dim_store_info",
                target_field="store_address",
                source_table="ods.store_b",
                source_field="address",
                transform_expression="coalesce(a.address, b.address)",
                confidence="HIGH",
            ),
            FieldLineageCandidate(
                target_table="default.dim_store_info",
                target_field="store_name",
                source_table="ods.store_a",
                source_field="name,id",
                transform_expression="concat(name, id)",
                confidence="HIGH",
            ),
        ],
        unresolved_fields=[],
        llm_model="deepseek-test",
    )

    summary = tmp_path / "summary.xlsx"
    write_batch_summary_workbook(batch_dir, summary)

    wb = load_workbook(summary)
    ws = wb["table_summary"]
    headers = [cell.value for cell in ws[1]]
    row = {headers[idx]: value for idx, value in enumerate(next(ws.iter_rows(min_row=2, values_only=True)))}
    assert row["total_candidate_count"] == 3
    assert row["auto_approved_count"] == 2
    assert row["auto_approved_percent"] == 2 / 3


def test_group_coalesce_rows_merge_sources_and_transform() -> None:
    rows = [
        FieldLineageCandidate(
            target_table="default.dim_store_info",
            target_field="store_address",
            source_table="ods.bach_store",
            source_field="store_address",
            transform_expression="coalesce(bach.store_address, hd.store_address)",
            transform_explanation=(
                "优先取 ods.bach_store 表中的 store_address 字段，"
                "为空时取 ods.hd_store 表中的 store_address 字段，"
                "表示门店地址按优先级兜底合并。"
            ),
        ),
        FieldLineageCandidate(
            target_table="default.dim_store_info",
            target_field="store_address",
            source_table="ods.hd_store",
            source_field="store_address",
            transform_expression="coalesce(bach.store_address, hd.store_address)",
            transform_explanation=(
                "优先取 ods.bach_store 表中的 store_address 字段，"
                "为空时取 ods.hd_store 表中的 store_address 字段，"
                "表示门店地址按优先级兜底合并。"
            ),
        ),
    ]
    grouped = group_approved_rows(rows)
    assert len(grouped) == 1
    assert grouped[0].target_field == "store_address"
    assert len(grouped[0].sources) == 2
    _coalesce_explanation = (
        "优先取 ods.bach_store 表中的 store_address 字段，"
        "为空时取 ods.hd_store 表中的 store_address 字段，"
        "表示门店地址按优先级兜底合并。"
    )
    assert grouped[0].transform_operation == (
        f"/* 中文解释：{_coalesce_explanation} */\n"
        "coalesce(bach.store_address, hd.store_address)"
    )
    assert grouped[0].transform_explanation == _coalesce_explanation


def test_group_approved_rows_skips_partition_fields() -> None:
    grouped = group_approved_rows(
        [
            FieldLineageCandidate(
                target_table="default.dim_store_info",
                target_field="store_id",
                source_table="ods.store_info",
                source_field="id",
                transform_expression="cast(id as bigint)",
            ),
            FieldLineageCandidate(
                target_table="default.dim_store_info",
                target_field="dt",
                source_table="ods.store_info",
                source_field="dt",
                transform_expression="'20260604'",
            ),
            FieldLineageCandidate(
                target_table="default.dim_store_info",
                target_field="hr",
                source_table="ods.store_info",
                source_field="hr",
                transform_expression="'11'",
            ),
        ]
    )

    assert [item.target_field for item in grouped] == ["store_id"]


def test_write_candidate_workbook_excludes_self_dependency_rows(tmp_path: Path) -> None:
    output = tmp_path / "review.xlsx"
    source = FieldLineageInput(
        dataset_urn=make_hive_dataset_urn("default.pdw_cvs_usercenter_weixin_id_mapping"),
        table_name="default.pdw_cvs_usercenter_weixin_id_mapping",
        etl_script="select 1",
        execute_shell="sh run.sh",
    )
    write_candidate_workbook(
        output,
        source_input=source,
        candidates=[
            FieldLineageCandidate(
                target_table="default.pdw_cvs_usercenter_weixin_id_mapping",
                target_field="user_id",
                source_table="default.pdw_cvs_usercenter_weixin_id_mapping",
                source_field="user_id",
                transform_expression="user_id",
                confidence="HIGH",
            ),
            FieldLineageCandidate(
                target_table="default.pdw_cvs_usercenter_weixin_id_mapping",
                target_field="user_id",
                source_table="default.ods_cvs_usercenter_weixin_id_mapping_di",
                source_field="user_id",
                transform_expression="user_id",
                confidence="HIGH",
            ),
        ],
        unresolved_fields=[],
        llm_model="deepseek-test",
    )

    wb = load_workbook(output)
    rows = list(wb["candidate_lineage"].iter_rows(min_row=2, values_only=True))

    assert len(rows) == 1
    headers = [cell.value for cell in wb["candidate_lineage"][1]]
    rec = dict(zip(headers, rows[0]))
    assert rec["source_table"] == "default.ods_cvs_usercenter_weixin_id_mapping_di"


def test_group_approved_rows_skips_self_dependency_rows() -> None:
    grouped = group_approved_rows(
        [
            FieldLineageCandidate(
                target_table="default.pdw_cvs_usercenter_weixin_id_mapping",
                target_field="user_id",
                source_table="default.pdw_cvs_usercenter_weixin_id_mapping",
                source_field="user_id",
                transform_expression="user_id",
            ),
            FieldLineageCandidate(
                target_table="default.pdw_cvs_usercenter_weixin_id_mapping",
                target_field="user_id",
                source_table="default.ods_cvs_usercenter_weixin_id_mapping_di",
                source_field="user_id",
                transform_expression="user_id",
            ),
        ]
    )

    assert len(grouped) == 1
    assert grouped[0].sources == (
        ("default.ods_cvs_usercenter_weixin_id_mapping_di", "user_id"),
    )


def test_group_approved_rows_keeps_null_constant_fields() -> None:
    grouped = group_approved_rows(
        [
            FieldLineageCandidate(
                target_table="default.plan_price",
                target_field="vendor_code",
                source_table="",
                source_field="",
                transform_expression="NULL",
                transform_explanation="目标字段 vendor_code 由常量 NULL 写入，无上游来源字段。",
                confidence="HIGH",
            )
        ]
    )

    assert len(grouped) == 1
    assert grouped[0].target_field == "vendor_code"
    assert grouped[0].sources == ()
    assert "NULL" in grouped[0].transform_operation


class _FakeFineGrainedLineage:
    def __init__(
        self,
        upstreams: list[str],
        downstreams: list[str],
        transform_operation: str = "",
    ) -> None:
        self.upstreams = upstreams
        self.downstreams = downstreams
        self.transformOperation = transform_operation


def test_merge_fine_grained_lineages_updates_same_downstream_field() -> None:
    existing_keep = _FakeFineGrainedLineage(
        upstreams=["urn:li:schemaField:(source,old_keep)"],
        downstreams=["urn:li:schemaField:(target,keep_field)"],
    )
    existing_replace = _FakeFineGrainedLineage(
        upstreams=["urn:li:schemaField:(source,old_store_id)"],
        downstreams=["urn:li:schemaField:(target,store_id)"],
        transform_operation="old expression",
    )
    new_store_id = _FakeFineGrainedLineage(
        upstreams=["urn:li:schemaField:(source,new_store_id)"],
        downstreams=["urn:li:schemaField:(target,store_id)"],
        transform_operation="new expression",
    )

    merged = merge_fine_grained_lineages(
        [existing_keep, existing_replace],
        [new_store_id],
        clear_existing=False,
    )

    assert merged == [existing_keep, new_store_id]


def test_field_name_from_schema_field_urn_strips_dataset_wrapper() -> None:
    assert (
        _field_name_from_schema_field_urn(
            "urn:li:schemaField:(urn:li:dataset:(urn:li:dataPlatform:hive,"
            "blf-prod-hive.default.dim_store_info,PROD),store_id)"
        )
        == "store_id"
    )


def test_verify_field_lineage_completeness_marks_data_availability_flag(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    patched: dict[str, object] = {}

    monkeypatch.setattr(
        "job_info_sync_datahub.field_lineage_writer.fetch_schema_fields_with_partitions",
        lambda *args, **kwargs: (["store_id", "store_name", "dt", "version"], ["version"]),
    )
    monkeypatch.setattr(
        "job_info_sync_datahub.field_lineage_writer.fetch_structured_properties",
        lambda *args, **kwargs: _structured_properties_payload(
            etl_script="select 1",
            execute_shell="sh run.sh",
            availability_flags=["表血缘", "DDL"],
        ),
    )

    def fake_patch(gms_url: str, dataset_urn: str, flags: list[str], token: str | None = None) -> None:
        patched["dataset_urn"] = dataset_urn
        patched["flags"] = flags

    monkeypatch.setattr(
        "job_info_sync_datahub.field_lineage_writer._patch_data_availability_flags",
        fake_patch,
    )

    dataset_urn = make_hive_dataset_urn("default.dim_store_info")
    result = verify_field_lineage_completeness(
        "http://localhost:8080",
        "default.dim_store_info",
        [
            _FakeFineGrainedLineage(
                upstreams=["urn:li:schemaField:(source,id)"],
                downstreams=[f"urn:li:schemaField:({dataset_urn},store_id)"],
            ),
            _FakeFineGrainedLineage(
                upstreams=["urn:li:schemaField:(source,name)"],
                downstreams=[f"urn:li:schemaField:({dataset_urn},store_name)"],
            ),
        ],
        token="token",
    )

    assert result["is_complete"] is True
    assert result["schema_field_count"] == 2
    assert result["partition_fields"] == ["version"]
    assert result["covered_field_count"] == 2
    assert result["missing_fields"] == []
    assert result["marked_data_availability_flag"] is True
    assert patched["dataset_urn"] == dataset_urn
    assert patched["flags"] == ["DDL", "表血缘", "字段血缘"]


def test_verify_field_lineage_completeness_does_not_mark_when_missing_field(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(
        "job_info_sync_datahub.field_lineage_writer.fetch_schema_fields_with_partitions",
        lambda *args, **kwargs: (["store_id", "store_name", "dt"], []),
    )
    monkeypatch.setattr(
        "job_info_sync_datahub.field_lineage_writer.fetch_structured_properties",
        lambda *args, **kwargs: pytest.fail("incomplete lineage should not read flags"),
    )
    monkeypatch.setattr(
        "job_info_sync_datahub.field_lineage_writer._patch_data_availability_flags",
        lambda *args, **kwargs: pytest.fail("incomplete lineage should not patch flags"),
    )

    dataset_urn = make_hive_dataset_urn("default.dim_store_info")
    result = verify_field_lineage_completeness(
        "http://localhost:8080",
        "default.dim_store_info",
        [
            _FakeFineGrainedLineage(
                upstreams=["urn:li:schemaField:(source,id)"],
                downstreams=[f"urn:li:schemaField:({dataset_urn},store_id)"],
            )
        ],
    )

    assert result["is_complete"] is False
    assert result["schema_field_count"] == 2
    assert result["covered_field_count"] == 1
    assert result["missing_fields"] == ["store_name"]
    assert result["marked_data_availability_flag"] is False


def test_read_fine_grained_lineages_with_retry_waits_for_persisted_fields(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    attempts = 0

    class FakeAspect:
        def __init__(self, fine_grained_lineages: list[_FakeFineGrainedLineage]) -> None:
            self.fineGrainedLineages = fine_grained_lineages

    class FakeGraph:
        def get_aspect(self, *args, **kwargs) -> FakeAspect:
            nonlocal attempts
            attempts += 1
            if attempts == 1:
                return FakeAspect([])
            return FakeAspect(
                [
                    _FakeFineGrainedLineage(
                        upstreams=["urn:li:schemaField:(source,id)"],
                        downstreams=["urn:li:schemaField:(target,store_id)"],
                    )
                ]
            )

    monkeypatch.setattr(
        "job_info_sync_datahub.field_lineage_writer.time.sleep",
        lambda seconds: None,
    )

    result = _read_fine_grained_lineages_with_retry(
        FakeGraph(),
        "urn:li:dataset:(urn:li:dataPlatform:hive,blf-prod-hive.default.dim_store_info,PROD)",
        expected_downstream_fields={"store_id"},
        max_attempts=3,
        sleep_seconds=1,
    )

    assert attempts == 2
    assert len(result) == 1


def test_build_transform_operation_for_ui() -> None:
    assert (
        build_transform_operation_for_ui(
            "coalesce(a, b)",
            "优先取 ods.a 表中的 x 字段，为空时取 ods.b 表中的 y 字段，表示按优先级兜底。",
        )
        == "/* 中文解释：优先取 ods.a 表中的 x 字段，为空时取 ods.b 表中的 y 字段，表示按优先级兜底。 */\n"
        "coalesce(a, b)"
    )


def test_export_cli_skips_when_field_lineage_already_confirmed(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    output = tmp_path / "confirmed.xlsx"
    monkeypatch.setattr(
        "job_info_sync_datahub.field_lineage_cli.fetch_deprecation",
        lambda *args, **kwargs: {},
    )
    monkeypatch.setattr(
        "job_info_sync_datahub.field_lineage_cli.fetch_structured_properties",
        lambda *args, **kwargs: _structured_properties_payload(
            etl_script="select 1",
            execute_shell="sh run.sh",
            availability_flags=["DDL", "表血缘", "字段血缘"],
        ),
    )
    monkeypatch.setattr(
        "job_info_sync_datahub.field_lineage_cli.call_llm_extract_field_lineage",
        lambda *args, **kwargs: pytest.fail("confirmed field lineage should skip LLM"),
    )

    exit_code = field_lineage_cli_main(
        [
            "export",
            "--table",
            "default.dim_store_info",
            "--output",
            str(output),
            "--gms-url",
            "http://localhost:8080",
        ]
    )

    assert exit_code == 3
    assert not output.exists()


def test_export_cli_skips_deprecated_dataset(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    output = tmp_path / "deprecated.xlsx"
    monkeypatch.setattr(
        "job_info_sync_datahub.field_lineage_cli.fetch_deprecation",
        lambda *args, **kwargs: {"deprecation": {"value": {"deprecated": True}}},
    )
    monkeypatch.setattr(
        "job_info_sync_datahub.field_lineage_cli.fetch_structured_properties",
        lambda *args, **kwargs: pytest.fail("deprecated dataset should skip before structuredProperties"),
    )
    monkeypatch.setattr(
        "job_info_sync_datahub.field_lineage_cli.call_llm_extract_field_lineage",
        lambda *args, **kwargs: pytest.fail("deprecated dataset should skip LLM"),
    )

    exit_code = field_lineage_cli_main(
        [
            "export",
            "--table",
            "default.deprecated_table",
            "--output",
            str(output),
            "--gms-url",
            "http://localhost:8080",
        ]
    )

    assert exit_code == 3
    assert not output.exists()


@pytest.mark.skipif(
    importlib.util.find_spec("datahub") is None,
    reason="acryl-datahub not installed",
)
def test_build_fine_grained_allows_constant_target_field_without_upstreams() -> None:
    group = group_approved_rows(
        [
            FieldLineageCandidate(
                target_table="default.plan_price",
                target_field="vendor_code",
                source_table="",
                source_field="",
                transform_expression="NULL",
                transform_explanation="目标字段 vendor_code 由常量 NULL 写入，无上游来源字段。",
                confidence="HIGH",
            )
        ]
    )[0]

    fg = build_fine_grained_lineage_class(group)

    assert fg.upstreams == []
    assert len(fg.downstreams) == 1
    assert "NULL" in fg.transformOperation


@pytest.mark.skipif(
    importlib.util.find_spec("datahub") is None,
    reason="acryl-datahub not installed",
)
def test_build_fine_grained_allows_offline_source_without_upstreams() -> None:
    group = group_approved_rows(
        [
            FieldLineageCandidate(
                target_table="data_factory.target",
                target_field="manual_store_name",
                source_table="",
                source_field="",
                transform_expression="pandas.read_csv('/data/manual/store.csv')['store_name']",
                transform_explanation=(
                    "字段来自本地 CSV 文件 /data/manual/store.csv 的 store_name 列，"
                    "无法映射到 Hive 物理字段。"
                ),
                confidence="MEDIUM",
            )
        ]
    )[0]

    fg = build_fine_grained_lineage_class(group)

    assert fg.upstreams == []
    assert len(fg.downstreams) == 1
    assert "read_csv" in fg.transformOperation


@pytest.mark.skipif(
    importlib.util.find_spec("datahub") is None,
    reason="acryl-datahub not installed",
)
def test_build_fine_grained_sets_transform_operation() -> None:
    group = group_approved_rows(
        [
            FieldLineageCandidate(
                target_table="default.dim_store_info",
                target_field="store_id",
                source_table="ods.store_info",
                source_field="id",
                transform_expression="cast(id as bigint)",
                transform_explanation="取 ods.store_info 表中的 id 字段，表示将门店 ID 转为 bigint 类型写入 store_id。",
            )
        ]
    )[0]
    fg = build_fine_grained_lineage_class(group)
    assert fg.transformOperation == (
        "/* 中文解释：取 ods.store_info 表中的 id 字段，表示将门店 ID 转为 bigint 类型写入 store_id。 */\n"
        "cast(id as bigint)"
    )
    assert fg.query is None
    assert len(fg.upstreams) == 1


def test_import_reviewed_cli_writes_approved_plan(
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
) -> None:
    workbook = tmp_path / "review.xlsx"
    output = tmp_path / "import_plan.json"
    source = FieldLineageInput(
        dataset_urn=make_hive_dataset_urn("default.dim_store_info"),
        table_name="default.dim_store_info",
        etl_script="select 1",
        execute_shell="sh run.sh",
    )
    write_candidate_workbook(
        workbook,
        source_input=source,
        candidates=[
            FieldLineageCandidate(
                target_table="default.dim_store_info",
                target_field="store_id",
                source_table="ods.store_info",
                source_field="id",
                transform_expression="cast(id as bigint)",
                confidence="HIGH",
            )
        ],
        unresolved_fields=[],
        llm_model="deepseek-test",
    )
    wb = load_workbook(workbook)
    wb["candidate_lineage"]["A2"] = "APPROVED"
    wb.save(workbook)

    exit_code = field_lineage_cli_main(
        ["import-reviewed", "--input", str(workbook), "--output", str(output)]
    )

    assert exit_code == 0
    payload = json.loads(output.read_text(encoding="utf-8"))
    assert payload["dry_run"] is True
    assert payload["approved_rows"] == 1
    assert payload["rows"][0]["target_field"] == "store_id"
    assert "write_result" in payload
    table_result = payload["write_result"]["tables"]["default.dim_store_info"]
    assert (
        table_result["replacement_mode"]
        == "clear_import_replace_all_fine_grained_lineages_for_table"
    )
    stdout = capsys.readouterr().out
    assert '"rows"' not in stdout
    assert "import_plan=" in stdout


def test_import_reviewed_cli_can_limit_import_statuses(tmp_path: Path) -> None:
    workbook = tmp_path / "review.xlsx"
    output = tmp_path / "import_plan.json"
    source = FieldLineageInput(
        dataset_urn=make_hive_dataset_urn("default.dim_store_info"),
        table_name="default.dim_store_info",
        etl_script="select 1",
        execute_shell="sh run.sh",
    )
    write_candidate_workbook(
        workbook,
        source_input=source,
        candidates=[
            FieldLineageCandidate(
                target_table="default.dim_store_info",
                target_field="auto_approved_field",
                source_table="ods.store_info",
                source_field="id",
                transform_expression="cast(id as bigint)",
                confidence="HIGH",
            ),
            FieldLineageCandidate(
                target_table="default.dim_store_info",
                target_field="manual_approved_field",
                source_table="ods.store_info",
                source_field="code",
                transform_expression="code",
                confidence="MEDIUM",
                review_status=FieldLineageReviewStatus.APPROVED,
            ),
        ],
        unresolved_fields=[],
        llm_model="deepseek-test",
    )

    exit_code = field_lineage_cli_main(
        [
            "import-reviewed",
            "--input",
            str(workbook),
            "--output",
            str(output),
            "--import-statuses",
            "APPROVED",
        ]
    )

    assert exit_code == 0
    payload = json.loads(output.read_text(encoding="utf-8"))
    assert payload["import_statuses"] == ["APPROVED"]
    assert payload["approved_rows"] == 1
    assert payload["rows"][0]["target_field"] == "manual_approved_field"


def test_import_reviewed_cli_requires_full_auto_approved(tmp_path: Path) -> None:
    from job_info_sync_datahub.field_lineage_cli import EXIT_REQUIRES_REVIEW

    workbook = tmp_path / "review.xlsx"
    output = tmp_path / "import_plan.json"
    source = FieldLineageInput(
        dataset_urn=make_hive_dataset_urn("default.dim_store_info"),
        table_name="default.dim_store_info",
        etl_script="select 1",
        execute_shell="sh run.sh",
    )
    write_candidate_workbook(
        workbook,
        source_input=source,
        candidates=[
            FieldLineageCandidate(
                target_table="default.dim_store_info",
                target_field="auto_approved_field",
                source_table="ods.store_info",
                source_field="id",
                transform_expression="cast(id as bigint)",
                confidence="HIGH",
            ),
            FieldLineageCandidate(
                target_table="default.dim_store_info",
                target_field="needs_review_field",
                source_table="ods.store_info",
                source_field="id,name",
                transform_expression="concat(id, name)",
                confidence="HIGH",
            ),
        ],
        unresolved_fields=[],
        llm_model="deepseek-test",
    )

    exit_code = field_lineage_cli_main(
        [
            "import-reviewed",
            "--input",
            str(workbook),
            "--output",
            str(output),
            "--import-statuses",
            "AUTO_APPROVED",
            "--require-full-auto-approved",
        ]
    )

    assert exit_code == EXIT_REQUIRES_REVIEW
    assert not output.exists()


def test_import_reviewed_cli_allows_full_auto_approved(tmp_path: Path) -> None:
    workbook = tmp_path / "review.xlsx"
    output = tmp_path / "import_plan.json"
    source = FieldLineageInput(
        dataset_urn=make_hive_dataset_urn("default.dim_store_info"),
        table_name="default.dim_store_info",
        etl_script="select 1",
        execute_shell="sh run.sh",
    )
    write_candidate_workbook(
        workbook,
        source_input=source,
        candidates=[
            FieldLineageCandidate(
                target_table="default.dim_store_info",
                target_field="auto_approved_field",
                source_table="ods.store_info",
                source_field="id",
                transform_expression="cast(id as bigint)",
                confidence="HIGH",
            ),
        ],
        unresolved_fields=[],
        llm_model="deepseek-test",
    )

    exit_code = field_lineage_cli_main(
        [
            "import-reviewed",
            "--input",
            str(workbook),
            "--output",
            str(output),
            "--import-statuses",
            "AUTO_APPROVED",
            "--require-full-auto-approved",
        ]
    )

    assert exit_code == 0
    payload = json.loads(output.read_text(encoding="utf-8"))
    assert payload["approved_rows"] == 1
    assert payload["auto_approved_percent"] == 1.0
    assert payload["unresolved_field_count"] == 0


def test_import_reviewed_write_without_approved_exits_4(tmp_path: Path) -> None:
    from job_info_sync_datahub.field_lineage_cli import EXIT_NO_APPROVED_ROWS

    workbook = tmp_path / "review.xlsx"
    source = FieldLineageInput(
        dataset_urn=make_hive_dataset_urn("default.dim_store_info"),
        table_name="default.dim_store_info",
        etl_script="select 1",
        execute_shell="sh run.sh",
    )
    write_candidate_workbook(
        workbook,
        source_input=source,
            candidates=[
                FieldLineageCandidate(
                    target_table="default.dim_store_info",
                    target_field="store_id",
                    source_table="ods.store_info",
                    source_field="id",
                    confidence="MEDIUM",
                )
            ],
            unresolved_fields=[],
            llm_model="deepseek-test",
        )

    exit_code = field_lineage_cli_main(
        [
            "import-reviewed",
            "--input",
            str(workbook),
            "--write",
            "--gms-url",
            "http://localhost:8080",
        ]
    )

    assert exit_code == EXIT_NO_APPROVED_ROWS


def test_import_reviewed_write_blocks_confirmed_field_lineage(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    workbook = tmp_path / "review.xlsx"
    source = FieldLineageInput(
        dataset_urn=make_hive_dataset_urn("default.dim_store_info"),
        table_name="default.dim_store_info",
        etl_script="select 1",
        execute_shell="sh run.sh",
    )
    write_candidate_workbook(
        workbook,
        source_input=source,
        candidates=[
            FieldLineageCandidate(
                target_table="default.dim_store_info",
                target_field="store_id",
                source_table="ods.store_info",
                source_field="id",
                transform_expression="cast(id as bigint)",
                confidence="HIGH",
            )
        ],
        unresolved_fields=[],
        llm_model="deepseek-test",
    )
    monkeypatch.setattr(
        "job_info_sync_datahub.field_lineage_cli.fetch_structured_properties",
        lambda *args, **kwargs: _structured_properties_payload(
            etl_script="select 1",
            execute_shell="sh run.sh",
            availability_flags=["DDL", "表血缘", "字段血缘"],
        ),
    )

    exit_code = field_lineage_cli_main(
        [
            "import-reviewed",
            "--input",
            str(workbook),
            "--write",
            "--gms-url",
            "http://localhost:8080",
        ]
    )

    assert exit_code == 5


@pytest.mark.skipif(
    importlib.util.find_spec("datahub") is None,
    reason="acryl-datahub not installed",
)
def test_merge_fine_grained_lineages_can_clear_or_keep_existing() -> None:
    old = build_fine_grained_lineage_class(
        group_approved_rows(
            [
                FieldLineageCandidate(
                    target_table="default.dim_store_info",
                    target_field="old_field",
                    source_table="ods.store_info",
                    source_field="old_id",
                )
            ]
        )[0]
    )
    new = build_fine_grained_lineage_class(
        group_approved_rows(
            [
                FieldLineageCandidate(
                    target_table="default.dim_store_info",
                    target_field="new_field",
                    source_table="ods.store_info",
                    source_field="new_id",
                )
            ]
        )[0]
    )

    assert merge_fine_grained_lineages([old], [new], clear_existing=True) == [new]
    assert merge_fine_grained_lineages([old], [new], clear_existing=False) == [old, new]
