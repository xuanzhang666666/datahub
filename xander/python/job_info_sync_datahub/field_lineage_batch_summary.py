"""Batch summary workbook for field-lineage review exports."""

from __future__ import annotations

import argparse
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, Iterable, List

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


SUMMARY_HEADERS = [
    "table_name",
    "workbook",
    "total_candidate_count",
    "auto_approved_count",
    "auto_approved_percent",
    "approved_count",
    "needs_review_count",
    "pending_count",
    "rejected_count",
    "needs_fix_count",
    "other_status_count",
    "unresolved_field_count",
]


@dataclass(frozen=True)
class TableLineageReviewSummary:
    table_name: str
    workbook: str
    total_candidate_count: int
    status_counts: Dict[str, int]
    unresolved_field_count: int

    @property
    def auto_approved_count(self) -> int:
        return self.status_counts.get("AUTO_APPROVED", 0)

    @property
    def auto_approved_percent(self) -> float:
        if self.total_candidate_count == 0:
            return 0
        return self.auto_approved_count / self.total_candidate_count

    @property
    def is_fully_auto_approved(self) -> bool:
        return (
            self.total_candidate_count > 0
            and self.auto_approved_percent == 1
            and self.unresolved_field_count == 0
        )


def _style_sheet(ws) -> None:
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9E1F2")
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for idx, _ in enumerate(ws[1], start=1):
        letter = get_column_letter(idx)
        max_len = 12
        for cell in ws[letter]:
            if cell.value:
                max_len = min(max(max_len, len(str(cell.value))), 80)
        ws.column_dimensions[letter].width = max_len + 2
    ws.freeze_panes = "A2"
    headers = [cell.value for cell in ws[1]]
    if "auto_approved_percent" in headers:
        percent_col = headers.index("auto_approved_percent") + 1
        for cell in ws.iter_cols(
            min_col=percent_col,
            max_col=percent_col,
            min_row=2,
        ):
            for item in cell:
                item.number_format = "0.00%"


def _candidate_records(path: Path) -> List[Dict[str, str]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    if "candidate_lineage" not in wb.sheetnames:
        return []
    ws = wb["candidate_lineage"]
    rows = ws.iter_rows(values_only=True)
    headers = [str(value or "").strip() for value in next(rows, [])]
    records = []
    for row in rows:
        rec = {}
        for idx, header in enumerate(headers):
            value = row[idx] if idx < len(row) else ""
            rec[header] = "" if value is None else str(value).strip()
        records.append(rec)
    return records


def _unresolved_field_count(path: Path) -> int:
    wb = load_workbook(path, read_only=True, data_only=True)
    if "unresolved_fields" not in wb.sheetnames:
        return 0
    ws = wb["unresolved_fields"]
    rows = ws.iter_rows(values_only=True)
    headers = [str(value or "").strip() for value in next(rows, [])]
    try:
        target_field_idx = headers.index("target_field")
    except ValueError:
        target_field_idx = 0
    count = 0
    for row in rows:
        value = row[target_field_idx] if target_field_idx < len(row) else ""
        if value is not None and str(value).strip():
            count += 1
    return count


def _table_name_from_records(path: Path, records: List[Dict[str, str]]) -> str:
    for rec in records:
        table = rec.get("target_table", "").strip()
        if table:
            return table
    name = path.stem
    if "_" in name:
        return name.split("_", 1)[1]
    return name


def summarize_review_workbook(path: Path) -> TableLineageReviewSummary:
    records = _candidate_records(path)
    status_counts: Dict[str, int] = {}
    for rec in records:
        if not rec.get("target_field", "").strip():
            continue
        status = rec.get("review_status", "").strip().upper() or "<EMPTY>"
        status_counts[status] = status_counts.get(status, 0) + 1
    return TableLineageReviewSummary(
        table_name=_table_name_from_records(path, records),
        workbook=str(path),
        total_candidate_count=sum(status_counts.values()),
        status_counts=status_counts,
        unresolved_field_count=_unresolved_field_count(path),
    )


def _summary_to_row(summary: TableLineageReviewSummary) -> List[object]:
    known_statuses = {
        "AUTO_APPROVED",
        "APPROVED",
        "NEEDS_REVIEW",
        "PENDING",
        "REJECTED",
        "NEEDS_FIX",
    }
    other_count = sum(
        count
        for status, count in summary.status_counts.items()
        if status not in known_statuses
    )
    return [
        summary.table_name,
        summary.workbook,
        summary.total_candidate_count,
        summary.auto_approved_count,
        summary.auto_approved_percent,
        summary.status_counts.get("APPROVED", 0),
        summary.status_counts.get("NEEDS_REVIEW", 0),
        summary.status_counts.get("PENDING", 0),
        summary.status_counts.get("REJECTED", 0),
        summary.status_counts.get("NEEDS_FIX", 0),
        other_count,
        summary.unresolved_field_count,
    ]


def write_batch_summary_workbook(batch_dir: Path, output: Path) -> None:
    summaries = [
        summarize_review_workbook(path)
        for path in sorted(batch_dir.glob("*.xlsx"))
        if not path.name.endswith("_summary.xlsx")
    ]
    write_summaries(output, summaries)


def write_summaries(output: Path, summaries: Iterable[TableLineageReviewSummary]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "table_summary"
    ws.append(SUMMARY_HEADERS)
    for summary in summaries:
        ws.append(_summary_to_row(summary))
    _style_sheet(ws)
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)


def main(argv: List[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--batch-dir", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args(argv)
    write_batch_summary_workbook(args.batch_dir, args.output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
