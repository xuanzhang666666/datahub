#!/usr/bin/env python3
"""将 Hive 表清单 xlsx（首行为表头）切分为多个 xlsx。

默认 **按库拆分**（--mode db）：每个 Hive 库一个输出文件，文件名 ``{prefix}.bydb.{db}.xlsx``。
可选 **按行数拆分**（--mode rows）：与旧版一致，每文件最多 --chunk-size 条数据行。

默认排除库名 ``data_finance``（--exclude-dbs，空字符串表示不排除）。

识别库名列：与 dim 导出一致时读 ``db_name`` 列；否则若有 ``fqtn`` 列则取 ``db.table`` 的库前缀。

依赖: pip install openpyxl
"""

from __future__ import annotations

import argparse
import re
import sys
from collections import defaultdict
from pathlib import Path
from typing import Any, DefaultDict, List, Sequence, Set, Tuple

_SAFE_DB = re.compile(r"[^a-zA-Z0-9_-]+")


def cell_str(val: Any) -> str:
    if val is None:
        return ""
    if isinstance(val, float) and val == int(val):
        return str(int(val))
    return str(val).strip()


def _dim_export_header_row(cells: List[str]) -> bool:
    if len(cells) < 5:
        return False
    return (
        cells[0].lower() == "db_id"
        and cells[1].lower() == "db_name"
        and cells[2].lower() == "table_id"
        and cells[3].lower() == "table_name"
        and cells[4].lower() == "fqtn"
    )


def _has_fqtn_header(cells: List[str]) -> bool:
    return any(c.lower() == "fqtn" for c in cells)


def detect_fqtn_column_1based(header_cells: List[str]) -> int:
    for i, c in enumerate(header_cells):
        if c.lower() == "fqtn":
            return i + 1
    if _dim_export_header_row(header_cells):
        return 5
    return 1


def detect_db_column_1based(header_cells: List[str]) -> int | None:
    if _dim_export_header_row(header_cells):
        return 2
    for i, c in enumerate(header_cells):
        if c.lower() == "db_name":
            return i + 1
    return None


def _parse_exclude_dbs(s: str) -> Set[str]:
    t = (s or "").strip()
    if not t:
        return set()
    return {x.strip() for x in t.split(",") if x.strip()}


def _sanitize_db_for_filename(db: str) -> str:
    s = _SAFE_DB.sub("_", db.strip())
    return s or "unknown_db"


def parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--in", dest="in_path", type=Path, required=True, help="输入 .xlsx")
    p.add_argument(
        "--out-dir",
        type=Path,
        required=True,
        help="输出目录（将创建）",
    )
    p.add_argument(
        "--mode",
        choices=("db", "rows"),
        default="db",
        help="db=按 Hive 库各一个文件；rows=按固定行数分块（旧行为）",
    )
    p.add_argument(
        "--chunk-size",
        type=int,
        default=1000,
        help="mode=rows 时每个输出文件的数据行数上限（不含表头），默认 1000",
    )
    p.add_argument(
        "--prefix",
        type=str,
        default="",
        help="输出文件名前缀，默认使用输入文件 stem",
    )
    p.add_argument(
        "--exclude-dbs",
        type=str,
        default="data_finance",
        metavar="CSV",
        help="逗号分隔库名，整行跳过。空字符串不排除。",
    )
    return p.parse_args()


def _normalize_row(row: Tuple[Any, ...], width: int) -> List[Any]:
    cells = list(row)
    if len(cells) < width:
        cells.extend([None] * (width - len(cells)))
    return cells[:width]


def _write_chunk(
    path: Path,
    header: Sequence[Any],
    rows: Sequence[Sequence[Any]],
) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(list(header))
    for r in rows:
        ws.append(list(r))
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    wb.close()


def _row_db(
    row: List[Any],
    db_col_1: int | None,
    fqtn_col_1: int,
) -> str | None:
    if db_col_1 is not None:
        i = db_col_1 - 1
        if i < len(row):
            d = cell_str(row[i])
            return d or None
    fi = fqtn_col_1 - 1
    if fi >= len(row):
        return None
    fq = cell_str(row[fi])
    if "." not in fq:
        return None
    return fq.split(".", 1)[0] or None


def main() -> int:
    args = parse_args()
    exclude = _parse_exclude_dbs(args.exclude_dbs)
    if args.chunk_size < 1 and args.mode == "rows":
        print("ERROR: --chunk-size 须 >= 1", file=sys.stderr)
        return 2
    if not args.in_path.is_file():
        print(f"ERROR: 输入不存在: {args.in_path}", file=sys.stderr)
        return 2

    try:
        from openpyxl import load_workbook
    except ImportError:
        print("ERROR: pip install openpyxl", file=sys.stderr)
        return 2

    prefix = args.prefix.strip() or args.in_path.stem
    out_dir = args.out_dir

    wb = load_workbook(args.in_path, read_only=True, data_only=True)
    part = 0
    skipped_exclude = 0
    try:
        ws = wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        header = next(it, None)
        if not header:
            print("ERROR: 空表", file=sys.stderr)
            return 2
        width = len(header)
        header_strs = [cell_str(x) for x in header]
        db_col_1 = detect_db_column_1based(header_strs)
        fqtn_col_1 = detect_fqtn_column_1based(header_strs)

        if args.mode == "db":
            groups: DefaultDict[str, List[List[Any]]] = defaultdict(list)
            for row in it:
                if row is None:
                    continue
                norm = _normalize_row(tuple(row), width)
                db = _row_db(norm, db_col_1, fqtn_col_1)
                if not db:
                    continue
                if db in exclude:
                    skipped_exclude += 1
                    continue
                groups[db].append(norm)
            if not groups:
                print("ERROR: 按库拆分后无数据行（可能全部被 exclude）", file=sys.stderr)
                return 2
            for db in sorted(groups.keys()):
                part += 1
                safe = _sanitize_db_for_filename(db)
                out = out_dir / f"{prefix}.bydb.{safe}.xlsx"
                _write_chunk(out, header, groups[db])
                print(out, flush=True)
        else:
            buf: List[List[Any]] = []
            for row in it:
                if row is None:
                    continue
                norm = _normalize_row(tuple(row), width)
                db = _row_db(norm, db_col_1, fqtn_col_1)
                if db and db in exclude:
                    skipped_exclude += 1
                    continue
                buf.append(norm)
                if len(buf) >= args.chunk_size:
                    part += 1
                    out = out_dir / f"{prefix}.part{part:04d}.xlsx"
                    _write_chunk(out, header, buf)
                    print(out, flush=True)
                    buf = []
            if buf:
                part += 1
                out = out_dir / f"{prefix}.part{part:04d}.xlsx"
                _write_chunk(out, header, buf)
                print(out, flush=True)
    finally:
        wb.close()

    if part == 0:
        print("ERROR: 无输出文件", file=sys.stderr)
        return 2

    ex_msg = ",".join(sorted(exclude)) if exclude else "-"
    print(
        f"[split_hive_tables_xlsx] mode={args.mode} files={part} "
        f"skipped_exclude={skipped_exclude} exclude_dbs={ex_msg} -> {out_dir}",
        file=sys.stderr,
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
