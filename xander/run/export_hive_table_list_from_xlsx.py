#!/usr/bin/env python3
"""从 .xlsx 导出 Hive 全表名列表（每行 db.table），供 ingest_hive_table_list_to_datahub.sh 使用。

与 dim_dc_hive_meta_info_di 导出列一致时（db_id, db_name, table_id, table_name, fqtn）:
  默认 --fqtn-col auto：识别表头并只读 fqtn 列（第 E 列），数据从第 2 行起。

其它:
  - 单列全名: --fqtn-col N（1=A）或 auto
  - 双列: --db-col / --table-col（1-based）

默认排除库 ``data_finance``（--exclude-dbs，传空字符串 ``--exclude-dbs ''`` 不排除）。

依赖: pip install openpyxl
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path
from typing import Any, Iterable, List, Set, Tuple

_FQTN_RE = re.compile(r"^\S+\.\S+$")


def _parse_exclude_dbs(s: str) -> Set[str]:
    """逗号分隔库名；空字符串表示不排除。"""
    t = (s or "").strip()
    if not t:
        return set()
    return {x.strip() for x in t.split(",") if x.strip()}


def cell_str(val: Any) -> str:
    if val is None:
        return ""
    if isinstance(val, float) and val == int(val):
        return str(int(val))
    return str(val).strip()


def is_likely_header_row(cells: List[str]) -> bool:
    blob = " ".join(c.lower() for c in cells if c)
    keys = (
        "database",
        "db_name",
        "dbname",
        "table_name",
        "tablename",
        "schema",
        "库名",
        "表名",
    )
    return any(k in blob for k in keys)


def _dim_export_header_row(cells: List[str]) -> bool:
    """与 SQL 导出列名一致的首行。"""
    if len(cells) < 5:
        return False
    return (
        cells[0].lower() == "db_id"
        and cells[1].lower() == "db_name"
        and cells[2].lower() == "table_id"
        and cells[3].lower() == "table_name"
        and cells[4].lower() == "fqtn"
    )


def _has_fqtn_header(cells: List[str]) -> bool:
    return any(c.lower() == "fqtn" for c in cells)


def detect_fqtn_column_1based(header_cells: List[str]) -> int:
    for i, c in enumerate(header_cells):
        if c.lower() == "fqtn":
            return i + 1
    if _dim_export_header_row(header_cells):
        return 5
    return 1


def detect_data_start_row(header_cells: List[str]) -> int:
    if _dim_export_header_row(header_cells) or _has_fqtn_header(header_cells):
        return 2
    if is_likely_header_row(header_cells):
        return 2
    return 1


def parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--in", dest="in_path", type=Path, required=True, help="输入 .xlsx")
    p.add_argument(
        "--out",
        type=Path,
        required=True,
        help="输出 .txt，每行一个 db.table",
    )
    p.add_argument(
        "--sheet",
        default="0",
        help="工作表：0-based 索引或工作表名称，默认 0",
    )
    p.add_argument(
        "--fqtn-col",
        type=str,
        default="auto",
        metavar="N|auto",
        help="单列模式：fqtn 所在列 1=A；auto 时读首行表头识别（dim 导出为第 5 列）",
    )
    p.add_argument(
        "--db-col",
        type=int,
        default=None,
        metavar="N",
        help="双列模式：库名列（1-based），与 --table-col 同时指定",
    )
    p.add_argument(
        "--table-col",
        type=int,
        default=None,
        metavar="N",
        help="双列模式：表名列（1-based）",
    )
    p.add_argument(
        "--max-rows",
        type=int,
        default=0,
        help="最多处理数据行数（0 表示不限制），便于试跑",
    )
    p.add_argument(
        "--exclude-dbs",
        type=str,
        default="data_finance",
        metavar="CSV",
        help="逗号分隔的 Hive 库名，命中则跳过该行（fqtn 的库前缀或 db 列）。"
        "传空字符串不排除，例如: --exclude-dbs ''",
    )
    return p.parse_args()


def _resolve_sheet(wb: Any, spec: str) -> Any:
    if spec.isdigit():
        idx = int(spec)
        if idx < 0 or idx >= len(wb.worksheets):
            raise SystemExit(f"ERROR: sheet 索引越界: {idx}")
        return wb.worksheets[idx]
    for ws in wb.worksheets:
        if ws.title == spec:
            return ws
    names = [ws.title for ws in wb.worksheets]
    raise SystemExit(f"ERROR: 未找到工作表 {spec!r}，可选: {names}")


def _col_letter_to_idx1(col: int) -> int:
    if col < 1:
        raise SystemExit("ERROR: 列号须 >= 1（1=A）")
    return col


def _parse_fqtn_spec(spec: str) -> object:
    s = spec.strip().lower()
    if s == "auto":
        return "AUTO"
    try:
        return int(spec.strip())
    except ValueError as e:
        raise SystemExit(f"ERROR: --fqtn-col 须为 auto 或整数: {spec!r}") from e


def _iter_rows_two_col(
    ws: Any,
    db_col: int,
    table_col: int,
    max_rows: int,
    min_row: int,
) -> Iterable[Tuple[str, str]]:
    count = 0
    for row in ws.iter_rows(min_row=min_row, values_only=True):
        if not row:
            continue
        cells = list(row)
        db_i = db_col - 1
        tbl_i = table_col - 1
        if db_i >= len(cells) or tbl_i >= len(cells):
            continue
        db = cell_str(cells[db_i])
        tbl = cell_str(cells[tbl_i])
        if not db or not tbl:
            continue
        yield db, tbl
        count += 1
        if max_rows and count >= max_rows:
            break


def _iter_rows_fqtn(
    ws: Any, col1: int, max_rows: int, min_row: int
) -> Iterable[str]:
    count = 0
    cidx = col1 - 1
    for row in ws.iter_rows(min_row=min_row, values_only=True):
        if not row or cidx >= len(row):
            continue
        s = cell_str(row[cidx])
        if not s or s.startswith("#"):
            continue
        yield s
        count += 1
        if max_rows and count >= max_rows:
            break


def _header_cells(ws: Any) -> List[str]:
    first = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
    if not first:
        return []
    return [cell_str(x) for x in first]


def main() -> int:
    args = parse_args()
    if not args.in_path.is_file():
        print(f"ERROR: 输入文件不存在: {args.in_path}", file=sys.stderr)
        return 2

    try:
        from openpyxl import load_workbook
    except ImportError:
        print("ERROR: 请安装 openpyxl: pip install openpyxl", file=sys.stderr)
        return 2

    wb = load_workbook(args.in_path, read_only=True, data_only=True)
    use_two = args.db_col is not None and args.table_col is not None
    if (args.db_col is None) ^ (args.table_col is None):
        print("ERROR: --db-col 与 --table-col 须同时指定或同时省略", file=sys.stderr)
        return 2

    db_col = _col_letter_to_idx1(args.db_col) if args.db_col else 0
    table_col = _col_letter_to_idx1(args.table_col) if args.table_col else 0
    fqtn_spec = _parse_fqtn_spec(args.fqtn_col)

    exclude_dbs = _parse_exclude_dbs(args.exclude_dbs)
    seen: Set[str] = set()
    out_lines: List[str] = []
    skipped_header = False
    fqtn_col = 1
    data_start = 1
    skipped_exclude = 0

    try:
        ws = _resolve_sheet(wb, str(args.sheet).strip())
        header = _header_cells(ws)

        if use_two:
            if fqtn_spec != "AUTO":
                print(
                    "WARN: 双列模式下忽略 --fqtn-col",
                    file=sys.stderr,
                )
            data_start = (
                2
                if len(header) >= 2 and is_likely_header_row(header[:2])
                else 1
            )
            skipped_header = data_start == 2
            raw: List[Tuple[str, str]] = []
            for db, tbl in _iter_rows_two_col(
                ws, db_col, table_col, args.max_rows, data_start
            ):
                raw.append((db, tbl))
            for db, tbl in raw:
                if db in exclude_dbs:
                    skipped_exclude += 1
                    continue
                line = f"{db}.{tbl}"
                if line in seen:
                    continue
                seen.add(line)
                out_lines.append(line)
        else:
            if fqtn_spec == "AUTO":
                fqtn_col = detect_fqtn_column_1based(header)
                data_start = detect_data_start_row(header)
            else:
                fqtn_col = int(fqtn_spec)
                hci = fqtn_col - 1
                hval = header[hci] if hci < len(header) else ""
                if hval and not _FQTN_RE.match(hval):
                    data_start = 2
                else:
                    data_start = 1
            skipped_header = data_start >= 2

            raw_s: List[str] = []
            for s in _iter_rows_fqtn(ws, fqtn_col, args.max_rows, data_start):
                raw_s.append(s)
            for s in raw_s:
                if not _FQTN_RE.match(s):
                    continue
                db_part = s.split(".", 1)[0]
                if db_part in exclude_dbs:
                    skipped_exclude += 1
                    continue
                if s in seen:
                    continue
                seen.add(s)
                out_lines.append(s)
    finally:
        wb.close()

    if not out_lines:
        print("ERROR: 未解析到任何 db.table 行，请检查列号/表头", file=sys.stderr)
        return 2

    args.out.parent.mkdir(parents=True, exist_ok=True)
    args.out.write_text("\n".join(out_lines) + "\n", encoding="utf-8")
    fqtn_col_disp = str(fqtn_col) if not use_two else "-"
    ex_msg = ",".join(sorted(exclude_dbs)) if exclude_dbs else "-"
    print(
        f"[export_hive_table_list_from_xlsx] rows={len(out_lines)} "
        f"unique={len(seen)} skipped_exclude={skipped_exclude} exclude_dbs={ex_msg} "
        f"fqtn_col={fqtn_col_disp} data_start_row={data_start} "
        f"skip_header={skipped_header} -> {args.out}",
        file=sys.stderr,
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
